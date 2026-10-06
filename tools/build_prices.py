import json,datetime as dt,math,statistics
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
TICK=[('NasdaqGM:IGF','IGF','iShares Global Infrastructure ETF','Infrastructure cotée monde','ETF','C1'),
('LSE:INFR','INFR.L','iShares Global Infrastructure UCITS ETF','Infrastructure cotée monde (UCITS)','ETF','C1'),
('LSE:3IN','3IN.L','3i Infrastructure plc','Infrastructure core-plus, trust coté','Investment trust','C1'),
('LSE:HICL','HICL.L','HICL Infrastructure plc','Infrastructure core (PPP, régulé), trust coté','Investment trust','C1'),
('LSE:INPP','INPP.L','International Public Partnerships Ltd','Infrastructure core (PPP, régulé), trust coté','Investment trust','C1'),
('LSE:SEQI','SEQI.L','Sequoia Economic Infrastructure Income Fund','Dette d\'infrastructure, trust coté','Investment trust','C2'),
('LSE:GCP','GCP.L','GCP Infrastructure Investments Ltd','Dette d\'infrastructure (UK), trust coté','Investment trust','C2'),
('NYSE:BIP','BIP','Brookfield Infrastructure Partners L.P.','Infrastructure value add, gérant coté','LP cotée','C1'),
('KOSE:A088980','088980.KS','Macquarie Korea Infrastructure Fund','Infrastructure core Corée, fonds Macquarie coté','Fonds coté','C1'),
('ARCA:BKLN','BKLN','Invesco Senior Loan ETF','Prêts seniors à taux variable (leveraged loans)','ETF','C4'),
('ARCA:HYG','HYG','iShares iBoxx $ High Yield Corporate Bond ETF','Haut rendement USD','ETF','C4'),
('LSE:IHYG','IHYG.L','iShares € High Yield Corp Bond UCITS ETF','Haut rendement EUR','ETF','C4'),
('NasdaqGS:ARCC','ARCC','Ares Capital Corporation','Direct lending, BDC cotée (Ares)','BDC','C3'),
('ARCA:BIZD','BIZD','VanEck BDC Income ETF','Direct lending, panier de BDC','ETF','C3')]
N=len(TICK)
COMP=[('C1','Infrastructure cotée, capital','IGF, INFR, 3IN, HICL, INPP, A088980 ; BIP exclu'),
      ('C2','Dette d\'infrastructure cotée','SEQI, GCP'),
      ('C3','Direct lending coté','ARCC, BIZD'),
      ('C4','Poche liquide : haut rendement et prêts seniors','BKLN, HYG, IHYG')]
STATS_FROM={'INFR.L':dt.date(2016,1,1)}
EXCLUDE={'NYSE:BIP'}
def load(fn):
    j=json.load(open(fn))['chart']['result'][0]
    meta=j['meta']; ts=j['timestamp']; q=j['indicators']['quote'][0]
    adj=j['indicators'].get('adjclose',[{}])[0].get('adjclose',q['close'])
    off=meta.get('gmtoffset',0)
    return meta,[(dt.datetime.fromtimestamp(t+off,dt.UTC).date(),c,a) for t,c,a in zip(ts,q['close'],adj) if c is not None and a is not None]
data={}
for user,ysym,name,strat,typ,comp in TICK:
    mm,mrows=load(f'm_{ysym}.json'); dm,drows=load(f'd_{ysym}.json')
    cur=mm['currency']; scale=0.01 if cur=='GBp' else 1; cur_out='GBP' if cur=='GBp' else cur
    last_date,last,_=drows[-1]
    def ret(days,idx=1):
        target=last_date-dt.timedelta(days=days)
        past=[r for r in drows if r[0]<=target] or [r for r in drows if r[0]<=target+dt.timedelta(days=10)]
        return (drows[-1][idx]/past[-1][idx]-1) if past else None
    a=[r[2] for r in drows]; lr=[math.log(a[i]/a[i-1]) for i in range(1,len(a))]
    ma=[(d,c*scale,adj*scale) for d,c,adj in mrows]
    ma_c=[r for r in ma if (r[0].year,r[0].month)!=(last_date.year,last_date.month)]
    data[user]=dict(ysym=ysym,name=name,strat=strat,typ=typ,comp=comp,cur=cur_out,last_date=last_date,last=last*scale,prev=drows[-2][1]*scale,
        hi52=max(r[1] for r in drows)*scale,lo52=min(r[1] for r in drows)*scale,r1m=ret(30),r3m=ret(91),r1y=ret(365),tr1y=ret(365,2),
        vol_d=statistics.pstdev(lr)*math.sqrt(252),monthly=ma_c,bip_ratio=None)
months=sorted({(dd.year,dd.month) for u in data for dd,_,_ in data[u]['monthly']})
M=len(months); R0=6; RL=R0+M-1   # data rows in history sheets
NAVY='01313E'; TEAL='0287AA'; GREY='6A7A85'; LINE='DCE5EA'; CREAM='FFF8E7'
F=lambda b=False,c='000000',s=9: Font(name='Calibri',size=s,bold=b,color=c)
thin=Side(style='thin',color=LINE)
def hdr(ws,r,cols,c0=2,left=0):
    for i,h in enumerate(cols):
        c=ws.cell(r,c0+i,h); c.font=F(True,'FFFFFF'); c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='left' if i<left else 'center')
    ws.row_dimensions[r].height=30
def title(ws,t,sub,last_col):
    ws.sheet_view.showGridLines=False
    ws['B2']=t; ws['B2'].font=F(True,NAVY,10)
    ws['B3']=sub; ws['B3'].font=F(False,GREY); ws['B3'].alignment=Alignment(wrap_text=True,vertical='top'); ws.merge_cells(start_row=3,start_column=2,end_row=3,end_column=last_col); ws.row_dimensions[3].height=42
    ws.column_dimensions['A'].width=2
wb=openpyxl.Workbook()
def name(nm,ws,coord): wb.defined_names[nm]=DefinedName(nm,attr_text=f"'{ws.title}'!{coord}")
PCT='0.0 %'; PX='#,##0.00'
# ---------------- Paramètres
wp=wb.active; wp.title='Paramètres'
title(wp,'PARAMÈTRES - PROXIES COTÉS','Cellules crème : saisies. Les onglets Historique ajusté, Rendements, Composites et Synthèse sont en formules et se recalculent à partir de cet onglet.',8)
hdr(wp,5,['Paramètre','Valeur','Unité','Commentaire'],left=1)
def prow(r,lab,val,unit,com,nm,fmt=None):
    wp.cell(r,2,lab).font=F(); c=wp.cell(r,3,val); c.font=F(); c.fill=PatternFill('solid',fgColor=CREAM); c.alignment=Alignment(horizontal='right')
    if fmt: c.number_format=fmt
    wp.cell(r,4,unit).font=F(False,GREY); c=wp.cell(r,5,com); c.font=F(False,GREY); c.alignment=Alignment(wrap_text=True,vertical='top'); wp.merge_cells(start_row=r,start_column=5,end_row=r,end_column=8)
    for col in range(2,9): wp.cell(r,col).border=Border(bottom=thin)
    wp.row_dimensions[r].height=30; name(nm,wp,f'$C${r}')
prow(6,'Brookfield (NYSE:BIP) - date du regroupement 3 pour 2',dt.date(2016,9,14),'date','Les cours de fin de mois antérieurs à cette date sont divisés par le facteur ci-dessous.','BIP_date','dd/mm/yyyy')
prow(7,'Brookfield (NYSE:BIP) - facteur appliqué avant cette date',1.0,'x','1,5 pour une source brute (CapIQ). Mis à 1,0 ici : la série Yahoo est déjà ajustée du regroupement (contrôle en Synthèse : rapport septembre / août 2016).','BIP_facteur','0.00')
prow(8,'Brookfield (NYSE:BIP) - inclusion dans les composites',0,'1/0','Exclu : les distributions d\'actions BIPC de mars 2020 et juin 2022 ne sont probablement pas ajustées dans la série de cours (rendement total sous-estimé).','BIP_inclus','0')
prow(9,'INFR (LSE) - début des statistiques',dt.date(2016,1,1),'date','Quatre points aberrants de la source (pence et livres mélangés) avant cette date ; mois antérieurs ignorés dans les rendements.','INFR_debut','dd/mm/yyyy')
prow(10,'Seuil d\'exclusion d\'un rendement mensuel aberrant',0.6,'log','Un rendement mensuel dont la valeur absolue en logarithme dépasse ce seuil est ignoré (erreur de source). Mars 2020 reste en deçà pour toutes les séries.','Seuil_aberrant','0.00')
prow(11,'Fenêtre glissante des pires performances',12,'mois','','Fenetre','0')
r=13; wp.cell(r,2,'INCLUSION DANS LES COMPOSITES   1 = inclus, 0 = exclu').font=F(True,TEAL); r+=1
hdr(wp,r,['Ticker','Composite','C1 Infrastructure capital','C2 Dette d\'infrastructure','C3 Direct lending','C4 Poche liquide','Commentaire'],left=2); r+=1
FLAG_R0=r
for user,ysym,nm_,strat,typ,comp in TICK:
    wp.cell(r,2,user).font=F(True); wp.cell(r,3,comp).font=F()
    for j,cc in enumerate(['C1','C2','C3','C4']):
        v=1 if comp==cc else 0
        if user in EXCLUDE and comp==cc: v='=BIP_inclus'
        c=wp.cell(r,4+j,v); c.font=F(); c.fill=PatternFill('solid',fgColor=CREAM); c.alignment=Alignment(horizontal='center'); c.number_format='0'
    wp.cell(r,8,'Exclu par défaut (voir paramètre BIP)' if user in EXCLUDE else '').font=F(False,GREY)
    for col in range(2,9): wp.cell(r,col).border=Border(bottom=thin)
    r+=1
FLAG_RL=r-1
r+=1; wp.cell(r,2,'PONDÉRATION DU CONTRÔLE COTÉ DU PORTEFEUILLE   allocation cible du Fonds projetée sur les composites').font=F(True,TEAL); r+=1
hdr(wp,r,['Ligne du Fonds','Poids','C1','C2','C3','C4','Commentaire'],left=1); r+=1
W0=r
for lab,w,c1,c2,c3,c4,com in [('MAPIF II',0.34,1,0,0,0,'Secondaires d\'infrastructure : capital'),('MSIG 3',0.17,0,0,1,0,'Dette privée d\'infrastructure : direct lending'),('PG NGI',0.17,1,0,0,0,'Value add : capital (BIP serait le proxy le plus proche, exclu)'),('Ares AGI',0.17,0.6,0,0.2,0.2,'60 % infrastructure, 20 % dette privée, 20 % liquide et HY'),('Actifs liquides',0.15,0,0,0,1,'Monétaire et haut rendement')]:
    wp.cell(r,2,lab).font=F(True); c=wp.cell(r,3,w); c.number_format='0 %'; c.font=F(); c.fill=PatternFill('solid',fgColor=CREAM)
    for j,v in enumerate([c1,c2,c3,c4]):
        c=wp.cell(r,4+j,v); c.number_format='0 %'; c.font=F(); c.fill=PatternFill('solid',fgColor=CREAM); c.alignment=Alignment(horizontal='center')
    wp.cell(r,8,com).font=F(False,GREY)
    for col in range(2,9): wp.cell(r,col).border=Border(bottom=thin)
    r+=1
WL=r-1
wp.cell(r,2,'Poids des composites dans le contrôle coté').font=F(True)
for j in range(4):
    col=L(4+j); c=wp.cell(r,4+j,f'=SUMPRODUCT($C${W0}:$C${WL},{col}{W0}:{col}{WL})'); c.number_format='0.0 %'; c.font=F(True); c.alignment=Alignment(horizontal='center'); name(f'Wc_C{j+1}',wp,f'${col}${r}')
c=wp.cell(r,3,f'=SUM(C{W0}:C{WL})'); c.number_format='0 %'; c.font=F(True)
for col in range(2,9): wp.cell(r,col).border=Border(top=Side(style='medium',color=NAVY))
for i,w in enumerate([2,44,14,16,16,16,16,44]): wp.column_dimensions[L(i+1)].width=w
# ---------------- Brut (adjusted close as downloaded) and Cours
def hist_sheet(tname,sub,idx):
    w=wb.create_sheet(tname); title(w,tname.upper()+' - clôture de fin de mois, devise de cotation (LSE en livres)',sub,N+2)
    hdr(w,5,['Mois']+[t[0] for t in TICK],left=1); w.column_dimensions['B'].width=10
    lookup={u:{(dd.year,dd.month):row for dd,*row in data[u]['monthly']} for u in data}
    for i,(y,m) in enumerate(months):
        rr=R0+i; c=w.cell(rr,2,dt.date(y,m,1)); c.number_format='mm/yyyy'; c.font=F()
        for j,(user,*_ ) in enumerate(TICK):
            v=lookup[user].get((y,m))
            if v: c=w.cell(rr,3+j,v[idx-1]); c.number_format=PX; c.font=F()
            w.column_dimensions[L(3+j)].width=12
    return w
wc=hist_sheet('Historique cours','Clôture non ajustée, telle que publiée par la source. Mois en cours exclu. Information seulement : les calculs utilisent l\'onglet Brut ajusté puis Historique ajusté.',1)
wbru=hist_sheet('Brut ajusté','Clôture ajustée des dividendes telle que téléchargée (Yahoo Finance). Aucune correction : les corrections paramétrées sont appliquées dans l\'onglet Historique ajusté.',2)
# ---------------- Historique ajusté (formulas)
wa=wb.create_sheet('Historique ajusté'); title(wa,'HISTORIQUE AJUSTÉ - base des rendements','Formules : Brut ajusté corrigé des paramètres (facteur Brookfield avant la date de regroupement, début des statistiques INFR). Une cellule vide signifie : pas de donnée.',N+2)
hdr(wa,5,['Mois']+[t[0] for t in TICK],left=1); wa.column_dimensions['B'].width=10
for i in range(M):
    rr=R0+i; c=wa.cell(rr,2,f"='Brut ajusté'!B{rr}"); c.number_format='mm/yyyy'; c.font=F()
    for j,(user,ysym,*_ ) in enumerate(TICK):
        col=L(3+j); src=f"'Brut ajusté'!{col}{rr}"
        if user=='NYSE:BIP': f=f'=IF({src}="","",IF($B{rr}<BIP_date,{src}/BIP_facteur,{src}))'
        elif ysym=='INFR.L': f=f'=IF(OR({src}="",$B{rr}<INFR_debut),"",{src})'
        else: f=f'=IF({src}="","",{src})'
        c=wa.cell(rr,3+j,f); c.number_format=PX; c.font=F(); wa.column_dimensions[col].width=12
# ---------------- Rendements (monthly simple returns)
wr=wb.create_sheet('Rendements'); title(wr,'RENDEMENTS MENSUELS - devise de cotation, dividendes réinvestis','Formules : rendement simple d\'un mois sur l\'autre à partir de l\'Historique ajusté ; vide si donnée manquante ou rendement aberrant au sens du seuil paramétré.',N+2)
hdr(wr,5,['Mois']+[t[0] for t in TICK],left=1); wr.column_dimensions['B'].width=10
for i in range(1,M):
    rr=R0+i; c=wr.cell(rr,2,f"='Historique ajusté'!B{rr}"); c.number_format='mm/yyyy'; c.font=F()
    for j in range(N):
        col=L(3+j); a=f"'Historique ajusté'!{col}{rr}"; b=f"'Historique ajusté'!{col}{rr-1}"
        c=wr.cell(rr,3+j,f'=IF(OR({a}="",{b}=""),"",IF(ABS(LN({a}/{b}))>Seuil_aberrant,"",{a}/{b}-1))'); c.number_format=PCT; c.font=F(); wr.column_dimensions[col].width=12
RR0=R0+1; RRL=RL
# ---------------- Composites
wk=wb.create_sheet('Composites'); title(wk,'COMPOSITES COTÉS - rendements mensuels équipondérés et indices base 100','Formules : moyenne simple des rendements disponibles des titres inclus (onglet Paramètres), puis contrôle coté du portefeuille pondéré par l\'allocation cible. Rendements en devise locale : l\'effet de change n\'est pas inclus.',16)
cols=['Mois','C1 Infrastructure capital','C2 Dette d\'infra','C3 Direct lending','C4 Poche liquide','Portefeuille (contrôle coté)','Titres C1','Titres C2','Titres C3','Titres C4','Indice C1','Indice C2','Indice C3','Indice C4','Indice portefeuille']
hdr(wk,5,cols,left=1); wk.column_dimensions['B'].width=10
for j in range(2,16): wk.column_dimensions[L(j+1)].width=13
flag=lambda k: f"Paramètres!${L(4+k)}${FLAG_R0}:${L(4+k)}${FLAG_RL}"
rng=lambda rr: f"TRANSPOSE(Rendements!$C{rr}:${L(2+N)}{rr})"
for i in range(1,M):
    rr=R0+i; c=wk.cell(rr,2,f"=Rendements!B{rr}"); c.number_format='mm/yyyy'; c.font=F()
    for k in range(4):
        # number of available included returns
        cnt=f'SUMPRODUCT({flag(k)}*ISNUMBER({rng(rr)}))'
        num=f'SUMPRODUCT({flag(k)}*ISNUMBER({rng(rr)})*N(+{rng(rr)}))'
        c=wk.cell(rr,3+k,f'=IF({cnt}=0,"",{num}/{cnt})'); c.number_format=PCT; c.font=F()
        c=wk.cell(rr,8+k,f'={cnt}'); c.number_format='0'; c.font=F(False,GREY)
    c=wk.cell(rr,7,f'=IF(OR(C{rr}="",D{rr}="",E{rr}="",F{rr}=""),"",Wc_C1*C{rr}+Wc_C2*D{rr}+Wc_C3*E{rr}+Wc_C4*F{rr})'); c.number_format=PCT; c.font=F(True)
    for k in range(5):
        src=L(3+k); col=L(12+k)
        if i==1: f=f'=IF({src}{rr}="","",100)'
        else: f=f'=IF({src}{rr}="","",IF({col}{rr-1}="",100,{col}{rr-1}*(1+{src}{rr})))'
        c=wk.cell(rr,12+k,f); c.number_format='0.0'; c.font=F()
# ---------------- Synthèse
ws=wb.create_sheet('Synthèse',0)
title(ws,'COURS DES PROXIES COTÉS - OPENSTONE INFRAWORLD',f"Source : Yahoo Finance, extraction du {dt.date.today():%d/%m/%Y} (séance américaine en cours). Cours en devise de cotation, LSE converti de pence en livres. Colonnes de cours et performances à un an : valeurs extraites. Colonnes de statistiques mensuelles : formules sur l'onglet Rendements (devise locale, dividendes réinvestis, paramètres de l'onglet Paramètres).",24)
cols=['Ticker','Yahoo','Nom','Stratégie / usage proxy','Composite','Devise','Date','Dernier cours','Veille','Var. jour','Plus haut 1 an','Plus bas 1 an','Perf. 1 mois','Perf. 3 mois','Perf. 1 an (prix)','Perf. 1 an (dividendes réinvestis)','Vol. quotidienne 1 an','Vol. mensuelle annualisée','Rendement annualisé','Pire fenêtre glissante','Perte maximale','Mois disponibles','Inclus']
hdr(ws,5,cols,left=6)
r=5
for j,(user,*_ ) in enumerate(TICK):
    d=data[user]; r+=1; col=L(3+j)
    rr=f'Rendements!{col}${RR0}:{col}${RRL}'
    ha=f"'Historique ajusté'!{col}${R0}:{col}${RL}"
    vol=f'=IF(COUNT({rr})<24,"",STDEVP(IF(ISNUMBER({rr}),LN(1+{rr})))*SQRT(12))'
    cagr=f'=IF(COUNT({rr})<24,"",EXP(SUMPRODUCT(ISNUMBER({rr})*N(+IF(ISNUMBER({rr}),LN(1+{rr}),0)))*12/COUNT({rr}))-1)'
    vals=[user,d['ysym'],d['name'],d['strat'],d['comp'],d['cur'],d['last_date'],d['last'],d['prev'],d['last']/d['prev']-1,d['hi52'],d['lo52'],d['r1m'],d['r3m'],d['r1y'],d['tr1y'],d['vol_d'],vol,cagr,None,None,f'=COUNT({rr})',f'=SUM(Paramètres!D{FLAG_R0+j}:G{FLAG_R0+j})']
    fm=[None,None,None,None,None,None,'dd/mm/yyyy',PX,PX,PCT,PX,PX,PCT,PCT,PCT,PCT,PCT,PCT,PCT,PCT,PCT,'0','0']
    for i,v in enumerate(vals):
        c=ws.cell(r,2+i,v); c.font=F(i==0); c.border=Border(bottom=thin)
        if fm[i]: c.number_format=fm[i]
        c.alignment=Alignment(horizontal='left' if i<6 else 'right',vertical='center',wrap_text=i in (2,3))
    ws.row_dimensions[r].height=24
    data[user]['row']=r
# worst window / max drawdown : helper sheet
wh=wb.create_sheet('Aide calcul'); title(wh,'AIDE AU CALCUL - fenêtres glissantes et pertes maximales','Formules intermédiaires : performance sur la fenêtre paramétrée et drawdown par rapport au plus haut historique, pour chaque titre puis pour les composites.',2+2*(N+5))
labels=[t[0] for t in TICK]+['C1','C2','C3','C4','Portefeuille']
hdr(wh,5,['Mois']+[f'{l} fenêtre' for l in labels]+[f'{l} drawdown' for l in labels],left=1); wh.column_dimensions['B'].width=10
K=len(labels)
def series_ref(k,rr):
    if k<N: return f"'Historique ajusté'!{L(3+k)}{rr}"
    return f"Composites!{L(12+k-N)}{rr}"
for i in range(M):
    rr=R0+i; c=wh.cell(rr,2,f"='Historique ajusté'!B{rr}"); c.number_format='mm/yyyy'; c.font=F()
    for k in range(K):
        s=series_ref(k,rr)
        # window perf: value / value Fenetre months earlier
        if i>=1:
            f=f'=IF(OR({s}="",ROW()-Fenetre<{R0}),"",IF(INDEX({series_ref(k,"$1").split("!")[0]}!{L(3+k) if k<N else L(12+k-N)}:{L(3+k) if k<N else L(12+k-N)},ROW()-Fenetre)="","",{s}/INDEX({series_ref(k,"$1").split("!")[0]}!{L(3+k) if k<N else L(12+k-N)}:{L(3+k) if k<N else L(12+k-N)},ROW()-Fenetre)-1))'
        else: f='=""'
        c=wh.cell(rr,3+k,f); c.number_format=PCT; c.font=F(); wh.column_dimensions[L(3+k)].width=11
        # drawdown vs running max
        sheet=series_ref(k,"$1").split("!")[0]; colx=L(3+k) if k<N else L(12+k-N)
        f=f'=IF({s}="","",{s}/MAX({sheet}!{colx}${R0}:{colx}{rr})-1)'
        c=wh.cell(rr,3+K+k,f); c.number_format=PCT; c.font=F(); wh.column_dimensions[L(3+K+k)].width=11
for j,(user,*_ ) in enumerate(TICK):
    r=data[user]['row']; colw=L(3+j); cold=L(3+K+j)
    c=ws.cell(r,21,f"=IF(COUNT('Aide calcul'!{colw}${R0}:{colw}${RL})=0,\"\",MIN('Aide calcul'!{colw}${R0}:{colw}${RL}))"); c.number_format=PCT; c.font=F(); c.border=Border(bottom=thin)
    c=ws.cell(r,22,f"=IF(COUNT('Aide calcul'!{cold}${R0}:{cold}${RL})=0,\"\",MIN('Aide calcul'!{cold}${R0}:{cold}${RL}))"); c.number_format=PCT; c.font=F(); c.border=Border(bottom=thin)
# composites block in Synthèse
r=5+N+2
ws.cell(r,2,'COMPOSITES COTÉS   équipondérés, devise locale, dividendes réinvestis ; contrôle coté du portefeuille pondéré par l\'allocation cible').font=F(True,TEAL); r+=1
hdr(ws,r,['Composite','Titres inclus','Poids dans le contrôle coté','Vol. mensuelle annualisée','Rendement annualisé','Pire fenêtre glissante','Perte maximale','Mois','Perf. 12 derniers mois','Référence DIC'],left=2); r+=1
refs={'C1':'Preqin Infrastructure désmoothé : vol. 3,34 % ; value add 4,27 %','C2':'Pas d\'indice Preqin dédié ; dette d\'infra traitée en direct lending (6,09 %)','C3':'Preqin Private Debt Direct Lending désmoothé : vol. 6,09 %','C4':'Morningstar LSTA : vol. 5,68 % (coté)','P':'VEV de contrôle coté retenue au DIC : 10,88 %'}
for k,(code,lab,incl) in enumerate(COMP+[('P','Portefeuille (contrôle coté)','Pondération de l\'onglet Paramètres')]):
    col=L(3+k); rr=f'Composites!{col}${RR0}:{col}${RRL}'
    colw=L(3+N+k); cold=L(3+K+N+k)
    vals=[lab,incl,(f'=Wc_{code}' if code!='P' else '=SUM(Wc_C1:Wc_C4)' if False else (f'=Wc_{code}' if code!='P' else '=Wc_C1+Wc_C2+Wc_C3+Wc_C4')),
          f'=IF(COUNT({rr})<24,"",STDEVP(IF(ISNUMBER({rr}),LN(1+{rr})))*SQRT(12))',
          f'=IF(COUNT({rr})<24,"",EXP(SUMPRODUCT(ISNUMBER({rr})*N(+IF(ISNUMBER({rr}),LN(1+{rr}),0)))*12/COUNT({rr}))-1)',
          f"=IF(COUNT('Aide calcul'!{colw}${R0}:{colw}${RL})=0,\"\",MIN('Aide calcul'!{colw}${R0}:{colw}${RL}))",
          f"=IF(COUNT('Aide calcul'!{cold}${R0}:{cold}${RL})=0,\"\",MIN('Aide calcul'!{cold}${R0}:{cold}${RL}))",
          f'=COUNT({rr})',
          f'=IF(COUNT(Composites!{col}{RRL-11}:{col}{RRL})<12,"",PRODUCT(1+Composites!{col}{RRL-11}:{col}{RRL})-1)',
          refs[code]]
    fm=[None,None,'0.0 %',PCT,PCT,PCT,PCT,'0',PCT,None]
    for i,v in enumerate(vals):
        c=ws.cell(r,2+i,v); c.font=F(i==0); c.border=Border(bottom=thin)
        if fm[i]: c.number_format=fm[i]
        c.alignment=Alignment(horizontal='left' if i in (0,1,9) else 'right',vertical='center',wrap_text=i in (1,9))
    ws.merge_cells(start_row=r,start_column=11,end_row=r,end_column=24); ws.row_dimensions[r].height=24; r+=1
r+=1
ws.cell(r,2,'CONTRÔLE BROOKFIELD   rapport des cours de fin de mois septembre / août 2016 dans la série brute').font=F(True,TEAL); r+=1
jb=[t[0] for t in TICK].index('NYSE:BIP'); colb=L(3+jb)
ws.cell(r,2,'Rapport septembre 2016 / août 2016, série Brut ajusté').font=F()
c=ws.cell(r,4,f"=INDEX('Brut ajusté'!{colb}${R0}:{colb}${RL},MATCH(DATE(2016,9,1),'Brut ajusté'!$B${R0}:$B${RL},0))/INDEX('Brut ajusté'!{colb}${R0}:{colb}${RL},MATCH(DATE(2016,8,1),'Brut ajusté'!$B${R0}:$B${RL},0))"); c.number_format='0.000'; c.font=F(True)
c=ws.cell(r,5,'="Facteur appliqué : "&SUBSTITUTE(ROUND(BIP_facteur,2)&"",".",",")&" - un rapport proche de 1 signifie que la source a déjà ajusté le regroupement 3 pour 2 ; un rapport proche de 1,5 appelle le facteur 1,5."'); c.font=F(False,GREY); ws.merge_cells(start_row=r,start_column=5,end_row=r,end_column=24)
r+=2
for n in ["Lecture risque : les trusts d'infrastructure cotés à Londres (3IN, HICL, INPP, SEQI, GCP) traitent avec une décote sur l'actif net ; leur cours intègre un risque de taux et de liquidité que les VL des fonds cibles ne montrent pas. Ils constituent le contrôle coté le plus proche des fonds non cotés du portefeuille.",
"Brookfield (BIP) est conservé en information mais exclu des composites : les distributions d'actions BIPC de mars 2020 et de juin 2022 ne sont pas reflétées dans la série de cours, dont le rendement total est donc sous-estimé. Le facteur de regroupement de septembre 2016 est paramétré (onglet Paramètres).",
"Les composites sont équipondérés en devise locale : ils mesurent le risque de marché des actifs, non le risque de change du Fonds (51 % de l'actif en dollars non couverts). Les volatilités mensuelles sont comparables aux volatilités désmoothées Preqin du DIC et à la VEV de contrôle coté de 10,88 %."]:
    c=ws.cell(r,2,n); c.font=F(False,GREY); c.alignment=Alignment(wrap_text=True,vertical='top'); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=24); ws.row_dimensions[r].height=30; r+=1
for i,w in enumerate([2,16,10,34,34,10,8,11,11,11,10,11,11,10,10,11,14,12,13,13,12,10,10,7]): ws.column_dimensions[L(i+1)].width=w
# array formulas: STDEVP(IF(...)) needs array entry -> use openpyxl ArrayFormula
from openpyxl.worksheet.formula import ArrayFormula
for w in (ws,):
    for row in w.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('=') and 'STDEVP(IF(' in c.value or (isinstance(c.value,str) and 'N(+IF(' in c.value) or (isinstance(c.value,str) and c.value.startswith('=IF(COUNT(Composites') and 'PRODUCT(1+' in c.value):
                c.value=ArrayFormula(c.coordinate,c.value)
for w in wb.worksheets: assert w.freeze_panes is None
wb.save('Cours_Proxies_Cotes_Openstone_Infraworld.xlsx'); print('saved',M,'months')
