# -*- coding: utf-8 -*-
from lib import *
from openpyxl.workbook.defined_name import DefinedName
import datetime as dt

def name(wb, ws, nm, coord):
    wb.defined_names[nm]=DefinedName(nm, attr_text=f"'{ws.title}'!${coord[0]}${coord[1:]}")

def build_params(wb):
    ws=wb.create_sheet('0 Paramètres')
    S=Sheet(ws,'PARAMÈTRES - OPENSTONE INFRAWORLD, COMPARTIMENT I · rapport de risque initial, octobre 2026',
            'Toutes les valeurs du rapport découlent de ce bloc ; aucune autre feuille ne redéfinit un paramètre. Saisies sur fond crème, valeurs dérivées en formules.',
            {'A':2,'B':44,'C':14,'D':10,'E':52,'F':11,'G':60})
    S.hdr(['Paramètre','Valeur','Unité','Source','Fiabilité','Réserve']); hdrcols=6
    S.ncols=6
    P={}
    def row(key,label,val,unit,src,fiab,res='',fmt=None,derived=False):
        r=S.row([label,val,unit,src,fiab,res],fmts=[None,fmt],input_cols=() if derived else (1,))
        P[key]=f"'0 Paramètres'!$C${r}"; name(wb,ws,key,f'C{r}'); return r
    S.block('FONDS ET DATES')
    row('DateArrete','Date d\'arrêté du rapport',dt.date(2026,10,6),'date','Date de l\'analyse ; prospectus projet du 30 septembre 2026','Élevée','Le Fonds n\'est pas constitué : rapport initial sur documentation, sans valeur liquidative ni passif.','dd/mm/yyyy')
    row('DateProspectus','Date du projet de prospectus',dt.date(2026,9,30),'date','Prospectus Openstone Infraworld, projet du 30 septembre 2026','Élevée','Plusieurs mentions restent à compléter (SRI, ISIN A3c, mentions de liquidité).','dd/mm/yyyy')
    row('AN_cible','Actif net cible',50000000,'euros','Encours cible retenu par la note de liquidité et l\'IRP','Moyenne','Prévision commerciale ; les simulations sont rejouées sur l\'encours réel une fois le Fonds lancé.','# ##0')
    row('AN_min','Actif net minimal réglementaire',300000,'euros','Règlement art. 2 ; dissolution si l\'actif reste en deçà 30 jours','Élevée','','# ##0')
    row('Seuil_Declenchement','Seuil de la Date de Déclenchement',10000000,'euros','Prospectus, définitions du Compartiment I : 10 M€ d\'engagements cumulés ou 2e anniversaire','Élevée','Bascule des parts B (conseil 0,57 %) vers le barème des parts A (1,07 %).','# ##0')
    row('Nominal','Valeur nominale de la part',100,'euros','Prospectus art. 26.3 - souscription au nominal tant que la VL lui est inférieure','Élevée','Souscription au-dessus de la VL si celle-ci passe sous 100 € : point de traitement équitable (registre, n° 9).','# ##0.00')
    S.block('ALLOCATION CIBLE  ·  synthèse Openstone de septembre 2026, non contractuelle')
    row('W_MAPIF','Macquarie Alliance Partners Infrastructure Fund II (MAPIF II)',0.34,'% AN','Synthèse Openstone septembre 2026, p. 2','Moyenne','Fonds fermé de secondaires, USD ; conditionné à l\'acceptation de la souscription et à la dérogation de ticket (10 M$).','0 %')
    row('W_MSIG','Macquarie Specialized Infrastructure Global 3 (MSIG 3)',0.17,'% AN','Synthèse Openstone septembre 2026, p. 2','Moyenne','Fonds de dette fermé hybride ; véhicule (USD / EUR) et manche (levered / unlevered) non arrêtés.','0 %')
    row('W_PG','Partners Group Next Generation Infrastructure (PG NGI)',0.17,'% AN','Synthèse Openstone septembre 2026, p. 2','Moyenne','Evergreen, devise de référence USD ; classe I ou classe H couverte à arrêter.','0 %')
    row('W_ARES','Ares Global Infrastructure ELTIF (Ares AGI)',0.17,'% AN','Synthèse Openstone septembre 2026, p. 2','Moyenne','Evergreen ELTIF en EUR ; souscription au nominal jusqu\'au 01/11/2026 attendu.','0 %')
    row('W_LIQ','Actifs liquides',0.15,'% AN','Synthèse septembre 2026 ; prospectus art. 21.8 (minimum 15 % de l\'actif)','Élevée','Plancher écrit comme obligation permanente, sans clause de dépassement passif (registre, n° 2).','0 %')
    r=S.row(['Total de l\'allocation','=W_MAPIF+W_MSIG+W_PG+W_ARES+W_LIQ','% AN','dérivé','-','Doit boucler à 100 %.'],fmts=[None,'0 %']); P['W_TOT']=f"'0 Paramètres'!$C${r}"; name(wb,ws,'W_TOT',f'C{r}')
    r=S.row(['Part des fonds fermés (MAPIF II + MSIG 3)','=W_MAPIF+W_MSIG','% AN','dérivé','-','Aucune liquidité avant échéance ; sortie par cession secondaire décotée.'],fmts=[None,'0 %']); P['W_FERMES']=f"'0 Paramètres'!$C${r}"; name(wb,ws,'W_FERMES',f'C{r}')
    r=S.row(['Part des fonds evergreen (PG NGI + Ares AGI)','=W_PG+W_ARES','% AN','dérivé','-','Rachat trimestriel plafonné à 5 % du fonds cible.'],fmts=[None,'0 %']); P['W_EVER']=f"'0 Paramètres'!$C${r}"; name(wb,ws,'W_EVER',f'C{r}')
    r=S.row(['Part gérée par Macquarie','=W_MAPIF+W_MSIG','% AN','dérivé','-','Concentration gérant : 51 % de l\'allocation cible chez un seul groupe.'],fmts=[None,'0 %']); P['W_MACQ']=f"'0 Paramètres'!$C${r}"; name(wb,ws,'W_MACQ',f'C{r}')
    r=S.row(['Part libellée en dollars - hypothèse basse (MAPIF II)','=W_MAPIF','% AN','dérivé','-','Si MSIG 3 est souscrit via le fonds EUR et PG NGI en classe couverte.'],fmts=[None,'0 %']); name(wb,ws,'W_USD_BAS',f'C{r}')
    r=S.row(['Part libellée en dollars - hypothèse haute (MAPIF II + MSIG 3 USD + PG NGI)','=W_MAPIF+W_MSIG+W_PG','% AN','dérivé','-','Devise de référence de PG NGI : USD (supplément 2) ; MSIG 3 (USD) sans classe EUR.'],fmts=[None,'0 %']); name(wb,ws,'W_USD_HAUT',f'C{r}')
    S.block('CONDITIONS DE LIQUIDITÉ DU FONDS  ·  prospectus du 30 septembre 2026')
    row('Blocage_mois','Période de Blocage par souscription',24,'mois','Prospectus art. 7.2.2 - 2 ans à compter de chaque souscription, imputation FIFO','Élevée','')
    row('Preavis_j','Préavis de rachat',91,'jours','Prospectus art. 7.2.3 - centralisation le 31 mars pour la VL du 30 juin, le 30 septembre pour celle du 31 décembre','Élevée','La note de liquidité du 24/09 et le classeur de stress tests retiennent 180 jours : à recaler.')
    row('Preavis_max_j','Préavis prolongé au plus',180,'jours','Prospectus art. 7.2.7 - 180 jours calendaires au plus, 12 mois au plus','Élevée','')
    row('Gate','Plafonnement des Rachats par date de centralisation',0.05,'% AN','Prospectus art. 7.2.6 - 5 % de l\'Actif Net, demandes nettes des souscriptions','Élevée','Assiette : actif net de la VL d\'exécution, inconnue à l\'activation (registre, n° 7).','0,0 %')
    row('Gate_dates','Reports successifs au plus',4,'dates','Prospectus art. 7.2.6 - 4 Dates de Centralisation ; Demandes Prioritaires après 2 reports','Élevée','L\'issue au-delà de 4 dates n\'est pas écrite ; la note de liquidité prévoit l\'examen de la suspension.')
    row('Reglement_JO','Délai de règlement des rachats',35,'j.o.','Prospectus art. 7.2.5 - 5 Jours Ouvrés après publication, publication sous 30 Jours Ouvrés','Élevée','')
    row('Centralisation_sous_JO','Centralisation des souscriptions avant la VL',25,'j.o.','Prospectus art. 7.1.2','Élevée','Paiement intégral à la centralisation ; sommes bloquées hors actif net.')
    row('Publication_JO','Publication de la valeur liquidative',30,'j.o.','Prospectus art. 6.2','Élevée','VL mensuelle au dernier Jour Ouvré du mois.')
    row('Liq_min','Plancher d\'actifs liquides',0.15,'% actif','Prospectus art. 21.8','Élevée','Assiette en actif, les autres limites en actif net (registre, n° 8).','0 %')
    row('Liq_cible','Poche liquide visée par la gestion',0.16,'% AN','Note de liquidité du 24/09/2026, § 1','Moyenne','','0 %')
    row('Liq_monetaire_min','Dont monétaire et dépôts au minimum (règle interne)',0.05,'% actif','Note de liquidité, § 2','Moyenne','Règle interne, non opposable.','0 %')
    row('Surengagement_max','Plafond de Sur-Engagement',1.30,'% AN','Prospectus art. 23 ; art. 3.1.1 (engagements hors bilan, 130 %)','Élevée','Deux plafonds à 130 % sans précision de cumul (registre, n° 6).','0 %')
    row('Emprunt_max','Emprunts d\'espèces',0,'% AN','Prospectus art. 22 - néant ; ni ligne de crédit ni instrument dérivé','Élevée','Les art. 3.1.1, 14.2 et 25.2.6 supposent encore un emprunt (registre, n° 1).','0 %')
    row('Levier_max','Levier AIFM, méthode de l\'engagement',1.00,'% AN','Note de liquidité § 4 ; à réintroduire au prospectus (art. 23 § 1 a) de la directive 2011/61/UE)','Moyenne','Le projet du 30/09 ne mentionne plus de niveau maximal de levier.','0 %')
    row('Entreprises_max','Entreprises Cibles (co-investissements directs)',0.30,'% AN','Prospectus art. 22','Élevée','Traitement d\'un dépassement passif non précisé.','0 %')
    row('Pers_phys_max','Détention maximale d\'une personne physique',0.10,'% parts','Prospectus art. 5.1 et 8.1','Élevée','Contrôle à la souscription et au transfert.','0 %')
    S.block('FRAIS  ·  prospectus art. 27')
    row('Fee_SGP','Commission de Gestion (tranche 0 - 100 M€)',0.0018,'% AN','Prospectus art. 27.2.1 - 0,18 % ; 0,15 % de 100 à 500 M€ ; 0,10 % au-delà ; minimum 60 000 € par an','Élevée','Le minimum de 60 000 € pèse 0,12 % à 50 M€ et 1,2 % à 5 M€.','0,00 %')
    row('Fee_CIF_A','Commission de Conseil, parts A (tranche 0 - 100 M€)',0.0107,'% AN','Prospectus art. 27.2.2 - 1,07 % ; 1,10 % ; 1,15 %','Élevée','Barème qualifié de dégressif alors que les taux croissent (registre, n° 11).','0,00 %')
    row('Fee_CIF_B','Commission de Conseil, parts B jusqu\'à la Date de Déclenchement',0.0057,'% AN','Prospectus art. 27.2.2','Élevée','','0,00 %')
    row('Fee_Distrib','Commission de Distribution, parts A1 et B1',0.0075,'% AN','Prospectus art. 27.2.3','Élevée','Rétrocédée au Conseiller pour les Distributeurs.','0,00 %')
    row('Fee_Entree','Droits d\'entrée au plus',0.05,'% prix','Prospectus art. 27.2.4','Élevée','Non acquis au Fonds.','0 %')
    row('Fee_Const','Frais de constitution au plus',0.0125,'% AN','Prospectus art. 27.1 - le plus élevé de 60 000 € ou 1,25 % de l\'actif net au 1er anniversaire, amortis sur 5 exercices','Élevée','','0,00 %')
    row('Fee_Total_A1','Frais récurrents directs, part A1 (hors frais fixes)','=Fee_SGP+Fee_CIF_A+Fee_Distrib','% AN','dérivé','-','Hors dépositaire, comptable, CAC et frais des fonds cibles (1,00 à 1,25 % plus carried 12,5 à 15 %).','0,00 %',derived=True)
    S.block('SEUILS D\'ALERTE DE LA NOTE DE LIQUIDITÉ  ·  § 5, vert / orange / rouge')
    row('S_file_rouge','File d\'attente (% AN) - rouge au-delà de',0.25,'% AN','Note de liquidité § 5','Moyenne','Orange entre 0 et 25 %.','0 %')
    row('S_resorption_rouge','Délai de résorption de la file - rouge au-delà de',25,'mois','Note de liquidité § 5','Moyenne','Orange de 12 à 25 mois.')
    row('S_poche_rouge','Poche d\'actifs liquides - rouge en deçà de',0.05,'% actif','Note de liquidité § 5','Moyenne','Orange de 5 à 15 %.','0 %')
    row('S_couv_vert','Couverture de liquidité à 12 mois - vert au-delà de',1.5,'ratio','Note de liquidité § 5','Moyenne','Orange de 1,0 à 1,5 ; rouge en deçà de 1,0.','0.0')
    row('S_couv_rouge','Couverture de liquidité à 12 mois - rouge en deçà de',1.0,'ratio','Note de liquidité § 5','Moyenne','','0.0')
    row('S_sureng_orange','Sur-engagement - orange au-delà de',1.20,'% AN','Note de liquidité § 5','Moyenne','Rouge au-delà de 130 %.','0 %')
    row('S_top1_orange','Premier porteur - orange au-delà de',0.10,'% AN','Note de liquidité § 5','Moyenne','Rouge au-delà de 15 %.','0 %')
    row('S_top5_orange','Cinq premiers porteurs - orange au-delà de',0.25,'% AN','Note de liquidité § 5','Moyenne','Rouge au-delà de 40 %.','0 %')
    S.block('INITIAL RISK PROFILE  ·  classeur du 06/10/2026, scores 0 à 100')
    row('IRP_Fonds','Score IRP du Fonds, caractéristiques propres',71.4,'/100','Initial Risk Profile Openstone Infraworld, onglet Openstone_Infraworld','Élevée','ÉLEVÉ ; liquidité, crédit saturés ; contrepartie et levier faibles.','0.0')
    row('IRP_Transp','Score IRP en transparence, pondéré par l\'allocation cible',62.1,'/100','Idem, bloc Transparence','Élevée','73,1 hors poche liquide.','0.0')
    row('IRP_MAPIF','Score IRP MAPIF II',52.1,'/100','Idem, onglet MAPIF_II','Élevée','MOYEN - fonds fermé sans risque de liquidité propre ; levier et opérationnel élevés.','0.0')
    row('IRP_MSIG','Score IRP MSIG 3 (USD, manche levered)',82.3,'/100','Idem, onglet MSIG_3','Élevée','ÉLEVÉ - crédit et levier saturés ; fonds non lancé à la date du PPM.','0.0')
    row('IRP_PG','Score IRP PG Next Generation Infrastructure',92.6,'/100','Idem, onglet PG_NGI','Élevée','ÉLEVÉ - liquidité, opérationnel, levier saturés ; seul fonds cible en activité.','0.0')
    row('IRP_ARES','Score IRP Ares Global Infrastructure ELTIF',86.4,'/100','Idem, onglet Ares_AGI','Élevée','ÉLEVÉ - blind pool en montée en charge, levier AIFM 400 % / 500 %.','0.0')
    row('SRI','Indicateur synthétique de risque (DIC PRIIPs)',4,'/7','DIC du 29/09/2026, parts A1 à B3 - classe 4, relèvement prudentiel d\'une classe','Élevée','Le prospectus du 30/09 et la synthèse laissent le SRI à compléter (X/7).')
    row('PDR','Période de détention recommandée',10,'ans','Prospectus art. 4.2 (supérieure à 10 ans) ; DIC du 29/09 (10 ans)','Élevée','La synthèse Openstone affiche 8 ans : à aligner (registre, n° 10).')
    S.blank()
    S.note('Fiabilité : Élevée = texte contractuel ou classeur validé ; Moyenne = prévision, hypothèse interne ou document commercial. Les montants en euros sont des cibles : le Fonds n\'a ni valeur liquidative ni porteur à la date d\'arrêté.',6)
    return P

def build_params_stress(wb, ws_name='0 Paramètres'):
    """Bloc d'hypothèses du moteur de stress tests, ajouté en bas de l'onglet Paramètres."""
    ws=wb[ws_name]
    r=ws.max_row+2
    from lib import F, fill, al, B_THIN, WHITE, NAVY, TEAL, INK, GREY
    import datetime as dt
    def block(text):
        nonlocal r
        c=ws.cell(r,2,text); c.font=F(True,WHITE); c.fill=fill(NAVY)
        for col in range(3,8): ws.cell(r,col).fill=fill(NAVY)
        r+=1
    def row(key,label,val,unit,src,fiab,res='',fmt=None,derived=False):
        nonlocal r
        for i,v in enumerate([label,val,unit,src,fiab,res]):
            c=ws.cell(r,2+i,v); c.font=F(False,INK); c.alignment=al('left',v='top')
            if i==1 and fmt: c.number_format=fmt
            if i==1 and not derived: c.fill=fill('FFF8F2')
        for col in range(2,8): ws.cell(r,col).border=B_THIN
        name(wb,ws,key,f'C{r}'); r+=1
    block('HYPOTHÈSES DU MOTEUR DE STRESS TESTS  ·  onglet 3b, reprises de la note de liquidité § 6 et du classeur du 24/09/2026')
    row('Date_depart','Date de départ de la projection',dt.date(2028,9,30),'date','Fin du blocage des premières souscriptions (hypothèse de constitution fin 2026)','Moyenne','Toutes les parts sont supposées rachetables dès le départ : prudent.','dd/mm/yyyy')
    row('NA_depart','Non-appelé de départ (régime normal)',0.15,'% AN','Note de liquidité § 6 et § 7 : non-appelé ramené à 15 % de l\'actif net à la fin du blocage','Moyenne','Soit un sur-engagement de 115 % au départ.','0 %')
    row('Part_MAPIF_NA','Part de MAPIF II dans le non-appelé',0.60,'%','Hypothèse (classeur du 24/09)','Moyenne','','0 %')
    row('Appels_base','Appels en régime normal, part du non-appelé initial par trimestre','=1/12','% / trim.','Tirage sur trois ans','Moyenne','',"0.0 %",derived=True)
    row('Capacite_evg','Capacité de rachat des fonds evergreen, part de la ligne par trimestre',0.05,'% / trim.','Supplément 2 PG NGI ; annexe 7 Ares','Élevée','Gate de 5 % par trimestre chez chaque fonds cible.','0 %')
    row('Decote','Décote de cession secondaire',0.10,'%','Hypothèse','Faible','Optimiste pour MAPIF II (accord du GP), sans objet pour MSIG 3 (run-off). Sensibilité à 25 % dans la note.','0 %')
    row('Poche_min_exec','Poche minimale après exécution au-delà du plafond',0.10,'% AN','Note de liquidité § 5','Moyenne','','0 %')
    row('Seuil_reinv','Seuil de réinvestissement de la trésorerie dans les fonds evergreen',0.20,'% AN','Hypothèse','Moyenne','','0 %')
    row('Taux_treso','Rendement de la poche liquide (monétaire et haut rendement pondérés)','=(Liq_mon_w*TauxMon+Liq_hy_w*Rend_HY)/W_LIQ','par an','dérivé des entrées du DIC (bloc PRIIPs)','Moyenne','5 % en monétaire à 2,5 %, 10 % en haut rendement à 6 %.','0.00 %',derived=True)
    row('Taux_emprunt','Taux d\'emprunt',0.0,'par an','Sans objet : aucun emprunt','Élevée','','0.0 %')
    row('Frais_an','Frais récurrents du Fonds, prestataires compris (part A1)','=Fee_Fonds_A1','% AN / an','Classeur KID PRIIPs v2.3, Fees H20 ; frais des fonds cibles exclus (déjà dans les rendements nets)','Élevée','Hypothèse prudente : toutes les parts au barème A1.','0.00 %',derived=True)
    row('Rachats_norm','Rachats demandés en régime normal, par date',0.02,'% AN','Hypothèse','Faible','Aucun historique : taux postulé.','0 %')
    row('Collecte_norm','Collecte en régime normal, par trimestre',0.05,'% AN','Prévision du Conseiller, à confirmer','Faible','Nulle dans tous les scénarios de crise.','0 %')
    row('Dist_porteurs','Acompte cible des parts de distribution',0.05,'% AN / an','Synthèse : 4 à 6 % ; DIC : 5 %','Moyenne','','0 %')
    row('Part_d','Part des parts de distribution dans l\'actif net',0.20,'% AN','Hypothèse','Faible','','0 %')
    row('Distrib_tension','Acomptes suspendus en cas de file, de poche sous le plancher ou de suspension (1 = oui)',1,'1/0','Note de liquidité § 7 ; prospectus art. 14 (examen discrétionnaire)','Élevée','')
    row('Deduc_ares','Déduction de sortie Ares (parts détenues moins de 24 mois)',0.05,'%','Annexe 7 : Early Redemption Deduction','Élevée','','0 %')
    row('Deduc_trim','Trimestres de projection concernés par la déduction Ares',4,'trim.','Hypothèse : à caler sur les dates de souscription Ares','Moyenne','')
    row('MSIG_USD','MSIG 3 détenu en dollars (1 = oui)',0,'1/0','Hypothèse : fonds EUR retenu ; le PPM lu est celui du fonds USD','Faible','Si 1, l\'exposition dollar monte à 51 % hors PG NGI.')
    row('Top5','Part des cinq premiers porteurs',0.35,'% AN','Seuil orange des indicateurs','Faible','Aucun registre : hypothèse de structure du passif.','0 %')
    row('Choc_cap','Choc de valorisation retenu, capital infrastructure',0.20,'%','Note de liquidité § 6 ; indices Preqin corrigés du lissage','Moyenne','Supérieur au choc à 99 % à un an des indices désamorcés (7,8 à 9,9 %).','0 %')
    row('Choc_dette','Choc de valorisation retenu, dette infrastructure',0.10,'%','Note de liquidité § 6','Moyenne','','0 %')
    row('Choc_usd','Variation du dollar dans les scénarios de change',0.15,'%','Note de liquidité § 6','Moyenne','Signe selon le scénario.','0 %')
    block('LIGNES DU PORTEFEUILLE  ·  rendement net annuel, distribution annuelle, part en capital (classeur DIC et documents des fonds cibles)')
    row('Rend_MAPIF','MAPIF II - rendement net annuel',0.135,'par an','Classeur KID PRIIPs v2.3, Fonds cibles D17 : TRI net cible retenu (gérant 12 à 15 %, relevé de 0,5 point), avant frais du Fonds','Faible','Net des frais du fonds cible, avant frais du Fonds et change.','0.00 %')
    row('Dist_MAPIF','MAPIF II - distribution annuelle',0.06,'par an','Hypothèse (secondaires matures)','Faible','','0.00 %')
    row('Cap_MAPIF','MAPIF II - part en capital',1,'%','Equity','Élevée','','0 %')
    row('Rend_MSIG','MSIG 3 - rendement net annuel',0.11,'par an','Classeur KID PRIIPs v2.3, Fonds cibles F17 : manche levered 10 à 12 %','Faible','Manche unlevered : 8 à 10 %.','0.00 %')
    row('Dist_MSIG','MSIG 3 - distribution annuelle',0.07,'par an','PPM : revenu distribué trimestriellement','Moyenne','','0.00 %')
    row('Cap_MSIG','MSIG 3 - part en capital',0,'%','Dette','Élevée','','0 %')
    row('Rend_PG','PG NGI - rendement net annuel',0.115,'par an','Classeur KID PRIIPs v2.3, Fonds cibles E17 ; aucune fourchette publiée par le gérant','Faible','DIC PG NGI : 9,2 % en scénario intermédiaire.','0.00 %')
    row('Dist_PG','PG NGI - distribution annuelle',0.02,'par an','Hypothèse','Faible','Imputée sur le gate du fonds cible.','0.00 %')
    row('Cap_PG','PG NGI - part en capital',1,'%','Equity','Élevée','','0 %')
    row('Rend_ARES','Ares AGI - rendement net annuel',0.095,'par an','Classeur KID PRIIPs v2.3, Fonds cibles G17 : documentation commerciale 8 à 10 %, non contractuelle','Faible','','0.00 %')
    row('Dist_ARES','Ares AGI - distribution annuelle',0.05,'par an','Fiche : distributions mensuelles attendues','Moyenne','','0.00 %')
    row('Cap_ARES','Ares AGI - part en capital',0.7,'%','Mix equity et dette','Moyenne','','0 %')
    row('Ws_MAPIF','MAPIF II - poids de départ (cible moins non-appelé)','=W_MAPIF-NA_depart*Part_MAPIF_NA','% AN','dérivé','-','La part non encore appelée est portée par les fonds evergreen en attendant.','0.0 %',derived=True)
    row('Ws_MSIG','MSIG 3 - poids de départ','=W_MSIG-NA_depart*(1-Part_MAPIF_NA)','% AN','dérivé','-','','0.0 %',derived=True)
    row('Ws_PG','PG NGI - poids de départ','=W_PG+NA_depart/2','% AN','dérivé','-','','0.0 %',derived=True)
    row('Ws_ARES','Ares AGI - poids de départ','=W_ARES+NA_depart/2','% AN','dérivé','-','','0.0 %',derived=True)
    row('Ws_LIQ','Actifs liquides - poids de départ','=W_LIQ','% AN','dérivé','-','','0.0 %',derived=True)
    row('Ws_TOT','Total des poids de départ','=Ws_MAPIF+Ws_MSIG+Ws_PG+Ws_ARES+Ws_LIQ','% AN','dérivé','-','Doit boucler à 100 %.','0 %',derived=True)
