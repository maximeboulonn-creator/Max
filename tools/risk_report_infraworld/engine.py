# -*- coding: utf-8 -*-
"""Moteur de stress tests de liquidité en formules : projection trimestrielle sur 36 mois, un bloc par scénario."""
from lib import *
from params import name
from openpyxl.utils import get_column_letter as L

NQ=12  # trimestres T1..T12 ; T0 en colonne C
COL0=3 # colonne C = T0
def c(t): return L(COL0+t)   # lettre de colonne du trimestre t (0..12)

# Scénarios : (nom, famille, r1, r2, rsuiv, collecte, coefDist, trimSansDist, NA, part4, usd, chocCap, chocDette, capEvg14, capPG5, capAres5, credit, gate, beyond, lecture)
S=[
 ('Base : rachats 2 % par date, collecte 5 % par trimestre','Passif','=Rachats_norm','=Rachats_norm','=Rachats_norm','=Collecte_norm',1,0,'=NA_depart','=1/3',0,0,0,1,1,1,0,1,1,'Régime de croisière.'),
 ('Adverse : rachats de 10 % sur deux dates','Passif',0.10,0.10,0,0,1,0,'=NA_depart','=1/3',0,0,0,1,1,1,0,1,1,'Deux dates de demandes doubles du plafond, collecte nulle.'),
 ('Extrême : rachats de 15 % puis 20 %','Passif',0.15,0.20,0,0,1,0,'=NA_depart','=1/3',0,0,0,1,1,1,0,1,1,'Sept fois le plafond sur deux dates.'),
 ('Concentration : sortie des cinq premiers porteurs','Passif','=Top5+Rachats_norm','=Rachats_norm','=Rachats_norm',0,1,0,'=NA_depart','=1/3',0,0,0,1,1,1,0,1,1,'Les cinq premiers porteurs sortent à la première date ; les autres demandent 2 % par date. Le classeur du 24/09 laissait cette demande à zéro.'),
 ('Distributions des fonds cibles divisées par deux et retardées d\'un an','Actif','=Rachats_norm','=Rachats_norm','=Rachats_norm',0,0.5,4,'=NA_depart','=1/3',0,0,0,1,1,1,0,1,1,'Aucune distribution la première année, puis moitié.'),
 ('Appels de fonds accélérés, dollar + 15 %','Actif','=Rachats_norm','=Rachats_norm','=Rachats_norm',0,1,0,'=NA_depart',0.5,'=Choc_usd',0,0,1,1,1,0,1,1,'Moitié du non-appelé tirée la première année, renchérie par le dollar.'),
 ('Valorisation - 20 % (capital), - 10 % (dette), dollar - 15 %','Actif','=Rachats_norm','=Rachats_norm','=Rachats_norm',0,1,0,'=NA_depart','=1/3','=-Choc_usd','=Choc_cap','=Choc_dette',1,1,1,0,1,1,'Choc de valeur au premier trimestre.'),
 ('Crise combinée : chocs de passif et d\'actif simultanés','Combiné',0.15,0.20,0,0,0.5,4,'=NA_depart',0.5,0,0,0,0.5,0.5,0,0,1,1,'Rachats extrêmes, distributions divisées par deux et retardées, appels accélérés, fonds evergreen au prorata puis Ares fermé.'),
 ('Fonds evergreen au prorata','Actif','=Rachats_norm','=Rachats_norm','=Rachats_norm',0,1,0,'=NA_depart','=1/3',0,0,0,0.5,0.5,0,0,1,1,'PG NGI et Ares servent la moitié des demandes la première année ; Ares ne sert plus rien ensuite.'),
]
for x in [0.05,0.075,0.10,0.125,0.15]:
    S.append((f'Rachats persistants {x*100:g} % par date','Inversé',x,x,x,0,1,0,'=NA_depart','=1/3',0,0,0,1,1,1,0,1,1,''))
for x in [0.05,0.075,0.10,0.125,0.15]:
    S.append((f'Sans plafonnement {x*100:g} % par date','Inversé',x,x,x,0,1,0,'=NA_depart','=1/3',0,0,0,1,1,1,0,0,0,''))
for x,na in [(0.075,0.15),(0.125,0.25),(0.175,0.35),(0.225,0.45),(0.45,0.45)]:
    S.append((f'Appels à douze mois {x*100:g} % de l\'actif net','Inversé','=Rachats_norm','=Rachats_norm','=Rachats_norm',0,1,0,na,(1.0 if x==0.45 else 0.5),0,0,0,1,1,1,0,1,1,''))
for x in [0.10,0.15,0.20,0.25]:
    S.append((f'Porteur unique {x*100:g} % à une date','Porteur',x,'=Rachats_norm','=Rachats_norm',0,1,0,'=NA_depart','=1/3',0,0,0,1,1,1,0,1,1,''))
NS=len(S)
PROW={'r1':6,'r2':7,'rs':8,'col':9,'dco':10,'dtr':11,'na':12,'p4':13,'usd':14,'ccap':15,'cdet':16,'cevg':17,'cpg':18,'cares':19,'cred':20,'gate':21,'bey':22}
LAB={'r1':'Rachats demandés, 1re date (% AN)','r2':'Rachats demandés, 2e date (% AN)','rs':'Rachats demandés, dates suivantes (% AN)','col':'Collecte (% AN par trimestre)','dco':'Distributions des fonds cibles, coefficient','dtr':'Distributions : trimestres sans versement','na':'Non-appelé de départ (% AN)','p4':'Part du non-appelé appelée sur les quatre premiers trimestres','usd':'Variation du dollar','ccap':'Choc de valorisation, capital','cdet':'Choc de valorisation, dette','cevg':'Fonds evergreen : capacité T1 à T4','cpg':'PG NGI : capacité à partir de T5','cares':'Ares : capacité à partir de T5','cred':'Ligne de crédit (% AN de départ)','gate':'Plafonnement appliqué (1/0)','bey':'Exécution au-delà du plafond (1/0)'}
BLOCK0=28; BSIZE=66

def scol(i): return L(3+i)  # colonne du scénario i dans la table

def build_engine(wb):
    ws=wb.create_sheet('3b Moteur')
    widths={'A':2,'B':46}; 
    for t in range(0,NQ+1): widths[c(t)]=11
    for i in range(NS): widths.setdefault(scol(i),11)
    S_=Sheet(ws,'MOTEUR DE STRESS TESTS DE LIQUIDITÉ - projection trimestrielle sur 36 mois, un bloc par scénario',
             'Toutes les cellules sont des formules alimentées par l\'onglet 0 (hypothèses) et la table des scénarios ci-dessous. Mécanique reprise du classeur du 24/09/2026, recalée sur le prospectus du 30/09/2026.',widths)
    # table des scénarios
    r=5
    ws.cell(r,2,'Hypothèse par scénario').font=F(True,WHITE); ws.cell(r,2).fill=fill(TEAL)
    for i,s in enumerate(S):
        cc=ws.cell(r,3+i,s[0]); cc.font=F(True,WHITE); cc.fill=fill(TEAL); cc.alignment=al('left',wrap=True)
    ws.row_dimensions[r].height=60
    keys=['r1','r2','rs','col','dco','dtr','na','p4','usd','ccap','cdet','cevg','cpg','cares','cred','gate','bey']
    for k in keys:
        rr=PROW[k]; ws.cell(rr,2,LAB[k]).font=F(False,INK)
        for i,s in enumerate(S):
            v=s[2+keys.index(k)]
            cc=ws.cell(rr,3+i,v); cc.font=F(False,INK); cc.fill=fill('FFF8F2')
            cc.number_format='0.0 %' if k in('r1','r2','rs','col','na','usd','ccap','cdet','cred') else ('0.00' if k=='p4' else '0.0')
        for col in range(2,3+NS): ws.cell(rr,col).border=B_THIN
    ws.cell(23,2,'Famille').font=F(False,INK); ws.cell(24,2,'Lecture').font=F(False,INK)
    for i,s in enumerate(S):
        ws.cell(23,3+i,s[1]).font=F(False,GREY); ws.cell(24,3+i,s[-1]).font=F(False,GREY); ws.cell(24,3+i).alignment=al('left',wrap=True)
    ws.row_dimensions[24].height=52
    # blocs
    summary=[]
    for i,s in enumerate(S):
        b=BLOCK0+i*BSIZE; sc=scol(i)
        def P(k): return f'${sc}${PROW[k]}'
        R=lambda o: b+o
        def put(o,label,fn,t0=None,fmt=None,bold=False):
            lab=ws.cell(R(o),2,label); lab.font=F(bold,NAVY if bold else INK)
            if t0 is not None:
                cc=ws.cell(R(o),COL0,t0); cc.font=F(False,INK)
                if fmt: cc.number_format=fmt
            for t in range(1,NQ+1):
                cc=ws.cell(R(o),COL0+t,fn(t)); cc.font=F(False,INK)
                if fmt: cc.number_format=fmt
        # en-tête de bloc
        h=ws.cell(R(0),2,f'={sc}$5'); h.font=F(True,WHITE); h.fill=fill(NAVY)
        for col in range(3,COL0+NQ+1): ws.cell(R(0),col).fill=fill(NAVY)
        ws.cell(R(0),3,'N°').font=F(True,WHITE); ws.cell(R(0),4,i+1).font=F(True,WHITE)
        put(1,'Trimestre',lambda t:t,0)
        put(2,'Date',lambda t:f'=EOMONTH(Date_depart,3*{c(t)}{R(1)})','=Date_depart','dd/mm/yyyy')
        put(3,'Date de rachat (30 juin, 31 décembre)',lambda t:f'=IF(OR(MONTH({c(t)}{R(2)})=6,MONTH({c(t)}{R(2)})=12),1,0)',0)
        put(4,'N° de date de rachat',lambda t:f'={c(t-1)}{R(4)}+{c(t)}{R(3)}',0)
        put(5,'Taux de rachat demandé (% AN)',lambda t:f'=IF({c(t)}{R(3)}=1,IF({c(t-1)}{R(15)}=1,0,IF({c(t)}{R(4)}=1,{P("r1")},IF({c(t)}{R(4)}=2,{P("r2")},{P("rs")}))),0)',None,'0.0 %')
        put(6,'Actif net en début de trimestre',lambda t:f'={c(t-1)}{R(50)}',None,'0.00')
        put(7,'Souscriptions',lambda t:f'=IF({c(t-1)}{R(15)}=1,0,{P("col")}*{c(t)}{R(6)})',None,'0.00')
        put(8,'Demandes de rachat nouvelles',lambda t:f'={c(t)}{R(5)}*{c(t)}{R(6)}',None,'0.00')
        put(9,'File reportée',lambda t:f'={c(t-1)}{R(17)}',None,'0.00')
        put(10,'Demande totale',lambda t:f'={c(t)}{R(8)}+{c(t)}{R(9)}',None,'0.00')
        put(11,'Plafond (net des souscriptions)',lambda t:f'=IF({P("gate")}=1,Gate*{c(t)}{R(6)}+{c(t)}{R(7)},{c(t)}{R(10)})',None,'0.00')
        put(12,'Exécution dans le plafond',lambda t:f'=IF({c(t)}{R(3)}=1,MIN({c(t)}{R(10)},{c(t)}{R(11)}),0)',None,'0.00')
        def marge(t):
            nxt=f'SUM({c(t+1)}{R(38)}:{c(min(t+2,NQ))}{R(38)})' if t<NQ else '0'
            return f'=MAX(0,{c(t)}{R(41)}+{c(t)}{R(32)}+{c(t)}{R(33)}-Poche_min_exec*{c(t)}{R(6)}-{c(t)}{R(12)}-{nxt}+MAX(0,MIN({P("cred")}*AN_cible,Emprunt_max*{c(t)}{R(6)})-{c(t-1)}{R(49)}))'
        put(13,'Marge pour exécuter au-delà du plafond',marge,None,'0.00')
        put(14,'Exécution au-delà du plafond',lambda t:f'=IF(AND({c(t)}{R(3)}=1,{P("gate")}=1,{P("bey")}=1),MIN({c(t)}{R(10)}-{c(t)}{R(12)},{c(t)}{R(13)}),0)',None,'0.00')
        put(15,'Suspension des rachats (1 = oui)',lambda t:f'=IF({c(t-1)}{R(15)}=1,1,IF(AND({c(t)}{R(3)}=1,{P("gate")}=1,{c(t-1)}{R(19)}>=Gate_dates,{c(t)}{R(10)}-{c(t)}{R(12)}-{c(t)}{R(14)}>0.000001),1,0))',0)
        put(16,'Rachats payés',lambda t:f'=IF({c(t)}{R(15)}=1,0,{c(t)}{R(12)}+{c(t)}{R(14)})',None,'0.00')
        put(17,'File après exécution',lambda t:f'={c(t)}{R(10)}-{c(t)}{R(16)}',0,'0.00')
        put(18,'Date plafonnée (1 = oui)',lambda t:f'=IF(AND({c(t)}{R(3)}=1,{c(t)}{R(15)}=0,{c(t)}{R(17)}>0.000001),1,0)',None)
        put(19,'Dates plafonnées consécutives',lambda t:f'=IF({c(t)}{R(3)}=1,IF({c(t)}{R(18)}=1,{c(t-1)}{R(19)}+1,0),{c(t-1)}{R(19)})',0)
        def choc(t,cap,usdflag): return f'IF({c(t)}{R(1)}=1,(1-({cap}*{P("ccap")}+(1-{cap})*{P("cdet")}))*(1+{usdflag}*{P("usd")}),1)'
        tot=lambda t: f'({c(t-1)}{R(20)}+{c(t-1)}{R(21)}+{c(t-1)}{R(22)}+{c(t-1)}{R(23)})'
        share=lambda t,o: f'IF({tot(t)}>0,{c(t-1)}{R(o)}/{tot(t)},0)'
        put(20,'MAPIF II - valeur',lambda t:f'={c(t-1)}{R(20)}*(1+Rend_MAPIF/4)*{choc(t,"Cap_MAPIF","1")}+{c(t)}{R(30)}-{c(t)}{R(24)}-{c(t)}{R(45)}*{share(t,20)}','=AN_cible*Ws_MAPIF','0.00')
        put(21,'MSIG 3 - valeur',lambda t:f'={c(t-1)}{R(21)}*(1+Rend_MSIG/4)*{choc(t,"Cap_MSIG","MSIG_USD")}+{c(t)}{R(31)}-{c(t)}{R(25)}-{c(t)}{R(45)}*{share(t,21)}','=AN_cible*Ws_MSIG','0.00')
        evg_share=lambda t,o: f'IF(({c(t)}{R(32)}+{c(t)}{R(33)})>0,{c(t)}{R(o)}/({c(t)}{R(32)}+{c(t)}{R(33)}),0)'
        put(22,'PG NGI - valeur',lambda t:f'={c(t-1)}{R(22)}*(1+Rend_PG/4)*{choc(t,"Cap_PG","0")}-{c(t)}{R(26)}-{c(t)}{R(42)}*{evg_share(t,32)}+{c(t)}{R(47)}/2-{c(t)}{R(45)}*{share(t,22)}','=AN_cible*Ws_PG','0.00')
        put(23,'Ares AGI - valeur',lambda t:f'={c(t-1)}{R(23)}*(1+Rend_ARES/4)*{choc(t,"Cap_ARES","0")}-{c(t)}{R(27)}-{c(t)}{R(42)}*{evg_share(t,33)}/(1-IF({c(t)}{R(1)}<=Deduc_trim,Deduc_ares,0))+{c(t)}{R(47)}/2-{c(t)}{R(45)}*{share(t,23)}','=AN_cible*Ws_ARES','0.00')
        dist=lambda t,o,nm: f'={c(t-1)}{R(o)}*{nm}/4*IF({c(t)}{R(1)}<={P("dtr")},0,{P("dco")})'
        put(24,'MAPIF II - distributions',lambda t:dist(t,20,'Dist_MAPIF'),None,'0.00')
        put(25,'MSIG 3 - distributions',lambda t:dist(t,21,'Dist_MSIG'),None,'0.00')
        put(26,'PG NGI - distributions',lambda t:dist(t,22,'Dist_PG'),None,'0.00')
        put(27,'Ares AGI - distributions',lambda t:dist(t,23,'Dist_ARES'),None,'0.00')
        put(28,'MAPIF II - non-appelé',lambda t:f'={c(t-1)}{R(28)}*IF({c(t)}{R(1)}=1,1+{P("usd")},1)-{c(t)}{R(30)}',f'=AN_cible*{P("na")}*Part_MAPIF_NA','0.00')
        put(29,'MSIG 3 - non-appelé',lambda t:f'={c(t-1)}{R(29)}*IF({c(t)}{R(1)}=1,1+MSIG_USD*{P("usd")},1)-{c(t)}{R(31)}',f'=AN_cible*{P("na")}*(1-Part_MAPIF_NA)','0.00')
        put(30,'MAPIF II - appels',lambda t:f'=MIN({c(t-1)}{R(28)}*IF({c(t)}{R(1)}=1,1+{P("usd")},1),$C{R(28)}*(1+{P("usd")})*IF({c(t)}{R(1)}<=4,{P("p4")}/4,Appels_base))',None,'0.00')
        put(31,'MSIG 3 - appels',lambda t:f'=MIN({c(t-1)}{R(29)}*IF({c(t)}{R(1)}=1,1+MSIG_USD*{P("usd")},1),$C{R(29)}*(1+MSIG_USD*{P("usd")})*IF({c(t)}{R(1)}<=4,{P("p4")}/4,Appels_base))',None,'0.00')
        put(32,'PG NGI - capacité de rachat',lambda t:f'={c(t-1)}{R(22)}*Capacite_evg*IF({c(t)}{R(1)}<=4,{P("cevg")},{P("cpg")})',None,'0.00')
        put(33,'Ares AGI - capacité de rachat, nette de la déduction',lambda t:f'={c(t-1)}{R(23)}*Capacite_evg*IF({c(t)}{R(1)}<=4,{P("cevg")},{P("cares")})*(1-IF({c(t)}{R(1)}<=Deduc_trim,Deduc_ares,0))',None,'0.00')
        put(34,'Trésorerie en début de trimestre',lambda t:f'={c(t-1)}{R(48)}',None,'0.00')
        put(35,'Produits de trésorerie',lambda t:f'={c(t)}{R(34)}*Taux_treso/4',None,'0.00')
        put(36,'Distributions reçues',lambda t:f'=SUM({c(t)}{R(24)}:{c(t)}{R(27)})',None,'0.00')
        put(37,'Acomptes aux parts de distribution',lambda t:f'=IF(AND(Distrib_tension=1,OR({c(t)}{R(9)}>0.000001,{c(t-1)}{R(51)}<Liq_min,{c(t-1)}{R(15)}=1)),0,MIN(Dist_porteurs/4*Part_d*{c(t)}{R(6)},{c(t)}{R(36)}))',None,'0.00')
        put(38,'Appels de fonds',lambda t:f'={c(t)}{R(30)}+{c(t)}{R(31)}',None,'0.00')
        put(39,'Frais',lambda t:f'={c(t)}{R(6)}*Frais_an/4',None,'0.00')
        put(40,'Intérêts d\'emprunt',lambda t:f'={c(t-1)}{R(49)}*Taux_emprunt/4',None,'0.00')
        put(41,'Trésorerie avant rachats',lambda t:f'={c(t)}{R(34)}+{c(t)}{R(35)}+{c(t)}{R(36)}-{c(t)}{R(37)}+{c(t)}{R(7)}-{c(t)}{R(38)}-{c(t)}{R(39)}-{c(t)}{R(40)}',None,'0.00')
        put(42,'Rachats auprès des fonds evergreen',lambda t:f'=MIN({c(t)}{R(32)}+{c(t)}{R(33)},MAX(0,Liq_cible*{c(t)}{R(6)}-({c(t)}{R(41)}-{c(t)}{R(16)})))',None,'0.00')
        put(43,'Trésorerie après rachats',lambda t:f'={c(t)}{R(41)}-{c(t)}{R(16)}+{c(t)}{R(42)}',None,'0.00')
        put(44,'Emprunt : tirage (+) ou remboursement (-)',lambda t:f'=MIN(MAX(0,MIN({P("cred")}*AN_cible,Emprunt_max*{c(t)}{R(6)})-{c(t-1)}{R(49)}),MAX(0,IF({P("cred")}>0,Poche_min_exec*{c(t)}{R(6)},0)-{c(t)}{R(43)}))-IF({c(t)}{R(43)}>Liq_cible*{c(t)}{R(6)},MIN({c(t-1)}{R(49)},{c(t)}{R(43)}-Liq_cible*{c(t)}{R(6)}),0)',None,'0.00')
        put(45,'Cession sur le marché secondaire (valeur cédée)',lambda t:f'=MIN({tot(t)},MAX(0,-({c(t)}{R(43)}+{c(t)}{R(44)}))/(1-Decote))',None,'0.00')
        put(46,'Décote de cession',lambda t:f'={c(t)}{R(45)}*Decote',None,'0.00')
        put(47,'Réinvestissement dans les fonds evergreen',lambda t:f'=IF(AND({c(t)}{R(17)}<0.000001,{c(t-1)}{R(49)}+{c(t)}{R(44)}<0.000001),MAX(0,{c(t)}{R(43)}+{c(t)}{R(44)}-Seuil_reinv*{c(t)}{R(6)}),0)',None,'0.00')
        put(48,'Trésorerie en fin de trimestre',lambda t:f'={c(t)}{R(43)}+{c(t)}{R(44)}+{c(t)}{R(45)}-{c(t)}{R(46)}-{c(t)}{R(47)}','=AN_cible*Ws_LIQ','0.00')
        put(49,'Emprunt en fin de trimestre',lambda t:f'={c(t-1)}{R(49)}+{c(t)}{R(44)}',0,'0.00')
        put(50,'Actif net en fin de trimestre',lambda t:f'={c(t)}{R(20)}+{c(t)}{R(21)}+{c(t)}{R(22)}+{c(t)}{R(23)}+{c(t)}{R(48)}-{c(t)}{R(49)}',f'=C{R(20)}+C{R(21)}+C{R(22)}+C{R(23)}+C{R(48)}-C{R(49)}','0.00',bold=True)
        put(51,'Poche d\'actifs liquides (% actif)',lambda t:f'=IF(({c(t)}{R(50)}+{c(t)}{R(49)})>0,{c(t)}{R(48)}/({c(t)}{R(50)}+{c(t)}{R(49)}),0)',f'=IF((C{R(50)}+C{R(49)})>0,C{R(48)}/(C{R(50)}+C{R(49)}),0)','0.0 %')
        put(52,'File d\'attente (% AN)',lambda t:f'=IF({c(t)}{R(50)}>0,{c(t)}{R(17)}/{c(t)}{R(50)},0)',f'=0','0.0 %')
        def couv(t):
            e=min(t+4,NQ)
            if t>=NQ: return '=9.99'
            rng=lambda o: f'SUM({c(t+1)}{R(o)}:{c(e)}{R(o)})'
            den=f'({c(t)}{R(17)}+{rng(8)}+{rng(38)}+{rng(39)}+{rng(37)})'
            num=f'({c(t)}{R(48)}+{rng(36)}+SUM({c(t+1)}{R(32)}:{c(e)}{R(33)})+MAX(0,MIN({P("cred")}*AN_cible,Emprunt_max*{c(t)}{R(50)})-{c(t)}{R(49)}))'
            return f'=IF({den}<=0,9.99,MIN(9.99,{num}/{den}))'
        put(53,'Couverture de liquidité à 12 mois',couv,None,'0.00')
        ws.cell(R(53),COL0,couv(0).replace(f'{c(0)}{R(17)}',f'C{R(17)}')).number_format='0.00'
        put(54,'Sur-engagement (% AN)',lambda t:f'=({c(t)}{R(20)}+{c(t)}{R(21)}+{c(t)}{R(22)}+{c(t)}{R(23)}+{c(t)}{R(28)}+{c(t)}{R(29)})/{c(t)}{R(50)}',f'=(C{R(20)}+C{R(21)}+C{R(22)}+C{R(23)}+C{R(28)}+C{R(29)})/C{R(50)}','0.0 %')
        put(55,'Levier, méthode de l\'engagement',lambda t:f'=({c(t)}{R(50)}+{c(t)}{R(49)})/{c(t)}{R(50)}',f'=(C{R(50)}+C{R(49)})/C{R(50)}','0.0 %')
        put(56,'Exposition de change non couverte (% AN)',lambda t:f'=({c(t)}{R(20)}+{c(t)}{R(21)}*MSIG_USD)/{c(t)}{R(50)}',f'=(C{R(20)}+C{R(21)}*MSIG_USD)/C{R(50)}','0.0 %')
        put(57,'Triple gel (collecte, distributions, fonds evergreen)',lambda t:f'=IF(AND({P("col")}<0.5*Collecte_norm,IF({c(t)}{R(1)}<={P("dtr")},0,{P("dco")})<0.5,IF({c(t)}{R(1)}<=4,{P("cevg")},({P("cpg")}+{P("cares")})/2)<0.7),1,0)',None)
        put(58,'aide : n° de la première date plafonnée',lambda t:f'=IF({c(t)}{R(18)}=1,{c(t)}{R(4)},99)',99)
        put(59,'aide : n° de la date de résorption',lambda t:f'=IF(AND({c(t)}{R(3)}=1,{c(t)}{R(4)}>MIN($D{R(58)}:${c(NQ)}{R(58)}),{c(t)}{R(17)}<=0.000001),{c(t)}{R(4)},99)',99)
        put(60,'aide : date de suspension',lambda t:f'=IF({c(t)}{R(15)}=1,{c(t)}{R(2)},9999999)',9999999)
        put(61,'aide : date du premier triple gel',lambda t:f'=IF({c(t)}{R(57)}=1,{c(t)}{R(2)},9999999)',9999999)
        for o in (58,59,60,61):
            for col in range(2,COL0+NQ+1): ws.cell(R(o),col).font=F(False,GREY)
        for o in range(1,62):
            for col in range(2,COL0+NQ+1): ws.cell(R(o),col).border=B_THIN
        rng=lambda o: f"'3b Moteur'!$D${R(o)}:${c(NQ)}${R(o)}"
        first=f"MIN({rng(58)})"
        summary.append(dict(
            nom=f"='3b Moteur'!{sc}$5", fam=f"='3b Moteur'!{sc}$23", lecture=f"='3b Moteur'!{sc}$24",
            dates=f"=SUM({rng(18)})", file=f"=MAX({rng(52)})",
            resorb=f"=IF({first}=99,0,IF(MIN({rng(59)})=99,99,6*(MIN({rng(59)})-{first})))",
            poche=f"=MIN({rng(51)})", couv=f"=MIN('3b Moteur'!$C${R(53)},{rng(53)})",
            sureng=f"=MAX({rng(54)})", change=f"=MAX({rng(56)})", cession=f"=SUM({rng(45)})/1000000",
            susp=f"=IF(MIN({rng(60)})>=9999999,\"-\",RIGHT(\"0\"&MONTH(MIN({rng(60)})),2)&\"/\"&YEAR(MIN({rng(60)})))",
            gel=f"=IF(MIN({rng(61)})>=9999999,\"-\",RIGHT(\"0\"&MONTH(MIN({rng(61)})),2)&\"/\"&YEAR(MIN({rng(61)})))",
            emprunt=f"=MAX({rng(49)})", an_fin=f"='3b Moteur'!${c(NQ)}${R(50)}",
            couv0=f"'3b Moteur'!$C${R(53)}", res0=f"('3b Moteur'!$C${R(48)}+SUM('3b Moteur'!$D${R(36)}:$G${R(36)})+SUM('3b Moteur'!$D${R(32)}:$G${R(33)}))"))
    ws.sheet_view.showGridLines=False
    return summary
