# -*- coding: utf-8 -*-
"""Complète le classeur KPI_Registre_Incidents_Fundcraft_France avec les indicateurs de service
(prise en compte, résolution, réouverture, escalades, DORA) du modèle Registre_Incidents_KPI_Fundcraft.

Le classeur source contient un graphique : il est modifié par LibreOffice (UNO) et non par openpyxl,
pour conserver graphique, validations, mises en forme conditionnelles et ajuster les références.
Usage : python complete_france.py source.xlsx sortie.xlsx sortie_redline.xlsx
"""
import subprocess
import sys
import time

import uno
from com.sun.star.beans import PropertyValue

NBSP = " "
CREAM, YELLOW = 0xFFF8F2, 0xFFFF00
FIRST, LAST = 7, 306  # lignes du registre (1-based)

FONDS = ["AirFund Conviction Value Capital", "Openstone Infraworld", "Aletheon Growth III S.L.P.",
         "Partners Group PEO Eltif Feeder", "Otentiq Private Equity X"]
RENOMMAGE = {"FPS Openstone Infraworld": "Openstone Infraworld", "Aletheon Growth III": "Aletheon Growth III S.L.P.",
             "Otentiq Private Equity X SLP": "Otentiq Private Equity X",
             "Partners Group Private Equity Opportunities ELTIF Feeder": "Partners Group PEO Eltif Feeder"}


def fr(s):
    for a in (" :", " ;", " %", " ?", " !"):
        s = s.replace(a, NBSP + a[1:])
    return s


def pv(n, v):
    p = PropertyValue()
    p.Name, p.Value = n, v
    return p


def conv(f):
    """Formule écrite avec des virgules -> séparateurs ';' de l'API LibreOffice (hors chaînes)."""
    out, q = [], False
    for ch in f:
        if ch == '"':
            q = not q
        out.append(";" if (ch == "," and not q) else ch)
    return "".join(out)


def col_letter(i):
    s, i = "", i + 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


class Calc:
    def __init__(self, path):
        self.proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--nologo", "--norestore",
                                      "--accept=pipe,name=fcpipe2;urp;"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        ctx = uno.getComponentContext()
        res = ctx.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", ctx)
        for _ in range(120):
            try:
                c = res.resolve("uno:pipe,name=fcpipe2;urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(0.5)
        self.desk = c.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", c)
        self.doc = self.desk.loadComponentFromURL(uno.systemPathToFileUrl(path), "_blank", 0, (pv("Hidden", True),))
        self.loc = uno.createUnoStruct("com.sun.star.lang.Locale")
        self.loc.Language, self.loc.Country = "en", "US"
        self.red = []  # (feuille, l, t, r, b) en 0-based pour la version redline

    def sheet(self, n):
        return self.doc.Sheets.getByName(n)

    def fmt(self, code):
        nf = self.doc.NumberFormats
        k = nf.queryKey(code, self.loc, False)
        return k if k != -1 else nf.addNew(code, self.loc)

    def addr(self, sh, c, r):
        a = uno.createUnoStruct("com.sun.star.table.CellAddress")
        a.Sheet, a.Column, a.Row = sh.RangeAddress.Sheet, c, r
        return a

    def copy(self, sh, l, t, r, b, dc, dr):
        sh.copyRange(self.addr(sh, dc, dr), sh.getCellRangeByPosition(l, t, r, b).getRangeAddress())

    def mark(self, name, l, t, r, b):
        self.red.append((name, l, t, r, b))

    def name(self, nm, content):
        nr = self.doc.NamedRanges
        if nr.hasByName(nm):
            nr.getByName(nm).setContent(content)
        else:
            nr.addNewByName(nm, content, self.addr(self.sheet("Données"), 0, 0), 0)

    def save(self, path):
        self.doc.calculateAll()
        self.doc.storeToURL(uno.systemPathToFileUrl(path), (pv("FilterName", "Calc MS Excel 2007 XML"),))

    def close(self):
        self.doc.close(True)
        try:
            self.desk.terminate()
        except Exception:
            pass
        self.proc.wait(timeout=30)


def find_col(sh, label, row=5):
    for c in range(0, 120):
        if sh.getCellByPosition(c, row).getString() == label:
            return c
    raise KeyError(label)


def find_row(sh, label, col=1):
    for r in range(0, 200):
        if sh.getCellByPosition(col, r).getString() == label:
            return r
    raise KeyError(label)


# --------------------------------------------------------------------------- registre
NOUVEAUX_CHAMPS = [  # (après la colonne, nouvelles colonnes [(libellé, gabarit)])
    ("Date de clôture", [("Réouvert", "Dépassement actif ou passif")]),
    ("Statut", [("Résolution (date et heure)", "Détection (date et heure)")]),
    ("Date d'escalade", [("Nombre de réaffectations", "Investisseurs lésés")]),
    ("Notification réalisée (date et heure)", [("Classification DORA (date et heure)", "Détection (date et heure)"),
                                                ("Rapport intermédiaire DORA (date et heure)", "Détection (date et heure)"),
                                                ("Rapport final DORA (date)", "Date de survenance")]),
    ("Déclaration au registre", [("Prise en compte (date et heure)", "Détection (date et heure)")]),
]


def inserer_champs(k):
    sh = k.sheet("Registre")
    for apres, cols in NOUVEAUX_CHAMPS:
        i = find_col(sh, apres) + 1
        sh.Columns.insertByIndex(i, len(cols))
        for j, (lab, gab) in enumerate(cols):
            c = i + j
            g = find_col(sh, gab)
            k.copy(sh, g, 4, g, LAST - 1, c, 4)
            sh.getCellRangeByPosition(c, FIRST - 1, c, LAST - 1).clearContents(1 | 2 | 4 | 16)
            sh.getCellByPosition(c, 5).setString(lab)
            sh.Columns.getByIndex(c).Width = max(sh.Columns.getByIndex(g).Width, 2600)
            if lab == "Réouvert":
                rng = sh.getCellRangeByPosition(c, FIRST - 1, c, LAST - 1)
                v = rng.Validation
                v.Type = uno.Enum("com.sun.star.sheet.ValidationType", "LIST")
                v.Formula1 = '"Oui";"Non"'
                v.ShowList = 1
                v.IgnoreBlankCells = True
                v.ShowErrorMessage = True
                v.ErrorTitle = "Valeur hors liste"
                v.ErrorMessage = "Choisir Oui ou Non."
                rng.Validation = v
    for _, cols in NOUVEAUX_CHAMPS:
        for lab, _ in cols:
            c = find_col(sh, lab)
            k.mark("Registre", c, 4, c, LAST - 1)
    for old, new in RENOMMAGE.items():
        cf = find_col(sh, "Fonds ou véhicule")
        for r in range(FIRST - 1, LAST):
            x = sh.getCellByPosition(cf, r)
            if x.getType().value == "TEXT" and x.getString() == old:
                x.setString(new)
                k.mark("Registre", cf, r, cf, r)


MOTEUR = [  # (libellé, nom, format, formule avec {X} = lettre de colonne et {r} = ligne)
    ("Prise en compte (h ouvrées)", "R_hPEC", "0.0",
     '=IF(OR({H}{r}="",{PEC}{r}=""),"",IF({PEC}{r}<{H}{r},"",ROUND(((NETWORKDAYS({H}{r},{PEC}{r})-1)*(HFer-HOuv)'
     '+IF(NETWORKDAYS({PEC}{r},{PEC}{r}),MEDIAN(MOD({PEC}{r},1),HFer,HOuv),HFer)'
     '-IF(NETWORKDAYS({H}{r},{H}{r}),MEDIAN(MOD({H}{r},1),HFer,HOuv),HOuv))*24,2)))'),
    ("Cible de prise en compte (h ouvrées)", None, "0", '=IF(${H}{r}="","",INDEX(SLA_PEC,{RANG}{r}))'),
    ("Prise en compte dans le délai", "R_okPEC", "0", '=IF(OR({hPEC}{r}="",{cPEC}{r}=""),"",IF({hPEC}{r}<={cPEC}{r},1,0))'),
    ("Délai de résolution (j.o.)", "R_jRes", "0",
     '=IF(OR(${H}{r}="",{RES}{r}=""),"",IF(INT({RES}{r})<INT({H}{r}),"",NETWORKDAYS(INT({H}{r}),INT({RES}{r}))-1))'),
    ("Cible de résolution (j.o.)", None, "0", '=IF(${H}{r}="","",INDEX(SLA_RES,{RANG}{r}))'),
    ("Résolution dans le délai", "R_okRes", "0", '=IF(OR({jRes}{r}="",{cRes}{r}=""),"",IF({jRes}{r}<={cRes}{r},1,0))'),
    ("Résolu sans escalade au-delà du métier", "R_ResN1", "0",
     '=IF(OR(${H}{r}="",{RES}{r}=""),"",IF({NREQ}{r}<=1,IF({NATT}{r}<=1,1,0),""))'),
    ("Réouvert (1/0)", "R_ReoN", "0", '=IF(OR(${H}{r}="",{RES}{r}=""),"",IF({REO}{r}="Oui",1,0))'),
    ("Temps en escalade (j.o.)", "R_jEsc", "0",
     '=IF(OR({Y}{r}="",{RES}{r}=""),"",IF(INT({RES}{r})<INT({Y}{r}),"",NETWORKDAYS(INT({Y}{r}),INT({RES}{r}))-1))'),
    ("Action réalisée à l'échéance", "R_ActOK", "0",
     '=IF(OR(${H}{r}="",{ECH}{r}=""),"",IF({REAL}{r}="",IF({ECH}{r}<DateArrete,0,""),IF({REAL}{r}<={ECH}{r},1,0)))'),
    ("DORA - échéance de notification initiale", None, "dd/mm/yyyy hh:mm",
     '=IF(OR(${H}{r}="",{NOTIF}{r}<>Notif_DORA),"",IF({DCLAS}{r}="",{H}{r}+INDEX(Notif_H,MATCH(Notif_DORA,Notif_Lib,0))/24,'
     'IF({DCLAS}{r}-{H}{r}<=INDEX(Notif_H,MATCH(Notif_DORA,Notif_Lib,0))/24,MIN({DCLAS}{r}+DORA_H_CLA/24,'
     '{H}{r}+INDEX(Notif_H,MATCH(Notif_DORA,Notif_Lib,0))/24),{DCLAS}{r}+DORA_H_CLA/24)))'),
    ("DORA - notification initiale", None, "0",
     '=IF({DoraDL}{r}="","",IF({V}{r}="",IF({DoraDL}{r}<DateArrete+1,0,""),IF({V}{r}<={DoraDL}{r},1,0)))'),
    ("DORA - rapport intermédiaire", None, "0",
     '=IF(OR({DoraDL}{r}="",{V}{r}=""),"",IF({DINTER}{r}="",IF({V}{r}+DORA_H_INT/24<DateArrete+1,0,""),'
     'IF({DINTER}{r}<={V}{r}+DORA_H_INT/24,1,0)))'),
    ("DORA - rapport final", None, "0",
     '=IF(OR({DoraDL}{r}="",{DINTER}{r}=""),"",IF({DFIN}{r}="",IF(EDATE(INT({DINTER}{r}),DORA_M_FIN)<DateArrete,0,""),'
     'IF({DFIN}{r}<=EDATE(INT({DINTER}{r}),DORA_M_FIN),1,0)))'),
    ("DORA - notifications conformes", "R_DoraOK", "0",
     '=IF({DoraDL}{r}="","",IF(COUNT({DoraI}{r}:{DoraF}{r})=0,"",MIN({DoraI}{r}:{DoraF}{r})))'),
]
CLES_MOTEUR = ["hPEC", "cPEC", "okPEC", "jRes", "cRes", "okRes", "ResN1", "ReoN", "jEsc", "ActOK",
               "DoraDL", "DoraI", "DoraM", "DoraF", "DoraOK"]


def ajouter_moteur(k):
    sh = k.sheet("Registre")
    L = lambda lab: col_letter(find_col(sh, lab))  # noqa: E731
    m = {"H": L("Détection (date et heure)"), "PEC": L("Prise en compte (date et heure)"), "RES": L("Résolution (date et heure)"),
         "CLO": L("Date de clôture"), "Y": L("Date d'escalade"), "REO": L("Réouvert"), "ECH": L("Échéance de l'action"),
         "REAL": L("Action réalisée le"), "NOTIF": L("Notification à une autorité"), "V": L("Notification réalisée (date et heure)"),
         "DCLAS": L("Classification DORA (date et heure)"), "DINTER": L("Rapport intermédiaire DORA (date et heure)"),
         "DFIN": L("Rapport final DORA (date)"), "RANG": L("Gravité retenue (rang)"), "NREQ": L("Niveau requis (rang)"),
         "NATT": L("Niveau atteint (rang)")}
    gab = find_col(sh, "Rang liste comité")
    debut_moteur = find_col(sh, "Score financier")
    for j, cle in enumerate(CLES_MOTEUR):
        m[cle] = col_letter(gab + 1 + j)
    for j, (lab, nm, code, f) in enumerate(MOTEUR):
        c = gab + 1 + j
        k.copy(sh, gab, 4, gab, LAST - 1, c, 4)
        sh.getCellRangeByPosition(c, FIRST - 1, c, LAST - 1).clearContents(1 | 2 | 4 | 16)
        sh.getCellByPosition(c, 5).setString(lab)
        sh.Columns.getByIndex(c).Width = sh.Columns.getByIndex(gab).Width
        rng = sh.getCellRangeByPosition(c, FIRST - 1, c, LAST - 1)
        rng.setFormulaArray(tuple((conv(f.format(r=r, **m)),) for r in range(FIRST, LAST + 1)))
        rng.NumberFormat = k.fmt(code)
        if nm:
            k.name(nm, f"$Registre.${col_letter(c)}${FIRST}:${col_letter(c)}${LAST}")
        k.mark("Registre", c, 4, c, LAST - 1)
    fin = gab + len(MOTEUR)
    ori = uno.Enum("com.sun.star.table.TableOrientation", "COLUMNS")
    a = sh.getCellRangeByPosition(debut_moteur, 0, gab, 0).getRangeAddress()
    try:
        sh.ungroup(a, ori)
    except Exception:
        pass
    sh.group(sh.getCellRangeByPosition(debut_moteur, 0, fin, 0).getRangeAddress(), ori)
    for c in range(debut_moteur, fin + 1):
        sh.Columns.getByIndex(c).IsVisible = False
    for nm, lab in (("R_PEC", "Prise en compte (date et heure)"), ("R_Res", "Résolution (date et heure)"), ("R_Reo", "Réouvert"),
                    ("R_Reaf", "Nombre de réaffectations"), ("R_DClas", "Classification DORA (date et heure)"),
                    ("R_DInter", "Rapport intermédiaire DORA (date et heure)"), ("R_DFin", "Rapport final DORA (date)"),
                    ("R_NivAtt", "Niveau atteint (rang)"), ("R_NivReq", "Niveau requis (rang)")):
        cl = L(lab)
        k.name(nm, f"$Registre.${cl}${FIRST}:${cl}${LAST}")


# --------------------------------------------------------------------------- données
KRI = [  # libellé, sens, vigilance, alerte, format, source
    ("Concentration : part de la première sous-catégorie", "Plafond", 0.30, 0.50, "0.0%", "Hypothèse - à confirmer"),
    ("Temps moyen passé en escalade", "Plafond", 3, 5, "0.0", "Hypothèse - à confirmer"),
    ("Incidents réaffectés au moins 2 fois", "Plafond", 0.10, 0.20, "0.0%", "Pratique ITSM (hop rate) - à confirmer"),
    ("Notifications DORA conformes (4 h, 24 h, 72 h, 1 mois)", "Plancher", 1, 1, "0.0%", "RD (UE) 2025/301, art. 5"),
    ("Délai moyen de prise en compte", "Plafond", 4, 9, "0.0", "Pratique ITSM - à confirmer"),
    ("Prises en compte dans le délai cible", "Plancher", 0.95, 0.85, "0.0%", "Pratique ITSM - à confirmer"),
    ("Délai moyen de résolution", "Plafond", 5, 10, "0.0", "Hypothèse - à confirmer"),
    ("Délai de résolution - 90e centile", "Plafond", 10, 20, "0.0", "Hypothèse - à confirmer"),
    ("Résolutions dans le délai cible", "Plancher", 0.90, 0.80, "0.0%", "Pratique ITSM - à confirmer"),
    ("Résolutions sans escalade au-delà du responsable métier", "Plancher", 0.80, 0.65, "0.0%", "Pratique ITSM - à confirmer"),
    ("Taux de réouverture", "Plafond", 0.05, 0.10, "0.0%", "Pratique ITSM - à confirmer"),
    ("Ratio de clôture (clos / déclarés)", "Plancher", 1.0, 0.8, '0.00"x"', "Pratique ITSM - à confirmer"),
    ("Analyse de cause documentée (significatifs et majeurs clos)", "Plancher", 1.0, 0.9, "0.0%", "Hypothèse - à confirmer"),
    ("Actions correctives réalisées à l'échéance", "Plancher", 0.90, 0.75, "0.0%", "Hypothèse - à confirmer"),
]


def completer_donnees(k):
    sh = k.sheet("Données")
    # Fonds et véhicules
    cf = find_col(sh, "Fonds et véhicules")
    r0 = find_row(sh, "Fundcraft France SAS", cf)
    for i, f in enumerate(FONDS):
        x = sh.getCellByPosition(cf, r0 + 1 + i)
        if x.getString() != f:
            x.setString(f)
            k.mark("Données", cf, r0 + 1 + i, cf, r0 + 1 + i)
    # Seuils des indicateurs
    rl = find_row(sh, "Réclamations traitées dans les délais")
    first = find_row(sh, "Incidents significatifs et majeurs")
    for j, (lab, sens, vig, al, code, src) in enumerate(KRI):
        r = rl + 1 + j
        k.copy(sh, 1, rl, 5, rl, 1, r)
        sh.getCellByPosition(1, r).setString(fr(lab))
        sh.getCellByPosition(2, r).setString(sens)
        for c, v in ((3, vig), (4, al)):
            x = sh.getCellByPosition(c, r)
            x.setValue(v)
            x.NumberFormat = k.fmt(code)
        sh.getCellByPosition(5, r).setString(src)
        k.mark("Données", 1, r, 5, r)
    last = rl + len(KRI)
    for nm, c in (("KRI_Lib", "B"), ("KRI_Sens", "C"), ("KRI_Vig", "D"), ("KRI_Alerte", "E")):
        k.name(nm, f"$Données.${c}${first + 1}:${c}${last + 1}")
    # Délais de service et DORA
    band = find_row(sh, "Délais cibles et seuils")
    hdr = find_row(sh, "Gravité")
    prm = find_row(sh, "Délai de déclaration au registre (j.o.)")
    r = last + 2
    k.copy(sh, 1, band, 8, band, 1, r)
    sh.getCellByPosition(1, r).setString("Délais de service et notifications DORA")
    k.mark("Données", 1, r, 8, r)
    r += 1
    k.copy(sh, 1, hdr, 4, hdr, 1, r)
    for c, t in enumerate(["Gravité", "Prise en compte depuis la détection (h ouvrées)", "Résolution depuis la détection (j.o.)", "Source"]):
        sh.getCellByPosition(1 + c, r).setString(t)
    k.mark("Données", 1, r, 4, r)
    sla = [(18, 20), (9, 10), (4, 5), (2, 2)]
    rs = r + 1
    for i, (pec, res) in enumerate(sla):
        rr = rs + i
        k.copy(sh, 1, hdr + 1, 3, hdr + 1, 1, rr)
        k.copy(sh, 8, hdr + 1, 8, hdr + 1, 4, rr)
        sh.getCellByPosition(1, rr).setFormula(conv(f"=INDEX(Grav_Lib,{i + 1})"))
        for c, v in ((2, pec), (3, res)):
            x = sh.getCellByPosition(c, rr)
            x.setValue(v)
            x.NumberFormat = k.fmt("0")
        sh.getCellByPosition(4, rr).setString("Hypothèse - à confirmer")
        k.mark("Données", 1, rr, 4, rr)
    k.name("SLA_PEC", f"$Données.$C${rs + 1}:$C${rs + 4}")
    k.name("SLA_RES", f"$Données.$D${rs + 1}:$D${rs + 4}")
    r = rs + 5
    params = [("Heure d'ouverture (calcul des heures ouvrées)", 9 / 24, "HH:MM", "Hypothèse - horaires de référence", "HOuv"),
              ("Heure de fermeture (calcul des heures ouvrées)", 18 / 24, "HH:MM", "Hypothèse - horaires de référence", "HFer"),
              ("DORA - notification initiale après classification (heures)", 4, "0", "RD (UE) 2025/301, art. 5", "DORA_H_CLA"),
              ("DORA - rapport intermédiaire après notification initiale (heures)", 72, "0", "RD (UE) 2025/301, art. 5", "DORA_H_INT"),
              ("DORA - rapport final après rapport intermédiaire (mois)", 1, "0", "RD (UE) 2025/301, art. 5", "DORA_M_FIN")]
    for lab, v, code, src, nm in params:
        k.copy(sh, 1, prm, 3, prm, 1, r)
        sh.getCellByPosition(1, r).setString(lab)
        x = sh.getCellByPosition(2, r)
        x.setValue(v)
        x.NumberFormat = k.fmt(code)
        sh.getCellByPosition(3, r).setString(src)
        k.name(nm, f"$Données.$C${r + 1}")
        k.mark("Données", 1, r, 3, r)
        r += 1
    k.copy(sh, 3, prm, 3, prm, 1, r)
    sh.getCellRangeByPosition(1, r, 8, r).merge(True)
    note = sh.getCellByPosition(1, r)
    note.IsTextWrapped = True
    sh.Rows.getByIndex(r).Height = 1300
    note.setString(fr("Conventions : délais mesurés depuis la détection ; résolution = impact contenu (VL corrigée, opération "
                      "régularisée, service rétabli), distincte de la clôture. Les indicateurs de résolution restent vides tant que la "
                      "date de résolution n'est pas saisie. Le délai de 24 h après détection pour la notification initiale DORA est celui "
                      "du tableau des notifications. Jours fériés non exclus, comme dans le reste du classeur."))
    k.mark("Données", 1, r, 1, r)


# --------------------------------------------------------------------------- dashboard
PER = {"D": ("DebutT", "FinTx"), "E": ("DebutT1", "DebutT"), "G": ("Debut12M", "FinTx")}
NOUVEAUX_KPI_ESC = [  # libellé, unité, gabarit, formule, matricielle
    ("Escalades au-delà du responsable métier", "%", "pct",
     '=IFERROR(COUNTIFS(R_Det,">="&{a},R_Det,"<"&{b},R_NivAtt,">=2")/{n},"-")', False),
    ("Escalades à la Direction générale ou au Comité des risques", "%", "pct",
     '=IFERROR(COUNTIFS(R_Det,">="&{a},R_Det,"<"&{b},R_NivAtt,">=3")/{n},"-")', False),
    ("Incidents notifiés à une autorité", "%", "pct",
     '=IFERROR(COUNTIFS(R_Det,">="&{a},R_Det,"<"&{b},R_Notif,"<>",R_Notif,"<>"&Notif_Aucune)/{n},"-")', False),
    ("Temps moyen passé en escalade", "j.o.", "jours", '=IFERROR(AVERAGEIFS(R_jEsc,R_Res,">="&{a},R_Res,"<"&{b}),"-")', False),
    ("Incidents réaffectés au moins 2 fois", "%", "pct",
     '=IFERROR(COUNTIFS(R_Det,">="&{a},R_Det,"<"&{b},R_Reaf,">=2")/COUNTIFS(R_Det,">="&{a},R_Det,"<"&{b},R_Reaf,">=0"),"-")', False),
    ("Notifications DORA conformes (4 h, 24 h, 72 h, 1 mois)", "%", "pct",
     '=IFERROR(AVERAGEIFS(R_DoraOK,R_Det,">="&{a},R_Det,"<"&{b}),"-")', False),
]
NOUVEAUX_KPI_RES = [
    ("Délai moyen de prise en compte", "h ouvr.", "jours", '=IFERROR(AVERAGEIFS(R_hPEC,R_Det,">="&{a},R_Det,"<"&{b}),"-")', False),
    ("Prises en compte dans le délai cible", "%", "pct", '=IFERROR(AVERAGEIFS(R_okPEC,R_Det,">="&{a},R_Det,"<"&{b}),"-")', False),
    ("Délai moyen de résolution", "j.o.", "jours", '=IFERROR(AVERAGEIFS(R_jRes,R_Res,">="&{a},R_Res,"<"&{b}),"-")', False),
    ("Délai de résolution - 90e centile", "j.o.", "jours",
     '=IFERROR(PERCENTILE(IF((R_Res>={a})*(R_Res<{b})*ISNUMBER(R_jRes),R_jRes),0.9),"-")', True),
    ("Résolutions dans le délai cible", "%", "pct", '=IFERROR(AVERAGEIFS(R_okRes,R_Res,">="&{a},R_Res,"<"&{b}),"-")', False),
    ("Résolutions sans escalade au-delà du responsable métier", "%", "pct",
     '=IFERROR(AVERAGEIFS(R_ResN1,R_Res,">="&{a},R_Res,"<"&{b}),"-")', False),
    ("Taux de réouverture", "%", "pct", '=IFERROR(AVERAGEIFS(R_ReoN,R_Res,">="&{a},R_Res,"<"&{b}),"-")', False),
    ("Ratio de clôture (clos / déclarés)", "x", "pct", '=IFERROR(COUNTIFS(R_Cloture,">="&{a},R_Cloture,"<"&{b})/{n},"-")', False),
    ("Analyse de cause documentée (significatifs et majeurs clos)", "%", "pct",
     '=IFERROR(COUNTIFS(R_Grav,">=3",R_Cloture,">="&{a},R_Cloture,"<"&{b},R_Cause,"<>")'
     '/COUNTIFS(R_Grav,">=3",R_Cloture,">="&{a},R_Cloture,"<"&{b}),"-")', False),
    ("Actions correctives réalisées à l'échéance", "%", "pct",
     '=IFERROR(AVERAGEIFS(R_ActOK,R_Echeance,">="&{a},R_Echeance,"<"&{b}),"-")', False),
]
NOUVEAU_KPI_CONC = ("Concentration : part de la première sous-catégorie", "%", "pct",
                    '=IFERROR(MAX(COUNTIFS(R_SC,Ref_SC,R_Det,">="&{a},R_Det,"<"&{b}))/{n},"-")', True)
FORMATS = {"%": ("0.0%", "\\+0.0%;\\-0.0%;0.0%"), "x": ('0.00"x"', '\\+0.00"x";\\-0.00"x";0.00"x"'),
           "j.o.": ("0.0", "\\+0.0;\\-0.0;0.0"), "h ouvr.": ("0.0", "\\+0.0;\\-0.0;0.0")}


def ecrire_kpi(k, sh, r, kpi, tpl, rI):
    lab, unit, _, f, arr = kpi
    k.copy(sh, 1, tpl, 9, tpl, 1, r)
    sh.getCellByPosition(1, r).setString(fr(lab))
    sh.getCellByPosition(2, r).setString(unit)
    v, var = FORMATS[unit]
    for col, (a, b) in PER.items():
        c = ord(col) - 65
        x = sh.getCellByPosition(c, r)
        txt = conv(f.format(a=a, b=b, n=f"{col}{rI + 1}"))
        if arr:
            x.setArrayFormula(txt)
        else:
            x.setFormula(txt)
        x.NumberFormat = k.fmt(v)
    sh.getCellByPosition(5, r).NumberFormat = k.fmt(var)
    for c in (7, 8):
        sh.getCellByPosition(c, r).NumberFormat = k.fmt(v)
    k.mark("Dashboard", 1, r, 9, r)


def completer_dashboard(k):
    sh = k.sheet("Dashboard")
    tpl = {"pct": find_row(sh, "Part d'incidents récurrents"), "jours": find_row(sh, "Délai moyen de détection")}
    band = find_row(sh, "Détection, déclaration et escalade")
    rn = find_row(sh, "Notifications aux autorités dans le délai")
    n_ins = len(NOUVEAUX_KPI_ESC) + 1 + len(NOUVEAUX_KPI_RES)
    sh.Rows.insertByIndex(rn + 1, n_ins)
    rI = find_row(sh, "Incidents déclarés")
    r = rn + 1
    for kpi in NOUVEAUX_KPI_ESC:
        ecrire_kpi(k, sh, r, kpi, tpl[kpi[2]], rI)
        r += 1
    k.copy(sh, 1, band, 9, band, 1, r)
    sh.getCellByPosition(1, r).setString("Prise en compte, résolution et clôture")
    k.mark("Dashboard", 1, r, 9, r)
    r += 1
    for kpi in NOUVEAUX_KPI_RES:
        ecrire_kpi(k, sh, r, kpi, tpl[kpi[2]], rI)
        r += 1
    rf = find_row(sh, "Incidents par fonds ou véhicule géré")
    sh.Rows.insertByIndex(rf + 1, 1)
    # les lignes situées sous l'insertion ont été décalées d'une ligne
    k.red = [(s, l, t + (1 if (s == "Dashboard" and t > rf) else 0), rr, b + (1 if (s == "Dashboard" and b > rf) else 0))
             for s, l, t, rr, b in k.red]
    tpl_pct = find_row(sh, "Part d'incidents récurrents")
    ecrire_kpi(k, sh, rf + 1, NOUVEAU_KPI_CONC, tpl_pct, rI)
    h = sh.Rows.getByIndex(rI).Height
    for rr in [rf + 1] + list(range(rn + 2, rn + 2 + n_ins)):
        sh.Rows.getByIndex(rr).Height = h
    src = None
    for r in range(0, 200):
        s = sh.getCellByPosition(1, r).getString()
        if s.startswith("Sources"):
            src = r
            break
    if src is not None:
        x = sh.getCellByPosition(1, src)
        x.setString(x.getString() + fr(" ; indicateurs de délai et de résolution : pratiques ITSM (ITIL 4), cibles internes à confirmer"))
        k.mark("Dashboard", 1, src, 1, src)


def matrices_origine(src):
    """Texte exact des formules matricielles du Dashboard d'origine, par libellé de ligne et colonne."""
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
    ws = openpyxl.load_workbook(src)["Dashboard"]
    out = {}
    for row in ws.iter_rows():
        for x in row:
            if isinstance(x.value, ArrayFormula):
                out[(ws.cell(x.row, 2).value, x.column - 1)] = x.value.text
    return out


def restaurer_matrices(k, matrices):
    """Les formules matricielles d'origine perdent leur statut lors du déplacement des lignes : on les réécrit
    à l'identique à partir du fichier source (relire la formule via l'API ajoute une branche TRUE() au SI)."""
    sh = k.sheet("Dashboard")
    for (lab, c), f in matrices.items():
        x = sh.getCellByPosition(c, find_row(sh, lab))
        try:
            x.setArrayFormula("")
        except Exception:
            pass
        x.setFormula("")
        x.setArrayFormula(conv(f))


def redline(k):
    for s, l, t, r, b in k.red:
        k.sheet(s).getCellRangeByPosition(l, t, r, b).CellBackColor = YELLOW


def main(src, out, out_red):
    matrices = matrices_origine(src)
    k = Calc(src)
    try:
        completer_donnees(k)
        inserer_champs(k)
        ajouter_moteur(k)
        completer_dashboard(k)
        restaurer_matrices(k, matrices)
        k.save(out)
        redline(k)
        k.save(out_red)
    finally:
        k.close()


if __name__ == "__main__":
    main(*sys.argv[1:4])
