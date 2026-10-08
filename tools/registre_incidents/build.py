# -*- coding: utf-8 -*-
"""Construit le classeur « Registre des incidents et des escalades » et son tableau de bord KPI.

Usage : python build.py [dossier_de_sortie]
Produit Registre_Incidents_KPI_Fundcraft.xlsx (vierge, noms réels des fonds) et
Registre_Incidents_KPI_Fundcraft_Demo.xlsx (données fictives de démonstration).
Charte Fundcraft : Calibri 9, quadrillage masqué, aucun volet figé, contenu en B2.
"""
import os
import sys
import datetime as dt

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.page import PageMargins
from dateutil.easter import easter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C  # noqa: E402

DARK, TEAL, CREAM, LIME, TXT, GREY, LINE, WHITE = "2F4858", "218E8E", "FFF8F2", "D9F785", "1A2027", "5A6B72", "E0E0E0", "FFFFFF"
FIRST, LAST = 7, 1006
DT, DATE = "dd/mm/yyyy hh:mm", "dd/mm/yyyy"
FMT = {"nb": "0", "%": "0.0%", "x": '0.0"x"', "j.o.": "0.0", "h": "0.0", "EUR": "#,##0"}
NBSP = " "
BOT = Border(bottom=Side(style="thin", color=LINE))


def fr(s):
    """Typographie française : espace insécable avant ; : ? ! % et à l'intérieur des guillemets."""
    if not isinstance(s, str) or s.startswith("="):
        return s
    for a in (" %", " :", " ;", " ?", " !", " »"):
        s = s.replace(a, NBSP + a[1:])
    return s.replace("« ", "«" + NBSP)


def F(b=False, c=TXT, s=9, i=False):
    return Font(name="Calibri", size=s, bold=b, italic=i, color=c)


def fill(c):
    return PatternFill("solid", fgColor=c)


def AL(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def put(ws, r, c, v=None, font=None, fl=None, fmt=None, al=None, border=None):
    col = CI(c) if isinstance(c, str) else c
    x = ws.cell(r, col)
    if v is not None:
        x.value = fr(v)
    x.font = font or F()
    if fl:
        x.fill = fill(fl)
    if fmt:
        x.number_format = fmt
    x.alignment = al or AL()
    if border:
        x.border = border
    return x


def band(ws, r, c1, c2, text):
    for c in range(CI(c1), CI(c2) + 1):
        ws.cell(r, c).fill = fill(DARK)
    put(ws, r, c1, text, F(True, WHITE), DARK)
    ws.row_dimensions[r].height = 15


def header(ws, r, c1, labels, h=None, wrap=True):
    for i, lab in enumerate(labels):
        put(ws, r, CI(c1) + i, lab, F(True, WHITE), TEAL, al=AL("left" if i == 0 else "center", wrap=wrap))
    if h:
        ws.row_dimensions[r].height = h


def title(ws, t, sub):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    put(ws, 2, "B", t, F(True, TXT, 10))
    put(ws, 3, "B", sub, F(False, GREY))


def page(ws, area, one_page=False):
    ws.print_area = area
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if one_page else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0, footer=0)


def row_height(texts_widths, base=11.5):
    lines = 1
    for t, w in texts_widths:
        if not t:
            continue
        cap = max(1.0, w * 1.1)
        n = sum(max(1, -(-len(p) // int(cap))) for p in str(t).split("\n"))
        lines = max(lines, n)
    return lines * base + 3


def jours_feries(years=range(2025, 2029)):
    out = []
    for y in years:
        e = easter(y)
        out += [(dt.date(y, 1, 1), "Jour de l'an"), (e + dt.timedelta(days=1), "Lundi de Pâques"),
                (dt.date(y, 5, 1), "Fête du travail"), (dt.date(y, 5, 8), "Victoire 1945"),
                (e + dt.timedelta(days=39), "Ascension"), (e + dt.timedelta(days=50), "Lundi de Pentecôte"),
                (dt.date(y, 7, 14), "Fête nationale"), (dt.date(y, 8, 15), "Assomption"),
                (dt.date(y, 11, 1), "Toussaint"), (dt.date(y, 11, 11), "Armistice 1918"), (dt.date(y, 12, 25), "Noël")]
    return sorted(out)


# --------------------------------------------------------------------------- formules
def BH(s, e):
    """Heures ouvrées entre deux dates-heures (calendrier et horaires de l'onglet Données)."""
    return (f'IF(OR({s}="",{e}=""),"",IF({e}<{s},"",ROUND(((NETWORKDAYS({s},{e},JoursFeries)-1)*(HFer-HOuv)'
            f'+IF(NETWORKDAYS({e},{e},JoursFeries),MEDIAN(MOD({e},1),HFer,HOuv),HFer)'
            f'-IF(NETWORKDAYS({s},{s},JoursFeries),MEDIAN(MOD({s},1),HFer,HOuv),HOuv))*24,4)))')


def helper_formulas(r):
    c = {k: f"{col}{r}" for k, col, *_ in C.CHAMPS}
    c.update({k: f"{col}{r}" for k, col, *_ in C.CALCULS})
    g = lambda k: c[k]  # noqa: E731
    ID = f"$B{r}"
    f = {}
    req = ",".join(f'{g(k)}=""' for k in ("Type", "Fonds", "Dom", "Src", "Desc", "Crit", "Statut", "Det", "Enr"))
    chrono = (f'AND({g("Surv")}<>"",{g("Surv")}>{g("Det")}),{g("Det")}>{g("Enr")},AND({g("PEC")}<>"",{g("PEC")}<{g("Enr")}),'
              f'AND({g("Res")}<>"",{g("Res")}<{g("Enr")}),AND({g("Clo")}<>"",{g("Clo")}<{g("Res")}),AND({g("Clo")}<>"",{g("Res")}=""),'
              f'AND({g("EscN2")}<>"",{g("EscN2")}<{g("Det")}),AND({g("EscN3")}<>"",{g("EscN3")}<{g("Det")})')
    f["Ctrl"] = (f'=IF({ID}="","",IF(OR({req}),"Champ obligatoire manquant",IF(OR({chrono}),"Chronologie incohérente",'
                 f'IF(AND(OR({g("Statut")}="Résolu",{g("Statut")}="Clos"),{g("Res")}=""),"Date de résolution manquante",'
                 f'IF(AND({g("Statut")}="Clos",{g("Clo")}=""),"Date de clôture manquante",'
                 f'IF(AND({g("DORA")}="Oui",{g("DClas")}=""),"Classification DORA manquante","OK"))))))')
    f["Cpt"] = f'=IF({ID}="","",IF(AND({g("Type")}="Incident",{g("Statut")}<>"Annulé"),1,0))'
    f["QI"] = f'=IF({ID}="","",IF(AND({g("Type")}="Quasi-incident",{g("Statut")}<>"Annulé"),1,0))'
    f["Actif"] = f'=IF({ID}="","",{g("Cpt")}+{g("QI")})'
    f["CritN"] = f'=IF(OR({ID}="",{g("Crit")}=""),"",IFERROR(VALUE(LEFT({g("Crit")},1)),""))'
    cpt = f'{g("Cpt")}<>1'
    f["Int"] = (f'=IF(OR({cpt},{g("Src")}=""),"",IFERROR(IF(INDEX(L_SrcInterne,MATCH({g("Src")},L_Sources,0))="Oui",1,0),""))')
    f["jDet"] = f'=IF({cpt},"",IFERROR(ROUND({BH(g("Surv"), g("Det"))}/HJour,4),""))'
    f["hDecl"] = f'=IF({cpt},"",{BH(g("Det"), g("Enr"))})'
    f["hPEC"] = f'=IF({cpt},"",{BH(g("Enr"), g("PEC"))})'
    f["jRes"] = f'=IF({cpt},"",IFERROR(ROUND({BH(g("Enr"), g("Res"))}/HJour,4),""))'
    f["jClo"] = f'=IF({cpt},"",IFERROR(ROUND({BH(g("Enr"), g("Clo"))}/HJour,4),""))'
    f["cPEC"] = f'=IF(OR({cpt},{g("CritN")}=""),"",INDEX(SLA_PEC,{g("CritN")}))'
    f["okPEC"] = f'=IF(OR({g("hPEC")}="",{g("cPEC")}=""),"",IF({g("hPEC")}<={g("cPEC")},1,0))'
    f["cRes"] = f'=IF(OR({cpt},{g("CritN")}=""),"",INDEX(SLA_RES,{g("CritN")}))'
    f["okRes"] = f'=IF(OR({g("jRes")}="",{g("cRes")}=""),"",IF({g("jRes")}<={g("cRes")},1,0))'
    f["DL"] = f'=IF(OR({g("cRes")}="",{g("Enr")}=""),"",WORKDAY(INT({g("Enr")}),{g("cRes")},JoursFeries)+MOD({g("Enr")},1))'
    f["Niv"] = (f'=IF({cpt},"",IF(OR({g("Not")}<>"",{g("DInit")}<>""),4,IF({g("EscN3")}<>"",3,IF({g("EscN2")}<>"",2,1))))')
    f["NivM"] = f'=IF(OR({cpt},{g("CritN")}=""),"",INDEX(NIV_REQ,{g("CritN")}))'
    f["NivR"] = f'=IF({g("NivM")}="","",MAX({g("NivM")},IF(OR({g("DORA")}="Oui",{g("NotReq")}="Oui"),4,1)))'
    f["EscOK"] = f'=IF(OR({g("NivR")}="",{g("Niv")}=""),"",IF({g("Niv")}>={g("NivR")},1,0))'
    f["DEsc"] = (f'=IF({g("NivM")}="","",IF({g("NivM")}<2,"",IF({g("NivM")}=2,'
                 f'IF(AND({g("EscN2")}="",{g("EscN3")}=""),"",MIN({g("EscN2")},{g("EscN3")})),'
                 f'IF({g("EscN3")}="","",{g("EscN3")}))))')
    f["hEsc"] = f'=IF({g("DEsc")}="","",{BH(g("Det"), g("DEsc"))})'
    f["EscDel"] = f'=IF(OR({g("hEsc")}="",{g("CritN")}=""),"",IF({g("hEsc")}<=INDEX(DELAI_ESC,{g("CritN")}),1,0))'
    f["D1Esc"] = f'=IF({cpt},"",IF(AND({g("EscN2")}="",{g("EscN3")}=""),"",MIN({g("EscN2")},{g("EscN3")})))'
    f["jEsc"] = f'=IF(OR({g("D1Esc")}="",{g("Res")}=""),"",IFERROR(ROUND({BH(g("D1Esc"), g("Res"))}/HJour,4),""))'
    f["ResN1"] = f'=IF(OR({cpt},{g("Res")}=""),"",IF({g("Niv")}=1,1,0))'
    f["ReoN"] = f'=IF({cpt},"",IF({g("Reo")}="Oui",1,0))'
    f["Rec"] = f'=IF({cpt},"",IF({g("Lien")}<>"",1,0))'
    f["RCAN"] = f'=IF({cpt},"",IF({g("RCA")}="Oui",1,0))'
    f["NotOK"] = (f'=IF(OR({cpt},{g("NotReq")}<>"Oui"),"",IF({g("Not")}="",IF(AND({g("NotLim")}<>"",NOW()>{g("NotLim")}),0,""),'
                  f'IF({g("NotLim")}="","",IF({g("Not")}<={g("NotLim")},1,0))))')
    f["DoraDL"] = (f'=IF(OR({cpt},{g("DORA")}<>"Oui",{g("Det")}=""),"",IF({g("DClas")}="",{g("Det")}+DORA_H_DET/24,'
                   f'IF({g("DClas")}-{g("Det")}<=DORA_H_DET/24,MIN({g("DClas")}+DORA_H_CLA/24,{g("Det")}+DORA_H_DET/24),'
                   f'{g("DClas")}+DORA_H_CLA/24)))')
    f["DoraI"] = (f'=IF({g("DoraDL")}="","",IF({g("DInit")}="",IF(NOW()>{g("DoraDL")},0,""),'
                  f'IF({g("DInit")}<={g("DoraDL")},1,0)))')
    f["DoraM"] = (f'=IF(OR({g("DoraDL")}="",{g("DInit")}=""),"",IF({g("DInter")}="",IF(NOW()>{g("DInit")}+DORA_H_INT/24,0,""),'
                  f'IF({g("DInter")}<={g("DInit")}+DORA_H_INT/24,1,0)))')
    f["DoraF"] = (f'=IF(OR({g("DoraDL")}="",{g("DInter")}=""),"",IF({g("DFin")}="",IF(NOW()>EDATE({g("DInter")},DORA_M_FIN)+1,0,""),'
                  f'IF({g("DFin")}<=EDATE({g("DInter")},DORA_M_FIN),1,0)))')
    f["DoraOK"] = (f'=IF({g("DoraDL")}="","",IF(COUNT({g("DoraI")}:{g("DoraF")})=0,"",MIN({g("DoraI")}:{g("DoraF")})))')
    f["PN"] = f'=IF({g("Actif")}<>1,"",IF(AND({g("PB")}="",{g("Recup")}=""),"",N({g("PB")})-N({g("Recup")})))'
    f["ActOK"] = (f'=IF(OR({g("Actif")}<>1,{g("ActReq")}<>"Oui",{g("Ech")}=""),"",IF({g("Real")}="",IF(NOW()>{g("Ech")}+1,0,""),'
                  f'IF({g("Real")}<={g("Ech")},1,0)))')
    return f


CALC_FMT = {"Ctrl": "@", "CritN": "0", "jDet": "0.0", "hDecl": "0.0", "hPEC": "0.0", "jRes": "0.0", "jClo": "0.0",
            "cPEC": "0", "cRes": "0", "DL": DT, "Niv": '"N"0', "NivM": '"N"0', "NivR": '"N"0', "DEsc": DT, "hEsc": "0.0",
            "D1Esc": DT, "jEsc": "0.0", "DoraDL": DT, "PN": "#,##0"}
INPUT_FMT = {"Surv": DT, "Det": DT, "Enr": DT, "PEC": DT, "Res": DT, "Clo": DT, "EscN2": DT, "EscN3": DT, "NotLim": DT,
             "Not": DT, "DClas": DT, "DInit": DT, "DInter": DT, "DFin": DATE, "Ech": DATE, "Real": DATE, "Reaf": "0",
             "PB": "#,##0", "Recup": "#,##0"}
INPUT_W = {"ID": 12, "Type": 12, "Fonds": 30, "Dom": 28, "Bale": 28, "Src": 28, "Presta": 18, "Desc": 36, "Crit": 11,
           "DORA": 10, "Statut": 10, "Resp": 15, "Reaf": 13, "Reo": 9, "Lien": 12, "RCA": 11, "NotReq": 12, "PB": 11,
           "Recup": 12, "ActReq": 12, "Com": 30}


def P(rng, D, Fn):
    return f'{rng},">="&{D},{rng},"<"&({Fn}+1)'


def kpi_formula(ref, D, Fn, v01, dom):
    """Renvoie (formule valeur, formule n ou None, matricielle)."""
    E = lambda r: P(r, D, Fn)  # noqa: E731
    ER, ERes, EClo, EEch = E("R_Enr"), E("R_Res"), E("R_Clo"), E("R_Ech")

    def ratio(num, den):
        return f'=IFERROR(({num})/({den}),"n.d.")', f"={den}"

    if ref == "V01":
        return f"=COUNTIFS(R_Cpt,1,{ER})", None, False
    if ref == "V02":
        return f'=IFERROR(COUNTIFS(R_QI,1,{ER})/{v01},"n.d.")', f"={v01}", False
    if ref == "V03":
        return f'=IFERROR(COUNTIFS(R_Cpt,1,R_CritN,"<=2",{ER})/{v01},"n.d.")', f"={v01}", False
    if ref == "V04":
        return f"=COUNTIFS(R_Cpt,1,R_CritN,1,{ER})", None, False
    if ref == "V05":
        return (*ratio(f"COUNTIFS(R_Int,1,{ER})", f'COUNTIFS(R_Int,">=0",{ER})'), False)
    if ref == "V06":
        return f'=IFERROR(MAX({dom})/{v01},"n.d.")', f"={v01}", False
    if ref == "V07":
        return f"=SUMIFS(R_PN,R_Cpt,1,{ER})", None, False
    if ref == "V08":
        return f'=COUNTIFS(R_Cpt,1,R_DORA,"Oui",{ER})', None, False
    if ref == "D01":
        return f'=IFERROR(AVERAGEIFS(R_jDet,R_Cpt,1,{ER}),"n.d.")', f'=COUNTIFS(R_jDet,">=0",{ER})', False
    if ref == "D02":
        return f'=IFERROR(AVERAGEIFS(R_hDecl,R_Cpt,1,{ER}),"n.d.")', f'=COUNTIFS(R_hDecl,">=0",{ER})', False
    if ref == "D03":
        return f'=IFERROR(AVERAGEIFS(R_hPEC,R_Cpt,1,{ER}),"n.d.")', f'=COUNTIFS(R_hPEC,">=0",{ER})', False
    if ref == "D04":
        return (*ratio(f"COUNTIFS(R_okPEC,1,{ER})", f'COUNTIFS(R_okPEC,">=0",{ER})'), False)
    if ref == "D05":
        return f'=IFERROR(AVERAGEIFS(R_jRes,{ERes}),"n.d.")', f'=COUNTIFS(R_jRes,">=0",{ERes})', False
    if ref == "D06":
        return (f'=IFERROR(PERCENTILE(IF(ISNUMBER(R_jRes)*(R_Res>={D})*(R_Res<{Fn}+1),R_jRes),0.9),"n.d.")',
                f'=COUNTIFS(R_jRes,">=0",{ERes})', True)
    if ref == "D07":
        return (*ratio(f"COUNTIFS(R_okRes,1,{ERes})", f'COUNTIFS(R_okRes,">=0",{ERes})'), False)
    if ref == "D08":
        return f'=IFERROR(AVERAGEIFS(R_jClo,{EClo}),"n.d.")', f'=COUNTIFS(R_jClo,">=0",{EClo})', False
    if ref == "E01":
        return (*ratio(f'COUNTIFS(R_ResN1,1,R_CritN,">=3",{ERes})', f'COUNTIFS(R_ResN1,">=0",R_CritN,">=3",{ERes})'), False)
    if ref == "E02":
        return f'=IFERROR(COUNTIFS(R_Cpt,1,{ERes})/{v01},"n.d.")', f"={v01}", False
    if ref == "E03":
        return (f'=COUNTIFS(R_Cpt,1,R_Enr,"<"&({Fn}+1))-COUNTIFS(R_Cpt,1,R_Enr,"<"&({Fn}+1),R_Res,"<"&({Fn}+1))', None, False)
    if ref == "E04":
        bl = f'(COUNTIFS(R_Cpt,1,R_Enr,"<"&({Fn}+1))-COUNTIFS(R_Cpt,1,R_Enr,"<"&({Fn}+1),R_Res,"<"&({Fn}+1)))'
        late = (f'(COUNTIFS(R_Cpt,1,R_Enr,"<"&({Fn}+1),R_DL,"<"&({Fn}+1))'
                f'-COUNTIFS(R_Cpt,1,R_Enr,"<"&({Fn}+1),R_DL,"<"&({Fn}+1),R_Res,"<"&({Fn}+1)))')
        return (*ratio(late, bl), False)
    if ref == "E05":
        return (*ratio(f"COUNTIFS(R_ReoN,1,{ERes})", f"COUNTIFS(R_Cpt,1,{ERes})"), False)
    if ref == "E06":
        return f'=IFERROR(COUNTIFS(R_Rec,1,{ER})/{v01},"n.d.")', f"={v01}", False
    if ref == "E07":
        return (*ratio(f'COUNTIFS(R_RCAN,1,R_CritN,"<=2",{EClo})', f'COUNTIFS(R_Cpt,1,R_CritN,"<=2",{EClo})'), False)
    if ref == "E08":
        return (*ratio(f"COUNTIFS(R_ActOK,1,{EEch})", f'COUNTIFS(R_ActOK,">=0",{EEch})'), False)
    if ref == "X01":
        return f'=IFERROR(COUNTIFS(R_Niv,">=2",{ER})/{v01},"n.d.")', f"={v01}", False
    if ref == "X02":
        return f'=IFERROR(COUNTIFS(R_Niv,">=3",{ER})/{v01},"n.d.")', f"={v01}", False
    if ref == "X03":
        return f'=IFERROR(COUNTIFS(R_Niv,4,{ER})/{v01},"n.d.")', f"={v01}", False
    if ref == "X04":
        return (*ratio(f"COUNTIFS(R_EscOK,1,{ER})", f'COUNTIFS(R_EscOK,">=0",{ER})'), False)
    if ref == "X05":
        return (*ratio(f"COUNTIFS(R_EscDel,1,{ER})", f'COUNTIFS(R_EscDel,">=0",{ER})'), False)
    if ref == "X06":
        return f'=IFERROR(AVERAGEIFS(R_hEsc,R_Cpt,1,{ER}),"n.d.")', f'=COUNTIFS(R_hEsc,">=0",{ER})', False
    if ref == "X07":
        return f'=IFERROR(AVERAGEIFS(R_jEsc,{ERes}),"n.d.")', f'=COUNTIFS(R_jEsc,">=0",{ERes})', False
    if ref == "X08":
        return (*ratio(f'COUNTIFS(R_Cpt,1,R_Reaf,">=2",{ER})', f'COUNTIFS(R_Cpt,1,R_Reaf,">=0",{ER})'), False)
    if ref == "X09":
        return (*ratio(f"COUNTIFS(R_NotOK,1,{ER})", f'COUNTIFS(R_NotOK,">=0",{ER})'), False)
    if ref == "X10":
        return (*ratio(f"COUNTIFS(R_DoraOK,1,{ER})", f'COUNTIFS(R_DoraOK,">=0",{ER})'), False)
    raise KeyError(ref)


# --------------------------------------------------------------------------- classeur
class Book:
    def __init__(self):
        self.wb = openpyxl.Workbook()
        self.wb._fonts = IndexedList([Font(name="Calibri", sz=9, family=2, scheme="minor")])
        self.wb.remove(self.wb.active)

    def name(self, nm, ws, ref):
        self.wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{ws.title}'!{ref}")


def build_donnees(bk, ws, fonds, demo):
    title(ws, "Données et paramètres",
          "Cellules crème : paramètres modifiables. Tous les calculs du classeur lisent ces cellules par plages nommées. "
          "Les cibles par défaut sont des standards de marché ou des hypothèses à valider en Comité des risques.")
    for col, w in zip("BCDEFGHI", (7, 46, 13, 13, 13, 13, 13, 70)):
        ws.column_dimensions[col].width = w
    r = 5
    band(ws, r, "B", "I", "1. Période d'analyse et calendrier ouvré")
    header(ws, r + 1, "B", ["Code", "Paramètre", "Valeur", "Unité", "", "", "", "Source / statut"])
    params = [
        ("P01", "Début de période", dt.date(2026, 7, 1), DATE, "date", "Saisie - premier jour de la période analysée", "DebutPeriode", True),
        ("P02", "Durée de la période", 3, "0", "mois", "Saisie - 1 = mois, 3 = trimestre, 12 = année", "DureeMois", True),
        ("P03", "Fin de période", "=EDATE(D7,D8)-1", DATE, "date", "Calcul", "FinPeriode", False),
        ("P04", "Base minimale pour juger un ratio", 5, "0", "incidents",
         "Hypothèse à confirmer - en dessous, le statut affiché est « Base insuffisante »", "NMin", True),
        ("P05", "Heure d'ouverture", dt.time(9, 0), "hh:mm", "heure", "Hypothèse - horaires de référence de la SGP", "HOuv", True),
        ("P06", "Heure de fermeture", dt.time(18, 0), "hh:mm", "heure", "Hypothèse - horaires de référence de la SGP", "HFer", True),
        ("P07", "Heures ouvrées par jour", "=(D12-D11)*24", "0.0", "h", "Calcul - 1 j.o. = ce nombre d'heures ouvrées", "HJour", False),
    ]
    r += 2
    for code, lab, val, fmt, unit, src, nm, inp in params:
        put(ws, r, "B", code, border=BOT)
        put(ws, r, "C", lab, border=BOT)
        put(ws, r, "D", val, F(not inp), CREAM if inp else None, fmt, AL("right"), BOT)
        put(ws, r, "E", unit, F(c=GREY), border=BOT)
        for c in "FGH":
            put(ws, r, c, None, border=BOT)
        put(ws, r, "I", src, F(c=GREY), border=BOT)
        bk.name(nm, ws, f"$D${r}")
        r += 1
    assert r == 14

    r += 1
    band(ws, r, "B", "I", "2. Matrice de criticité, SLA et escalade")
    header(ws, r + 1, "B", ["Code", "Criticité", "Prise en compte (h ouvrées)", "Résolution (j.o.)", "Niveau d'escalade requis",
                           "Délai max d'escalade (h ouvrées)", "", "Critères indicatifs"], h=40)
    r += 2
    r0 = r
    for i, (lab, pec, res, niv, dl, crit) in enumerate(C.CRITICITES, 1):
        put(ws, r, "B", f"C{i}", border=BOT)
        put(ws, r, "C", lab, border=BOT)
        put(ws, r, "D", pec, None, CREAM, "0", AL("right"), BOT)
        put(ws, r, "E", res, None, CREAM, "0", AL("right"), BOT)
        put(ws, r, "F", niv, None, CREAM, '"N"0', AL("right"), BOT)
        put(ws, r, "G", dl, None, CREAM, "0", AL("right"), BOT)
        put(ws, r, "H", None, border=BOT)
        put(ws, r, "I", crit, F(c=GREY), al=AL(wrap=True), border=BOT)
        ws.row_dimensions[r].height = row_height([(crit, 70)])
        r += 1
    for nm, col in (("L_Crit", "C"), ("SLA_PEC", "D"), ("SLA_RES", "E"), ("NIV_REQ", "F"), ("DELAI_ESC", "G")):
        bk.name(nm, ws, f"${col}${r0}:${col}${r - 1}")
    note = ("Source : hypothèses de calibrage pour une SGP de fonds non cotés, à valider en Comité des risques. Références ITSM usuelles "
            "(service desk informatique) : prise en compte de 15 min à 8 h et résolution de 4 h à 5 jours selon la priorité, non "
            "transposables telles quelles aux incidents de valorisation, de passif ou d'appels de fonds. Les seuils financiers de "
            "criticité (perte, impact VL) sont à fixer par la Direction.")
    put(ws, r, "C", note, F(c=GREY), al=AL(wrap=True, v="top"))
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9)
    ws.row_dimensions[r].height = row_height([(note, 190)])
    r += 2

    band(ws, r, "B", "I", "3. Niveaux d'escalade")
    header(ws, r + 1, "B", ["Code", "Niveau", "Valeur", "", "", "", "", "Définition"])
    r += 2
    r0 = r
    for code, lab, num, desc in C.NIVEAUX:
        put(ws, r, "B", code, border=BOT)
        put(ws, r, "C", lab, border=BOT)
        put(ws, r, "D", num, F(True), None, '"N"0', AL("right"), BOT)
        for c in "EFGH":
            put(ws, r, c, None, border=BOT)
        put(ws, r, "I", desc, F(c=GREY), al=AL(wrap=True), border=BOT)
        ws.row_dimensions[r].height = row_height([(desc, 70)])
        r += 1
    bk.name("L_NivLib", ws, f"$C${r0}:$C${r - 1}")
    bk.name("L_NivNum", ws, f"$D${r0}:$D${r - 1}")
    r += 1

    band(ws, r, "B", "I", "4. Paramètres DORA (incident majeur lié aux TIC)")
    header(ws, r + 1, "B", ["Code", "Paramètre", "Valeur", "Unité", "", "", "", "Source / statut"])
    r += 2
    dora = [("DORA1", "Notification initiale - délai après classification", 4, "h", "DORA_H_CLA"),
            ("DORA2", "Notification initiale - délai maximal après détection", 24, "h", "DORA_H_DET"),
            ("DORA3", "Rapport intermédiaire - délai après notification initiale", 72, "h", "DORA_H_INT"),
            ("DORA4", "Rapport final - délai après le dernier rapport intermédiaire", 1, "mois", "DORA_M_FIN")]
    src = ("Règlement (UE) 2022/2554 (DORA), art. 19 ; Règlement délégué (UE) 2025/301 (contenu et délais des notifications) ; "
           "classification : Règlement délégué (UE) 2024/1772. Heures calendaires.")
    for code, lab, val, unit, nm in dora:
        put(ws, r, "B", code, border=BOT)
        put(ws, r, "C", lab, border=BOT)
        put(ws, r, "D", val, None, CREAM, "0", AL("right"), BOT)
        put(ws, r, "E", unit, F(c=GREY), border=BOT)
        for c in "FGH":
            put(ws, r, c, None, border=BOT)
        put(ws, r, "I", src if code == "DORA1" else "Idem", F(c=GREY), al=AL(wrap=True), border=BOT)
        if code == "DORA1":
            ws.row_dimensions[r].height = row_height([(src, 70)])
        bk.name(nm, ws, f"$D${r}")
        r += 1
    note = ("Si la classification intervient plus de 24 h après la détection, le délai de 4 h court depuis la classification. "
            "Les assouplissements éventuels pour les échéances tombant un week-end ou un jour férié ne sont pas modélisés : à vérifier "
            "dans le texte du règlement délégué avant usage.")
    put(ws, r, "C", note, F(c=GREY), al=AL(wrap=True, v="top"))
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9)
    ws.row_dimensions[r].height = row_height([(note, 190)])
    r += 2

    band(ws, r, "B", "I", "5. Cibles des indicateurs")
    header(ws, r + 1, "B", ["Réf", "Indicateur", "Unité", "Sens", "Cible", "Seuil d'alerte", "Base minimale", "Source / statut"], h=26)
    r += 2
    r0 = r
    kpi_rows = {}
    for k in C.KPIS:
        put(ws, r, "B", k["ref"], border=BOT)
        put(ws, r, "C", k["lib"], border=BOT)
        put(ws, r, "D", k["unit"], F(c=GREY), border=BOT, al=AL("center"))
        put(ws, r, "E", k["sens"], None, CREAM, None, AL("center"), BOT)
        put(ws, r, "F", k["cible"], None, CREAM, FMT[k["unit"]], AL("right"), BOT)
        put(ws, r, "G", k["alerte"], None, CREAM, FMT[k["unit"]], AL("right"), BOT)
        put(ws, r, "H", k["base"], None, CREAM, None, AL("center"), BOT)
        put(ws, r, "I", k["src"], F(c=GREY), border=BOT)
        kpi_rows[k["ref"]] = r
        r += 1
    for nm, col in (("T_Ref", "B"), ("T_Lib", "C"), ("T_Unite", "D"), ("T_Sens", "E"), ("T_Cible", "F"),
                    ("T_Alerte", "G"), ("T_Base", "H")):
        bk.name(nm, ws, f"${col}${r0}:${col}${r - 1}")
    note = ("Sens : « ≥ » la valeur doit être supérieure ou égale à la cible, « ≤ » inférieure ou égale, « Suivi » sans cible. "
            "Entre la cible et le seuil d'alerte : Vigilance ; au-delà : Hors cible. Base minimale = Oui : le ratio n'est jugé que si "
            "son dénominateur atteint la base minimale (P04) ; Non pour les indicateurs absolus et de conformité.")
    put(ws, r, "C", note, F(c=GREY), al=AL(wrap=True, v="top"))
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9)
    ws.row_dimensions[r].height = row_height([(note, 190)])
    r += 2

    band(ws, r, "B", "I", "6. Listes de saisie du registre")
    r += 1
    lists = [("Type", C.TYPES, "L_Type", None), ("Statut", C.STATUTS, "L_Statuts", None), ("Oui / Non", C.OUINON, "L_OuiNon", None),
             ("Fonds / véhicule", fonds, "L_Fonds",
              "Données fictives de démonstration" if demo else "Liste à tenir à jour (onboarding, nouveaux véhicules)"),
             ("Domaine", C.DOMAINES, "L_Domaines", None), ("Catégorie d'événement (Bâle II)", C.BALE, "L_Bale",
                                                           "Comité de Bâle, classification des événements de perte (2006)"),
             ("Source de détection", [s for s, _ in C.SOURCES], "L_Sources",
              "Interne = détection par le dispositif de la SGP ou de ses délégataires (convention à valider)")]
    for lab, items, nm, src in lists:
        header(ws, r, "B", ["Liste", lab, "Interne" if nm == "L_Sources" else "", "", "", "", "", "Source / statut"])
        r += 1
        r0 = r
        for i, it in enumerate(items, 1):
            put(ws, r, "B", i, F(c=GREY), border=BOT, al=AL("left"))
            put(ws, r, "C", it, None, CREAM, border=BOT)
            if nm == "L_Sources":
                put(ws, r, "D", C.SOURCES[i - 1][1], None, CREAM, None, AL("center"), BOT)
            else:
                put(ws, r, "D", None, border=BOT)
            for c in "EFGH":
                put(ws, r, c, None, border=BOT)
            put(ws, r, "I", src if (src and i == 1) else None, F(c=GREY), border=BOT)
            r += 1
        bk.name(nm, ws, f"$C${r0}:$C${r - 1}")
        if nm == "L_Sources":
            bk.name("L_SrcInterne", ws, f"$D${r0}:$D${r - 1}")
        r += 1

    band(ws, r, "B", "I", "7. Jours fériés (France)")
    header(ws, r + 1, "B", ["Code", "Jour férié", "Date", "", "", "", "", "Source / statut"])
    r += 2
    r0 = r
    for i, (d, lab) in enumerate(jours_feries(), 1):
        put(ws, r, "B", f"JF{i:02d}", F(c=GREY), border=BOT)
        put(ws, r, "C", lab, border=BOT)
        put(ws, r, "D", d, None, CREAM, DATE, AL("right"), BOT)
        for c in "EFGH":
            put(ws, r, c, None, border=BOT)
        put(ws, r, "I", "Code du travail, art. L3133-1 ; Alsace-Moselle non intégrée ; à prolonger au-delà de 2028" if i == 1 else None,
            F(c=GREY), border=BOT)
        r += 1
    bk.name("JoursFeries", ws, f"$D${r0}:$D${r - 1}")
    page(ws, f"B2:I{r}")
    return kpi_rows


def build_registre(bk, ws, rows, demo):
    sub = ("Une ligne par incident ou quasi-incident. Saisie dans les cellules crème (listes déroulantes alimentées par l'onglet Données), "
           "dates au format jj/mm/aaaa hh:mm. Colonnes AN à BX calculées : ne pas saisir. Définitions et exemples : onglet Champs.")
    if demo:
        sub = "Données fictives de démonstration, à supprimer avant utilisation. " + sub
    title(ws, "Registre des incidents et des escalades", sub)
    blocks = [("B", "M", "Identification et qualification"), ("N", "S", "Chronologie"), ("T", "V", "Escalade"),
              ("W", "Y", "Qualité de résolution"), ("Z", "AB", "Notifications externes"), ("AC", "AF", "DORA"),
              ("AG", "AH", "Impact financier"), ("AI", "AL", "Actions correctives"), ("AN", "BX", "Calculs automatiques - ne pas saisir")]
    for a, b, t in blocks:
        band(ws, 5, a, b, t)
    for k, col, bloc, lab, *_ in C.CHAMPS:
        put(ws, 6, col, lab, F(True, WHITE), TEAL, al=AL("left", "center", True))
        ws.column_dimensions[col].width = INPUT_W.get(k, 16 if INPUT_FMT.get(k) == DT else 11)
    for k, col, lab, _ in C.CALCULS:
        put(ws, 6, col, lab, F(True, WHITE), TEAL, al=AL("left", "center", True))
        ws.column_dimensions[col].width = 22 if k == "Ctrl" else (16 if CALC_FMT.get(k) == DT else 11)
    ws.column_dimensions["AM"].width = 2
    ws.row_dimensions[6].height = 48
    ws.column_dimensions.group("AN", "BX", outline_level=1, hidden=False)

    keys = [k for k, *_ in C.CHAMPS]
    for r in range(FIRST, LAST + 1):
        data = rows[r - FIRST] if r - FIRST < len(rows) else {}
        for k, col, *_ in C.CHAMPS:
            v = data.get(k)
            x = ws[f"{col}{r}"]
            if v is not None:
                x.value = v
            x.fill = fill(CREAM)
            x.border = BOT
            x.font = F()
            if k in INPUT_FMT:
                x.number_format = INPUT_FMT[k]
        for k, f in helper_formulas(r).items():
            col = dict((kk, cc) for kk, cc, *_ in C.CALCULS)[k]
            x = ws[f"{col}{r}"]
            x.value = f
            x.font = F()
            x.border = BOT
            if k in CALC_FMT:
                x.number_format = CALC_FMT[k]
    assert len(keys) == len(set(keys))

    for k, col, *_ in C.CHAMPS + [(k, c) for k, c, *_ in C.CALCULS]:
        bk.name(f"R_{k}", ws, f"${col}${FIRST}:${col}${LAST}")

    def dv(lst, cols):
        d = DataValidation(type="list", formula1=lst, allow_blank=True, showErrorMessage=True,
                           errorTitle="Valeur hors liste", error="Choisir une valeur de la liste (onglet Données).")
        ws.add_data_validation(d)
        for c in cols:
            d.add(f"{c}{FIRST}:{c}{LAST}")
    dv("L_Type", ["C"])
    dv("L_Fonds", ["D"])
    dv("L_Domaines", ["E"])
    dv("L_Bale", ["F"])
    dv("L_Sources", ["G"])
    dv("L_Crit", ["J"])
    dv("L_Statuts", ["L"])
    dv("L_OuiNon", ["K", "W", "Y", "Z", "AI"])
    d = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True,
                       error="Nombre entier positif ou nul.")
    ws.add_data_validation(d)
    d.add(f"V{FIRST}:V{LAST}")
    d = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True,
                       error="Montant positif ou nul, en EUR.")
    ws.add_data_validation(d)
    d.add(f"AG{FIRST}:AH{LAST}")
    page(ws, f"B2:AL{LAST}")


def build_indicateurs(ws, kpi_rows):
    title(ws, "Dictionnaire des indicateurs",
          "Indicateurs strictement quantitatifs, regroupés en quatre catégories. Libellé, unité, sens, cible et seuil d'alerte sont lus "
          "dans l'onglet Données ; R_xxx désigne la colonne du registre de même nom (plage nommée), Début et Fin les bornes de la période.")
    widths = dict(zip("BCDEFGHIJKLMN", (6, 28, 7, 44, 40, 44, 6, 8, 8, 40, 40, 14, 40)))
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    header(ws, 5, "B", ["Réf", "Indicateur", "Unité", "Formule de calcul", "Méthode Excel", "Objectif métier", "Sens", "Cible",
                        "Seuil d'alerte", "Origine de la cible", "Champs requis du registre", "Fréquence", "Points de vigilance"], h=26)
    labels = {k: (col, lab) for k, col, _, lab, *_ in C.CHAMPS}
    r = 6
    rows = {}
    for cat, cat_lab in C.CATEGORIES:
        band(ws, r, "B", "N", cat_lab)
        r += 1
        for k in [k for k in C.KPIS if k["ref"].startswith(cat)]:
            m = f'MATCH($B{r},T_Ref,0)'
            champs = " ; ".join(f"{labels[f][1]} ({labels[f][0]})" for f in k["champs"])
            put(ws, r, "B", k["ref"], F(True), border=BOT, al=AL("left", "top"))
            put(ws, r, "C", f"=INDEX(T_Lib,{m})", F(True), border=BOT, al=AL("left", "top", True))
            put(ws, r, "D", f"=INDEX(T_Unite,{m})", F(c=GREY), border=BOT, al=AL("center", "top"))
            for col, key in (("E", "formule"), ("F", "excel"), ("G", "objectif"), ("K", "origine"), ("N", "vigilance")):
                x = put(ws, r, col, None, border=BOT, al=AL("left", "top", True))
                x.value = fr(k[key]) if key != "excel" else k[key]
                x.data_type = "s"
            put(ws, r, "H", f"=INDEX(T_Sens,{m})", border=BOT, al=AL("center", "top"))
            fm = FMT[k["unit"]]
            put(ws, r, "I", f'=IF(INDEX(T_Cible,{m})="","",INDEX(T_Cible,{m}))', border=BOT, fmt=fm, al=AL("right", "top"))
            put(ws, r, "J", f'=IF(INDEX(T_Alerte,{m})="","",INDEX(T_Alerte,{m}))', border=BOT, fmt=fm, al=AL("right", "top"))
            put(ws, r, "L", champs, border=BOT, al=AL("left", "top", True))
            put(ws, r, "M", k["freq"], border=BOT, al=AL("left", "top", True))
            ws.row_dimensions[r].height = row_height([(k["lib"], 28), (k["formule"], 44), (k["excel"], 40), (k["objectif"], 44),
                                                      (k["origine"], 40), (champs, 40), (k["vigilance"], 40), (k["freq"], 14)])
            rows[k["ref"]] = r
            r += 1
    r += 1
    note = ("Références réglementaires : Règlement délégué (UE) n° 231/2013 (niveau 2 AIFMD), art. 13 - risque opérationnel et base "
            "historique des pertes, art. 12 à 15 - fonds propres supplémentaires et assurance RC professionnelle ; Règlement (UE) "
            "2022/2554 (DORA), art. 17 à 19 ; Règlement délégué (UE) 2024/1772 (classification) ; Règlement délégué (UE) 2025/301 "
            "(délais de notification). Les standards ITSM cités (FCR, taux de réouverture, respect des SLA) sont des ordres de grandeur "
            "de service desk informatique (pratique de marché, non réglementaire), transposés au contexte d'une SGP de fonds non cotés.")
    put(ws, r, "B", note, F(c=GREY), al=AL(wrap=True, v="top"))
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
    ws.row_dimensions[r].height = row_height([(note, 320)])
    page(ws, f"B2:N{r}")
    return rows


def build_champs(ws, fonds):
    title(ws, "Champs du registre",
          "Modèle de données du registre : définition, règle de saisie et exemple de format pour chaque colonne, puis champs calculés.")
    for c, w in zip("BCDEFGHI", (6, 14, 34, 16, 16, 60, 34, 40)):
        ws.column_dimensions[c].width = w
    header(ws, 5, "B", ["Col.", "Bloc", "Champ", "Saisie", "Format", "Définition et règle de saisie", "Exemple", "Indicateurs alimentés"],
           h=15)
    r = 6
    band(ws, r, "B", "I", "Champs saisis")
    r += 1
    for k, col, bloc, lab, obl, typ, desc, ex in C.CHAMPS:
        used = ", ".join(x["ref"] for x in C.KPIS if k in x["champs"])
        extra = {"Fonds": "répartition par fonds", "Dom": "répartition par domaine", "Bale": "répartition Bâle II",
                 "Crit": "répartition par criticité"}.get(k)
        if extra:
            used = (used + " ; " if used else "") + "Dashboard : " + extra
        ex = fonds[1] if k == "Fonds" else ex
        vals = [col, bloc, lab, obl, typ, desc, ex, used or "-"]
        for i, v in enumerate(vals):
            x = put(ws, r, 2 + i, None, F(True) if i == 2 else F(), CREAM if i == 6 else None, border=BOT, al=AL("left", "top", True))
            x.value = fr(v)
            x.data_type = "s"
        ws.row_dimensions[r].height = row_height([(lab, 34), (desc, 60), (ex, 34), (used, 40)])
        r += 1
    r += 1
    band(ws, r, "B", "I", "Champs calculés (colonnes groupées AN à BX)")
    r += 1
    for k, col, lab, desc in C.CALCULS:
        used = ", ".join(x["ref"] for x in C.KPIS if f"R_{k}" in x["excel"]) or ("Dashboard" if k in ("Ctrl", "PN", "Niv") else "Calcul intermédiaire")
        if k in ("Ctrl",):
            used = "Dashboard (contrôle de saisie)"
        vals = [col, "Calcul", lab, "Calculé", CALC_FMT.get(k, "0 / 1").replace('"N"0', "N1 à N4"), desc, "", used]
        for i, v in enumerate(vals):
            x = put(ws, r, 2 + i, None, F(True) if i == 2 else F(), border=BOT, al=AL("left", "top", True))
            x.value = fr(v) if v else None
            x.data_type = "s"
        ws.row_dimensions[r].height = row_height([(lab, 34), (desc, 60), (used, 40)])
        r += 1
    page(ws, f"B2:I{r}")


def build_dashboard(ws, ind_rows, demo):
    sub = ("Fundcraft - indicateurs calculés à partir de l'onglet Registre ; paramètres, matrice SLA et cibles dans l'onglet Données ; "
           "définitions dans l'onglet Indicateurs (lien sur chaque référence).")
    if demo:
        sub = "Données fictives de démonstration. " + sub
    title(ws, "Tableau de bord du registre des incidents et des escalades", sub)
    widths = dict(zip(["B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V"],
                      (5, 42, 6, 11, 5, 15, 11, 11, 5, 15, 5, 8, 8, 2, 48, 10, 10, 10, 10, 14, 12)))
    for c, w in widths.items():
        ws.column_dimensions[c].width = w

    band(ws, 5, "B", "N", "Période d'analyse et qualité des données")
    header(ws, 6, "B", ["", "Bornes", "", "Période", "", "", "Période précédente", "12 mois glissants", "", "", "", "", ""], h=26)
    put(ws, 7, "C", "Début", border=BOT)
    put(ws, 8, "C", "Fin", border=BOT)
    put(ws, 7, "E", "=DebutPeriode", F(True), LIME, DATE, AL("right"), BOT)
    put(ws, 8, "E", "=FinPeriode", F(True), LIME, DATE, AL("right"), BOT)
    put(ws, 7, "H", "=EDATE(E7,-DureeMois)", None, None, DATE, AL("right"), BOT)
    put(ws, 8, "H", "=E7-1", None, None, DATE, AL("right"), BOT)
    put(ws, 7, "I", "=EDATE(E8+1,-12)", None, None, DATE, AL("right"), BOT)
    put(ws, 8, "I", "=E8", None, None, DATE, AL("right"), BOT)
    put(ws, 9, "C", "Lignes saisies dans le registre", border=BOT)
    put(ws, 9, "E", "=COUNTA(R_ID)", F(True), LIME, "0", AL("right"), BOT)
    put(ws, 10, "C", "Lignes en anomalie de saisie (voir le registre)", border=BOT)
    put(ws, 10, "E", '=SUMPRODUCT((R_Ctrl<>"")*(R_Ctrl<>"OK"))', F(True), LIME, "0", AL("right"), BOT)
    for rr in (7, 8, 9, 10):
        for c in ("B", "D", "F", "G", "J", "K", "L", "M", "N"):
            put(ws, rr, c, None, border=BOT)
        if rr >= 9:
            for c in ("H", "I"):
                put(ws, rr, c, None, border=BOT)

    hr = 12
    header(ws, hr, "B", ["Réf", "Indicateur", "Unité", "Période", "n", "Statut", "Période précédente", "12 mois glissants", "n",
                         "Statut 12 mois", "Sens", "Cible", "Seuil d'alerte"], h=26)
    cols = {"E": ("E$7", "E$8", "$Q$15:$Q$27"), "H": ("H$7", "H$8", "$R$15:$R$27"), "I": ("I$7", "I$8", "$S$15:$S$27")}
    r = hr + 1
    v01_row = None
    for cat, cat_lab in C.CATEGORIES:
        band(ws, r, "B", "N", cat_lab)
        r += 1
        for k in [k for k in C.KPIS if k["ref"].startswith(cat)]:
            ref = k["ref"]
            if ref == "V01":
                v01_row = r
            m = f"MATCH($B{r},T_Ref,0)"
            x = put(ws, r, "B", ref, F(True, TEAL), border=BOT)
            x.hyperlink = Hyperlink(ref=f"B{r}", location=f"'Indicateurs'!B{ind_rows[ref]}", display=ref)
            put(ws, r, "C", f"=INDEX(T_Lib,{m})", border=BOT)
            put(ws, r, "D", f"=INDEX(T_Unite,{m})", F(c=GREY), border=BOT, al=AL("center"))
            fm = FMT[k["unit"]]
            for col, (D, Fn, dom) in cols.items():
                v, n, arr = kpi_formula(ref, D, Fn, f"{col}${v01_row}", dom)
                cell = ws[f"{col}{r}"]
                cell.value = ArrayFormula(f"{col}{r}", v) if arr else v
                key = col == "E"
                cell.font = F(key)
                if key:
                    cell.fill = fill(LIME)
                cell.number_format = fm
                cell.alignment = AL("right")
                cell.border = BOT
                ncol = {"E": "F", "I": "J"}.get(col)
                if ncol:
                    put(ws, r, ncol, n, F(c=GREY), None, "0", AL("right"), BOT)
            put(ws, r, "L", f"=INDEX(T_Sens,{m})", border=BOT, al=AL("center"))
            put(ws, r, "M", f'=IF(INDEX(T_Cible,{m})="","",INDEX(T_Cible,{m}))', border=BOT, fmt=fm, al=AL("right"))
            put(ws, r, "N", f'=IF(INDEX(T_Alerte,{m})="","",INDEX(T_Alerte,{m}))', border=BOT, fmt=fm, al=AL("right"))
            for vcol, ncol, scol in (("E", "F", "G"), ("I", "J", "K")):
                V, N = f"{vcol}{r}", f"{ncol}{r}"
                core = (f'IF($L{r}="≥",IF({V}>=$M{r},"Conforme",IF({V}>=$N{r},"Vigilance","Hors cible")),'
                        f'IF({V}<=$M{r},"Conforme",IF({V}<=$N{r},"Vigilance","Hors cible")))')
                st = (f'=IF(NOT(ISNUMBER({V})),"n.d.",IF($L{r}="Suivi","Suivi",IF($M{r}="","Cible à définir",'
                      f'IF(AND(INDEX(T_Base,{m})="Oui",ISNUMBER({N}),N({N})<NMin),"Base insuffisante",{core}))))')
                put(ws, r, scol, st, border=BOT, al=AL("left"))
            r += 1
    last_kpi = r
    notes = ["Statuts : Conforme ; Vigilance (entre la cible et le seuil d'alerte) ; Hors cible ; Base insuffisante (dénominateur n "
             "inférieur à la base minimale, ratio non significatif) ; Suivi (indicateur sans cible) ; Cible à définir ; n.d. (aucune donnée).",
             "n : nombre d'incidents au dénominateur du ratio. Délais en heures ouvrées (h) et jours ouvrés (j.o.) selon le calendrier de "
             "l'onglet Données ; délais DORA en heures calendaires. Période précédente : même durée, immédiatement avant la période."]

    # Répartitions
    PE = lambda rng: P(rng, "$E$7", "$E$8")  # noqa: E731
    PP = lambda rng: P(rng, "$H$7", "$H$8")  # noqa: E731
    PY = lambda rng: P(rng, "$I$7", "$I$8")  # noqa: E731
    band(ws, 5, "P", "V", "Répartition par criticité")
    header(ws, 6, "P", ["Criticité", "Période", "%", "12 mois", "%", "MTTR 12 mois (j.o.)", "SLA résol. 12 mois"], h=26)
    for i in range(4):
        rr = 7 + i
        put(ws, rr, "P", f"=INDEX(L_Crit,{i + 1})", border=BOT)
        put(ws, rr, "Q", f"=COUNTIFS(R_Cpt,1,R_Crit,$P{rr},{PE('R_Enr')})", F(True), LIME, "0", AL("right"), BOT)
        put(ws, rr, "R", f'=IFERROR(Q{rr}/Q$11,"n.d.")', None, None, "0.0%", AL("right"), BOT)
        put(ws, rr, "S", f"=COUNTIFS(R_Cpt,1,R_Crit,$P{rr},{PY('R_Enr')})", None, None, "0", AL("right"), BOT)
        put(ws, rr, "T", f'=IFERROR(S{rr}/S$11,"n.d.")', None, None, "0.0%", AL("right"), BOT)
        put(ws, rr, "U", f'=IFERROR(AVERAGEIFS(R_jRes,R_Crit,$P{rr},{PY("R_Res")}),"n.d.")', None, None, "0.0", AL("right"), BOT)
        put(ws, rr, "V", f'=IFERROR(COUNTIFS(R_okRes,1,R_Crit,$P{rr},{PY("R_Res")})/COUNTIFS(R_okRes,">=0",R_Crit,$P{rr},{PY("R_Res")}),"n.d.")',
            None, None, "0.0%", AL("right"), BOT)
    put(ws, 11, "P", "Total", F(True), border=BOT)
    put(ws, 11, "Q", "=SUM(Q7:Q10)", F(True), None, "0", AL("right"), BOT)
    put(ws, 11, "S", "=SUM(S7:S10)", F(True), None, "0", AL("right"), BOT)
    for c in "RTUV":
        put(ws, 11, c, None, border=BOT)

    band(ws, 13, "P", "V", "Répartition par domaine")
    header(ws, 14, "P", ["Domaine", "Période", "Préc.", "12 mois", "% 12 mois", "", ""], wrap=False)
    for i in range(len(C.DOMAINES)):
        rr = 15 + i
        put(ws, rr, "P", f"=INDEX(L_Domaines,{i + 1})", border=BOT)
        put(ws, rr, "Q", f"=COUNTIFS(R_Cpt,1,R_Dom,$P{rr},{PE('R_Enr')})", F(True), LIME, "0", AL("right"), BOT)
        put(ws, rr, "R", f"=COUNTIFS(R_Cpt,1,R_Dom,$P{rr},{PP('R_Enr')})", None, None, "0", AL("right"), BOT)
        put(ws, rr, "S", f"=COUNTIFS(R_Cpt,1,R_Dom,$P{rr},{PY('R_Enr')})", None, None, "0", AL("right"), BOT)
        put(ws, rr, "T", f'=IFERROR(S{rr}/$I${v01_row},"n.d.")', None, None, "0.0%", AL("right"), BOT)
        for c in "UV":
            put(ws, rr, c, None, border=BOT)
    assert 15 + len(C.DOMAINES) - 1 == 27

    r2 = 29
    band(ws, r2, "P", "V", "Répartition par fonds / véhicule (perte nette en EUR)")
    header(ws, r2 + 1, "P", ["Fonds / véhicule", "Période", "12 mois", "% 12 mois", "", "Perte 12 mois", ""], wrap=False)
    nf = 6
    for i in range(nf):
        rr = r2 + 2 + i
        put(ws, rr, "P", f"=INDEX(L_Fonds,{i + 1})", border=BOT)
        put(ws, rr, "Q", f"=COUNTIFS(R_Cpt,1,R_Fonds,$P{rr},{PE('R_Enr')})", F(True), LIME, "0", AL("right"), BOT)
        put(ws, rr, "R", f"=COUNTIFS(R_Cpt,1,R_Fonds,$P{rr},{PY('R_Enr')})", None, None, "0", AL("right"), BOT)
        put(ws, rr, "S", f'=IFERROR(R{rr}/$I${v01_row},"n.d.")', None, None, "0.0%", AL("right"), BOT)
        put(ws, rr, "U", f"=SUMIFS(R_PN,R_Cpt,1,R_Fonds,$P{rr},{PY('R_Enr')})", None, None, "#,##0", AL("right"), BOT)
        for c in "TV":
            put(ws, rr, c, None, border=BOT)

    r3 = r2 + 2 + nf + 1
    band(ws, r3, "P", "V", "Répartition par niveau d'escalade maximal")
    header(ws, r3 + 1, "P", ["Niveau", "Période", "%", "12 mois", "%", "", ""], wrap=False)
    for i in range(4):
        rr = r3 + 2 + i
        put(ws, rr, "P", f'="N"&INDEX(L_NivNum,{i + 1})&" - "&INDEX(L_NivLib,{i + 1})', border=BOT)
        put(ws, rr, "Q", f"=COUNTIFS(R_Niv,INDEX(L_NivNum,{i + 1}),{PE('R_Enr')})", F(True), LIME, "0", AL("right"), BOT)
        put(ws, rr, "R", f'=IFERROR(Q{rr}/$E${v01_row},"n.d.")', None, None, "0.0%", AL("right"), BOT)
        put(ws, rr, "S", f"=COUNTIFS(R_Niv,INDEX(L_NivNum,{i + 1}),{PY('R_Enr')})", None, None, "0", AL("right"), BOT)
        put(ws, rr, "T", f'=IFERROR(S{rr}/$I${v01_row},"n.d.")', None, None, "0.0%", AL("right"), BOT)
        for c in "UV":
            put(ws, rr, c, None, border=BOT)

    r4 = r3 + 2 + 4 + 1
    band(ws, r4, "P", "V", "Répartition par catégorie d'événement (Bâle II)")
    header(ws, r4 + 1, "P", ["Catégorie", "Période", "12 mois", "% 12 mois", "", "", ""], wrap=False)
    for i in range(len(C.BALE)):
        rr = r4 + 2 + i
        put(ws, rr, "P", f"=INDEX(L_Bale,{i + 1})", border=BOT)
        put(ws, rr, "Q", f"=COUNTIFS(R_Cpt,1,R_Bale,$P{rr},{PE('R_Enr')})", F(True), LIME, "0", AL("right"), BOT)
        put(ws, rr, "R", f"=COUNTIFS(R_Cpt,1,R_Bale,$P{rr},{PY('R_Enr')})", None, None, "0", AL("right"), BOT)
        put(ws, rr, "S", f'=IFERROR(R{rr}/$I${v01_row},"n.d.")', None, None, "0.0%", AL("right"), BOT)
        for c in "TUV":
            put(ws, rr, c, None, border=BOT)
    r = max(last_kpi, r4 + 2 + len(C.BALE)) + 1
    for t in notes:
        put(ws, r, "B", t, F(c=GREY), al=AL(wrap=True, v="top"))
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=22)
        ws.row_dimensions[r].height = row_height([(t, 260)])
        r += 1
    page(ws, f"B2:V{r}", one_page=True)
    return v01_row


def build(out, demo):
    fonds = C.FONDS_DEMO if demo else C.FONDS_PROD
    rows = []
    if demo:
        import demo as D
        rows = D.generate(fonds)
    bk = Book()
    wd = bk.wb.create_sheet("Dashboard")
    wr = bk.wb.create_sheet("Registre")
    wi = bk.wb.create_sheet("Indicateurs")
    wc = bk.wb.create_sheet("Champs")
    wn = bk.wb.create_sheet("Données")
    build_donnees(bk, wn, fonds, demo)
    build_registre(bk, wr, rows, demo)
    ind_rows = build_indicateurs(wi, None)
    build_champs(wc, fonds)
    build_dashboard(wd, ind_rows, demo)
    bk.wb.active = 0
    for ws in bk.wb.worksheets:
        ws.sheet_view.zoomScale = 100
    bk.wb.save(out)
    return rows


if __name__ == "__main__":
    outdir = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")
    build(os.path.join(outdir, "Registre_Incidents_KPI_Fundcraft.xlsx"), demo=False)
    build(os.path.join(outdir, "Registre_Incidents_KPI_Fundcraft_Demo.xlsx"), demo=True)
    print("ok")
