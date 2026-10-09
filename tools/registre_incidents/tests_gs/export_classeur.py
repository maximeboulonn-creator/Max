# -*- coding: utf-8 -*-
"""Exporte le classeur (valeurs calculées, fonds de cellule, formats, plages nommées) en JSON pour test.js.

Le JSON contient les données du registre : il reste hors du dépôt.
Usage : python export_classeur.py classeur.xlsx classeur.json
"""
import datetime as dt
import json
import sys

import openpyxl


def enc(v):
    if isinstance(v, dt.datetime):
        return {"$d": v.strftime("%Y-%m-%dT%H:%M:%S")}
    if isinstance(v, dt.time):
        return {"$d": "1899-12-30T" + v.strftime("%H:%M:%S")}
    if isinstance(v, dt.date):
        return {"$d": v.strftime("%Y-%m-%dT00:00:00")}
    return "" if v is None else v


def fond(x):
    f = x.fill
    rgb = f.fgColor.rgb if f and f.fill_type == "solid" and isinstance(f.fgColor.rgb, str) else "FFFFFFFF"
    return "#" + rgb[-6:].lower()


def main(src, out):
    wv, wf = openpyxl.load_workbook(src, data_only=True), openpyxl.load_workbook(src)
    data = {"sheets": {}, "names": {}}
    for ws in wv:
        f = wf[ws.title]
        lignes = range(1, ws.max_row + 1)
        cols = range(1, ws.max_column + 1)
        data["sheets"][ws.title] = {
            "values": [[enc(ws.cell(r, c).value) for c in cols] for r in lignes],
            "bg": [[fond(f.cell(r, c)) for c in cols] for r in lignes],
            "fmt": [[f.cell(r, c).number_format for c in cols] for r in lignes],
        }
    for n, d in wf.defined_names.items():
        feuille, ref = list(d.destinations)[0]
        data["names"][n] = {"sheet": feuille, "ref": ref.replace("$", "")}
    json.dump(data, open(out, "w", encoding="utf-8"), ensure_ascii=False)


if __name__ == "__main__":
    main(*sys.argv[1:3])
