# -*- coding: utf-8 -*-
"""Export XML du registre des incidents et des escalades (classeur KPI_Registre_Incidents_Fundcraft_France).

Lit les valeurs calculées du classeur (colonnes repérées par leur libellé, référentiels par leurs plages
nommées), produit un XML conforme à registre_incidents.xsd et le valide.
Usage : python export_xml.py classeur.xlsx sortie.xml [nom_du_fichier_source]
"""
import datetime as dt
import os
import re
import sys

import openpyxl
from lxml import etree

NS = "urn:fundcraft:registre-incidents"
XSD = os.path.join(os.path.dirname(os.path.abspath(__file__)), "registre_incidents.xsd")


# --------------------------------------------------------------------------- conversions
def vide(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", "-"))


def txt(v):
    return None if vide(v) else str(v).strip()


def num(v, nd=6):
    if vide(v) or isinstance(v, bool) or not isinstance(v, (int, float)):
        return None
    if abs(v - round(v)) < 1e-9:
        return str(int(round(v)))
    return f"{v:.{nd}f}".rstrip("0").rstrip(".")


def montant(v):
    return num(round(v, 2), 2) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def entier(v):
    return str(int(round(v))) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def date(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    return v.isoformat() if isinstance(v, dt.date) else None


def dateheure(v):
    if isinstance(v, dt.datetime):
        return v.replace(microsecond=0).isoformat()
    return dt.datetime.combine(v, dt.time()).isoformat() if isinstance(v, dt.date) else None


def booleen(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, str) and v.strip() in ("Oui", "Non"):
        return "true" if v.strip() == "Oui" else "false"
    return None


# --------------------------------------------------------------------------- lecture du classeur
class Classeur:
    def __init__(self, path):
        self.f = openpyxl.load_workbook(path)
        self.v = openpyxl.load_workbook(path, data_only=True)

    def plage(self, nom):
        feuille, ref = self.f.defined_names[nom].attr_text.split("!")
        ws = self.v[feuille.strip("'")]
        cells = ws[ref.replace("$", "")]
        if not isinstance(cells, tuple):
            return [cells.value]
        return [c.value for row in cells for c in (row if isinstance(row, tuple) else (row,))]

    def cellule(self, nom):
        return self.plage(nom)[0]

    def adresse(self, nom):
        feuille, ref = self.f.defined_names[nom].attr_text.split("!")
        m = re.match(r"\$?([A-Z]+)\$?(\d+)", ref)
        return feuille.strip("'"), m.group(1), int(m.group(2))


def sous(parent, tag, texte=None, /, **attrs):
    e = etree.SubElement(parent, f"{{{NS}}}{tag}", {k: v for k, v in attrs.items() if v is not None})
    if texte is not None:
        e.text = texte
    return e


def opt(parent, tag, texte, /, **attrs):
    """Crée l'élément seulement si la valeur est renseignée."""
    if texte is not None:
        return sous(parent, tag, texte, **attrs)
    return None


# --------------------------------------------------------------------------- blocs
def entete(root, wb, source, n):
    d = wb.v["Dashboard"]
    e = sous(root, "Entete")
    sgp = str(d["B3"].value or "").split(" - ")[0] or "Fundcraft France SAS"
    sous(e, "SocieteDeGestion", sgp)
    arrete = wb.cellule("DateArrete")
    sous(e, "DateArrete", date(arrete))
    fin = wb.cellule("FinTx") - dt.timedelta(days=1)
    sous(e, "Periode", code=str(wb.cellule("LibT")), debut=date(wb.cellule("DebutT")), fin=date(fin))
    sous(e, "PeriodePrecedente", code=str(wb.cellule("LibT1")), debut=date(wb.cellule("DebutT1")),
         fin=date(wb.cellule("DebutT") - dt.timedelta(days=1)))
    sous(e, "DouzeMoisGlissants", code="12 mois glissants", debut=date(wb.cellule("Debut12M")), fin=date(fin))
    sous(e, "FichierSource", source)
    sous(e, "NombreIncidents", str(n))


PARAMETRES = ["DelaiDecl", "SeuilAge", "Seuil_VL", "Recl_AR", "Recl_Mois", "Dep_Actif", "Dep_Passif", "FondsPropres",
              "HOuv", "HFer", "DORA_H_CLA", "DORA_H_INT", "DORA_M_FIN"]


def referentiels(root, wb):
    r = sous(root, "Referentiels")
    ve = sous(r, "Vehicules")
    for nom, typ, reg in zip(wb.plage("Ref_Fonds"), wb.plage("Ref_Type"), wb.plage("Ref_Regime")):
        if vide(nom):
            continue
        v = sous(ve, "Vehicule")
        sous(v, "Nom", txt(nom))
        opt(v, "Type", txt(typ))
        opt(v, "RegimeLiquidite", txt(reg))
    sc = sous(r, "SousCategories")
    for s, c in zip(wb.plage("Ref_SC"), wb.plage("Ref_SC_Cat")):
        if not vide(s):
            sous(sc, "SousCategorie", txt(s), categorie=txt(c))
    og = sous(r, "OriginesDetection")
    for o, c in zip(wb.plage("Ref_Orig"), wb.plage("Ref_Orig_Ctrl")):
        if not vide(o):
            sous(og, "Origine", txt(o), dispositifControle=booleen(c) or "false")
    ca = sous(r, "CausesRacines")
    for c in wb.plage("Ref_Cause"):
        if not vide(c):
            sous(ca, "Cause", txt(c))
    ne = sous(r, "NiveauxEscalade")
    for lib, rang in zip(wb.plage("Esc_Lib"), wb.plage("Esc_Rang")):
        if not vide(lib):
            sous(ne, "Niveau", txt(lib), rang=entier(rang))
    no = sous(r, "Notifications")
    for lib, h in zip(wb.plage("Notif_Lib"), wb.plage("Notif_H")):
        if not vide(lib):
            sous(no, "Notification", txt(lib), delaiHeures=entier(h))
    ba = sous(r, "BaremeGravite")
    cols = zip(wb.plage("Grav_Lib"), wb.plage("Bareme_Fin"), wb.plage("Bareme_Inv"), wb.plage("Bareme_VL"),
               wb.plage("Grav_Esc"), wb.plage("Grav_Delai"), wb.plage("SLA_PEC"), wb.plage("SLA_RES"))
    for i, (lib, fin, inv, vl, esc, dl, pec, res) in enumerate(cols, 1):
        n = sous(ba, "Niveau", rang=str(i), libelle=txt(lib))
        sous(n, "ImpactFinancierBrutMinEUR", montant(fin))
        sous(n, "InvestisseursLesesMin", entier(inv))
        sous(n, "EcartVLMinBps", num(vl))
        sous(n, "NiveauEscaladeRequis", txt(esc))
        sous(n, "DelaiCibleEscaladeJO", entier(dl))
        opt(n, "DelaiCiblePriseEnCompteHeuresOuvrees", num(pec))
        opt(n, "DelaiCibleResolutionJO", entier(res))
    pa = sous(r, "Parametres")
    for code in PARAMETRES:
        feuille, col, row = wb.adresse(code)
        ws = wb.v[feuille]
        val = ws[f"{col}{row}"].value
        if isinstance(val, dt.time):
            val = val.strftime("%H:%M")
        elif isinstance(val, float) and code in ("HOuv", "HFer"):
            m = round(val * 1440)
            val = f"{m // 60:02d}:{m % 60:02d}"
        else:
            val = num(val) if isinstance(val, (int, float)) else txt(val)
        src_col = chr(ord(col) + 1)
        sous(pa, "Parametre", code=code, libelle=txt(ws[f"B{row}"].value), valeur=val, source=txt(ws[f"{src_col}{row}"].value))
    se = sous(r, "SeuilsIndicateurs")
    feuille, col, row = wb.adresse("KRI_Lib")
    ws = wb.v[feuille]
    for i, lib in enumerate(wb.plage("KRI_Lib")):
        rr = row + i
        if vide(lib):
            continue
        sous(se, "Seuil", indicateur=txt(lib), sens=txt(ws[f"C{rr}"].value), vigilance=num(ws[f"D{rr}"].value),
             alerte=num(ws[f"E{rr}"].value), source=txt(ws[f"F{rr}"].value))


def incidents(root, wb):
    ws = wb.v["Registre"]
    h = {ws.cell(6, c).value: c for c in range(2, ws.max_column + 1) if ws.cell(6, c).value}
    grav = wb.plage("Grav_Lib")
    notif_dora = wb.cellule("Notif_DORA") if "Notif_DORA" in wb.f.defined_names else "DORA - incident majeur"
    aucune = wb.cellule("Notif_Aucune") if "Notif_Aucune" in wb.f.defined_names else "Aucune"
    sc_recl, cat_lim = wb.cellule("SC_Recl"), wb.cellule("Cat_Lim")
    inc = sous(root, "Incidents")
    n = 0
    for r in range(7, ws.max_row + 1):
        g = lambda lab: ws.cell(r, h[lab]).value  # noqa: E731
        if vide(g("Réf.")) or g("Détection (date et heure)") is None:
            continue
        n += 1
        e = sous(inc, "Incident", reference=txt(g("Réf.")))
        sous(e, "Vehicule", txt(g("Fonds ou véhicule")))
        opt(e, "TypeVehicule", txt(g("Type de véhicule")))
        opt(e, "RegimeLiquidite", txt(g("Régime de liquidité")))
        sous(e, "SousCategorie", txt(g("Sous-catégorie")))
        sous(e, "Categorie", txt(g("Catégorie")))
        sous(e, "Description", txt(g("Description")))
        opt(e, "OrigineDetection", txt(g("Origine de la détection")))
        opt(e, "AutoDetection", booleen(g("Auto-détection")))
        opt(e, "CauseRacine", txt(g("Cause racine")))
        ch = sous(e, "Chronologie")
        opt(ch, "DateSurvenance", date(g("Date de survenance")))
        sous(ch, "DateHeureDetection", dateheure(g("Détection (date et heure)")))
        opt(ch, "DateDeclaration", date(g("Déclaration au registre")))
        opt(ch, "DateHeurePriseEnCompte", dateheure(g("Prise en compte (date et heure)")))
        opt(ch, "DateHeureResolution", dateheure(g("Résolution (date et heure)")))
        opt(ch, "DateCloture", date(g("Date de clôture")))
        sous(ch, "Trimestre", txt(g("Trimestre")))
        sous(e, "Statut", txt(g("Statut")))
        gr = sous(e, "Gravite")
        rc = int(g("Gravité calculée"))
        sous(gr, "Calculee", txt(grav[rc - 1]), rang=str(rc))
        opt(gr, "Forcee", txt(g("Gravité forcée")))
        sous(gr, "Retenue", txt(g("Gravité retenue")), rang=entier(g("Gravité retenue (rang)")))
        im = sous(e, "Impact")
        opt(im, "PerteBruteEUR", montant(g("Impact financier brut (EUR)")))
        opt(im, "RecouvrementEUR", montant(g("Recouvrement (EUR)")))
        sous(im, "PerteNetteEUR", montant(g("Perte nette (EUR)")) or "0")
        opt(im, "EcartVLBps", num(g("Écart de VL (bps)")))
        opt(im, "InvestisseursLeses", entier(g("Investisseurs lésés")))
        opt(im, "IndemnisationInvestisseursEUR", montant(g("Indemnisation investisseurs (EUR)")))
        opt(im, "ImpactReglementaire", txt(g("Impact réglementaire")))
        opt(im, "ImpactReputationnel", txt(g("Impact réputationnel")))
        es = sous(e, "Escalade")
        sous(es, "NiveauRequis", txt(g("Niveau d'escalade requis")), rang=entier(g("Niveau requis (rang)")))
        sous(es, "NiveauAtteint", txt(g("Niveau d'escalade atteint")) or "Aucun", rang=entier(g("Niveau atteint (rang)")) or "0")
        opt(es, "DateEscalade", date(g("Date d'escalade")))
        opt(es, "DelaiJO", entier(g("Délai d'escalade (j.o.)")))
        sous(es, "DelaiCibleJO", entier(g("Délai cible d'escalade (j.o.)")))
        opt(es, "Conforme", booleen(g("Escalade conforme")))
        opt(es, "NombreReaffectations", entier(g("Nombre de réaffectations")))
        aut = txt(g("Notification à une autorité"))
        if aut and aut != aucune:
            no = sous(e, "Notification")
            sous(no, "Autorite", aut)
            opt(no, "Echeance", dateheure(g("Échéance de notification")))
            opt(no, "DateHeureRealisee", dateheure(g("Notification réalisée (date et heure)")))
            opt(no, "DansLeDelai", booleen(g("Notification dans le délai")))
        dora = [g("Classification DORA (date et heure)"), g("Rapport intermédiaire DORA (date et heure)"), g("Rapport final DORA (date)")]
        if aut == notif_dora or any(x is not None for x in dora):
            do = sous(e, "DORA")
            opt(do, "DateHeureClassification", dateheure(dora[0]))
            opt(do, "EcheanceNotificationInitiale", dateheure(g("DORA - échéance de notification initiale")))
            opt(do, "DateHeureRapportIntermediaire", dateheure(dora[1]))
            opt(do, "DateRapportFinal", date(dora[2]))
            opt(do, "NotificationInitialeConforme", booleen(g("DORA - notification initiale")))
            opt(do, "RapportIntermediaireConforme", booleen(g("DORA - rapport intermédiaire")))
            opt(do, "RapportFinalConforme", booleen(g("DORA - rapport final")))
            opt(do, "Conforme", booleen(g("DORA - notifications conformes")))
        if txt(g("Sous-catégorie")) == sc_recl or g("Accusé de réception (réclamation)") is not None:
            rc_ = sous(e, "Reclamation")
            opt(rc_, "DateAccuseReception", date(g("Accusé de réception (réclamation)")))
            opt(rc_, "DelaiAccuseReceptionJO", entier(g("Délai d'accusé de réception (j.o.)")))
            opt(rc_, "TraiteeDansLesDelais", booleen(g("Réclamation dans les délais")))
        if txt(g("Catégorie")) == cat_lim:
            de = sous(e, "Depassement")
            opt(de, "Type", txt(g("Dépassement actif ou passif")))
            opt(de, "RegulariseDansLeDelai", booleen(g("Dépassement régularisé dans le délai")))
        dl = sous(e, "Delais")
        opt(dl, "DetectionJours", entier(g("Délai de détection (jours)")))
        opt(dl, "DeclarationJO", entier(g("Délai de déclaration (j.o.)")))
        opt(dl, "DeclarationDansLeDelai", booleen(g("Déclaration dans le délai")))
        opt(dl, "PriseEnCompteHeuresOuvrees", num(g("Prise en compte (h ouvrées)"), 2))
        opt(dl, "CiblePriseEnCompteHeuresOuvrees", num(g("Cible de prise en compte (h ouvrées)")))
        opt(dl, "PriseEnCompteDansLeDelai", booleen(g("Prise en compte dans le délai")))
        opt(dl, "ResolutionJO", entier(g("Délai de résolution (j.o.)")))
        opt(dl, "CibleResolutionJO", entier(g("Cible de résolution (j.o.)")))
        opt(dl, "ResolutionDansLeDelai", booleen(g("Résolution dans le délai")))
        opt(dl, "ResoluSansEscaladeAuDelaDuMetier", booleen(g("Résolu sans escalade au-delà du métier")))
        opt(dl, "TempsEnEscaladeJO", entier(g("Temps en escalade (j.o.)")))
        opt(dl, "ClotureJours", entier(g("Délai de clôture (jours)")))
        su = sous(e, "Suivi")
        sous(su, "OuvertALArrete", booleen(g("Ouvert à l'arrêté")) or "false")
        opt(su, "AgeALArreteJours", entier(g("Âge à l'arrêté (jours)")))
        opt(su, "Reouvert", booleen(g("Réouvert")))
        opt(su, "Recurrent", booleen(g("Incident récurrent")))
        opt(su, "ControleSaisie", txt(g("Contrôle de saisie")))
        act = [g("Action corrective"), g("Responsable"), g("Échéance de l'action"), g("Action réalisée le")]
        if any(not vide(x) for x in act):
            ac = sous(e, "ActionCorrective")
            opt(ac, "Libelle", txt(act[0]))
            opt(ac, "Responsable", txt(act[1]))
            opt(ac, "Echeance", date(act[2]))
            opt(ac, "DateRealisation", date(act[3]))
            opt(ac, "EnRetard", booleen(g("Action en retard")))
            opt(ac, "RealiseeALEcheance", booleen(g("Action réalisée à l'échéance")))
    return inc, n


def trouver(ws, texte, debut=1, prefixe=False):
    for r in range(debut, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and (v.startswith(texte) if prefixe else v == texte):
            return r
    raise KeyError(texte)


def indicateurs(root, wb):
    d = wb.v["Dashboard"]
    sens = {txt(lib): txt(s) for lib, s in zip(wb.plage("KRI_Lib"), wb.plage("KRI_Sens")) if not vide(lib)}
    hr = next(r for r in range(1, 40) if d.cell(r, 2).value == "Indicateur" and d.cell(r, 3).value == "Unité")
    fin = trouver(d, "Ventilation par", hr, prefixe=True)
    ind = sous(root, "Indicateurs")
    sec = None
    for r in range(hr + 1, fin):
        b, c = d.cell(r, 2).value, d.cell(r, 3).value
        if vide(b):
            continue
        if vide(c):
            sec = sous(ind, "Section", libelle=txt(b))
            continue
        k = sous(sec, "Indicateur")
        sous(k, "Libelle", txt(b))
        sous(k, "Unite", txt(c))
        for col, tag in (("D", "ValeurPeriode"), ("E", "ValeurPeriodePrecedente"), ("F", "Variation"), ("G", "Valeur12Mois")):
            opt(k, tag, num(d[f"{col}{r}"].value))
        vg, al = num(d[f"H{r}"].value), num(d[f"I{r}"].value)
        if vg is not None and al is not None and sens.get(txt(b)):
            sous(k, "Seuil", sens=sens[txt(b)], vigilance=vg, alerte=al)
        st = txt(d[f"J{r}"].value)
        if st in ("Conforme", "Vigilance", "Alerte"):
            sous(k, "Statut", st)


def ventilation(parent, d, tag, titre, perimetre, total=True, escalades=False):
    r0 = trouver(d, titre)
    cats = [txt(d.cell(r0 + 1, c).value) for c in range(3, 7)]
    v = sous(parent, tag, perimetre=perimetre)
    r = r0 + 2
    lignes = []
    while r <= d.max_row:
        b = d.cell(r, 2).value
        if b == "Total" or (vide(b) and not total):
            break
        if not vide(b):
            lignes.append((r, "Ligne"))
        r += 1
    if total and d.cell(r, 2).value == "Total":
        lignes.append((r, "Total"))
    for rr, t in lignes:
        lg = sous(v, t, libelle=txt(d.cell(rr, 2).value))
        for i, cat in enumerate(cats):
            sous(lg, "Categorie", entier(d.cell(rr, 3 + i).value) or "0", libelle=cat)
        sous(lg, "Total", entier(d.cell(rr, 7).value) or "0")
        sous(lg, "SignificatifsEtMajeurs", entier(d.cell(rr, 8).value) or "0")
        sous(lg, "PerteNetteEUR", montant(d.cell(rr, 9).value) or "0")
        if escalades:
            opt(lg, "EscaladesConformes", num(d.cell(rr, 10).value))
        else:
            sous(lg, "OuvertsALArrete", entier(d.cell(rr, 10).value) or "0")


def ventilations(root, wb):
    d = wb.v["Dashboard"]
    v = sous(root, "Ventilations")
    ventilation(v, d, "ParVehicule", "Ventilation par fonds ou véhicule - 12 mois glissants", "12 mois glissants")
    ventilation(v, d, "ParRegimeLiquidite", "Ventilation par régime de liquidité - 12 mois glissants", "12 mois glissants")
    ventilation(v, d, "EvolutionTrimestrielle", "Évolution trimestrielle", "Cinq derniers trimestres", total=False, escalades=True)


def signales(root, wb):
    d = wb.v["Dashboard"]
    r0 = trouver(d, "Incidents significatifs et majeurs", prefixe=True)
    s = sous(root, "IncidentsSignales", perimetre=txt(d.cell(r0, 2).value))
    for r in range(r0 + 2, d.max_row + 1):
        b = d.cell(r, 2).value
        if vide(b) or not str(b).startswith("INC-"):
            if not vide(b) and str(b).startswith("Sources"):
                break
            continue
        sous(s, "IncidentSignale", reference=str(b).split(" - ")[0].strip())


def exporter(path, out, source=None):
    wb = Classeur(path)
    root = etree.Element(f"{{{NS}}}RegistreIncidents", nsmap={None: NS},
                         dateGeneration=dt.datetime.now().replace(microsecond=0).isoformat())
    tmp = etree.Element("tmp")
    inc, n = incidents(tmp, wb)
    entete(root, wb, source or os.path.basename(path), n)
    referentiels(root, wb)
    root.append(inc)
    indicateurs(root, wb)
    ventilations(root, wb)
    signales(root, wb)
    doc = etree.ElementTree(root)
    schema = etree.XMLSchema(etree.parse(XSD))
    ok = schema.validate(doc)
    etree.indent(doc, space="  ")
    doc.write(out, xml_declaration=True, encoding="UTF-8", pretty_print=True)
    return ok, schema.error_log, n


if __name__ == "__main__":
    ok, log, n = exporter(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
    print(f"{n} incidents exportés ; conforme au schéma : {ok}")
    for e in list(log)[:20]:
        print(f"  ligne {e.line} : {e.message}")
    sys.exit(0 if ok else 1)
