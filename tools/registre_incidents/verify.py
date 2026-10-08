# -*- coding: utf-8 -*-
"""Contrôle indépendant du tableau de bord de démonstration.

Recalcule en Python chaque indicateur à partir des données brutes du jeu de
démonstration et compare aux valeurs du classeur recalculé par LibreOffice.
Usage : python verify.py Registre_Incidents_KPI_Fundcraft_Demo.xlsx (déjà recalculé)
"""
import os
import sys
import datetime as dt
import statistics

import numpy as np
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C  # noqa: E402
import demo as D  # noqa: E402
from build import jours_feries  # noqa: E402

HOL = {d for d, _ in jours_feries()}
O, CL = 9 / 24, 18 / 24
HJ = 9.0
CRIT = {lab: (i + 1, pec, res, niv, dl) for i, (lab, pec, res, niv, dl, _) in enumerate(C.CRITICITES)}
INTERNE = dict(C.SOURCES)


def wd(d):
    return d.weekday() < 5 and d not in HOL


def nwd(a, b):
    n, d = 0, a
    while d <= b:
        n += wd(d)
        d += dt.timedelta(days=1)
    return n


def frac(t):
    return (t.hour * 60 + t.minute) / 1440


def bh(s, e):
    if s is None or e is None or e < s:
        return None
    v = ((nwd(s.date(), e.date()) - 1) * (CL - O)
         + (statistics.median([frac(e), CL, O]) if wd(e.date()) else CL)
         - (statistics.median([frac(s), CL, O]) if wd(s.date()) else O)) * 24
    return round(v, 4)


def workday(d, n):
    while n > 0:
        d += dt.timedelta(days=1)
        if wd(d):
            n -= 1
    return d


def act(r, now):
    if r.get("Real"):
        return int(r["Real"] <= r["Ech"])
    return 0 if now > dt.datetime.combine(r["Ech"] + dt.timedelta(days=1), dt.time()) else None


def enrich(r, now):
    x = dict(r)
    x["cpt"] = r["Type"] == "Incident" and r["Statut"] != "Annulé"
    x["qi"] = r["Type"] == "Quasi-incident" and r["Statut"] != "Annulé"
    n, pec, res, niv_m, dl = CRIT[r["Crit"]]
    x["critn"] = n
    if not x["cpt"]:
        x["actok"] = None
        if x["qi"] and r.get("ActReq") == "Oui" and r.get("Ech"):
            x["actok"] = act(r, now)
        return x
    x["int"] = 1 if INTERNE[r["Src"]] == "Oui" else 0
    x["jdet"] = round(bh(r.get("Surv"), r["Det"]) / HJ, 4) if r.get("Surv") and bh(r.get("Surv"), r["Det"]) is not None else None
    x["hdecl"] = bh(r["Det"], r["Enr"])
    x["hpec"] = bh(r["Enr"], r.get("PEC"))
    x["jres"] = round(bh(r["Enr"], r.get("Res")) / HJ, 4) if r.get("Res") else None
    x["jclo"] = round(bh(r["Enr"], r.get("Clo")) / HJ, 4) if r.get("Clo") else None
    x["okpec"] = None if x["hpec"] is None else int(x["hpec"] <= pec)
    x["okres"] = None if x["jres"] is None else int(round(x["jres"], 4) <= res)
    x["dl"] = dt.datetime.combine(workday(r["Enr"].date(), res), r["Enr"].time())
    x["niv"] = 4 if (r.get("Not") or r.get("DInit")) else 3 if r.get("EscN3") else 2 if r.get("EscN2") else 1
    nivr = max(niv_m, 4 if (r.get("DORA") == "Oui" or r.get("NotReq") == "Oui") else 1)
    x["escok"] = int(x["niv"] >= nivr)
    desc = None
    if niv_m == 2:
        cands = [t for t in (r.get("EscN2"), r.get("EscN3")) if t]
        desc = min(cands) if cands else None
    elif niv_m >= 3:
        desc = r.get("EscN3")
    x["hesc"] = bh(r["Det"], desc) if desc else None
    x["escdel"] = None if x["hesc"] is None else int(x["hesc"] <= dl)
    first = min([t for t in (r.get("EscN2"), r.get("EscN3")) if t], default=None)
    x["jesc"] = round(bh(first, r["Res"]) / HJ, 4) if first and r.get("Res") and bh(first, r["Res"]) is not None else None
    x["resn1"] = None if not r.get("Res") else int(x["niv"] == 1)
    x["reo"] = int(r.get("Reo") == "Oui")
    x["rec"] = int(bool(r.get("Lien")))
    x["rca"] = int(r.get("RCA") == "Oui")
    x["notok"] = None
    if r.get("NotReq") == "Oui":
        if r.get("Not"):
            x["notok"] = int(r["Not"] <= r["NotLim"])
        elif r.get("NotLim") and now > r["NotLim"]:
            x["notok"] = 0
    x["doraok"] = None
    if r.get("DORA") == "Oui":
        det, cla = r["Det"], r.get("DClas")
        if cla is None:
            dlt = det + dt.timedelta(hours=24)
        elif cla - det <= dt.timedelta(hours=24):
            dlt = min(cla + dt.timedelta(hours=4), det + dt.timedelta(hours=24))
        else:
            dlt = cla + dt.timedelta(hours=4)
        checks = [int(r["DInit"] <= dlt) if r.get("DInit") else (0 if now > dlt else None)]
        if r.get("DInit"):
            lim = r["DInit"] + dt.timedelta(hours=72)
            checks.append(int(r["DInter"] <= lim) if r.get("DInter") else (0 if now > lim else None))
        if r.get("DInter"):
            m = r["DInter"].date()
            lim = dt.date(m.year + (m.month == 12), m.month % 12 + 1, m.day)
            checks.append(int(r["DFin"] <= lim) if r.get("DFin") else (0 if now.date() > lim else None))
        checks = [c for c in checks if c is not None]
        x["doraok"] = min(checks) if checks else None
    x["pn"] = (r.get("PB") or 0) - (r.get("Recup") or 0) if (r.get("PB") is not None or r.get("Recup") is not None) else None
    x["actok"] = None
    if r.get("ActReq") == "Oui" and r.get("Ech"):
        x["actok"] = act(r, now)
    return x


def kpis(rows, a, b):
    b1 = dt.datetime.combine(b + dt.timedelta(days=1), dt.time())
    a0 = dt.datetime.combine(a, dt.time())
    inp = lambda t: t is not None and a0 <= (t if isinstance(t, dt.datetime) else dt.datetime.combine(t, dt.time())) < b1  # noqa: E731
    inc = [r for r in rows if r["cpt"]]
    per = [r for r in inc if inp(r["Enr"])]
    resd = [r for r in inc if inp(r.get("Res"))]
    clo = [r for r in inc if inp(r.get("Clo"))]
    n = len(per)
    mean = lambda v: (sum(v) / len(v)) if v else "n.d."  # noqa: E731
    rat = lambda num, den: (num / den) if den else "n.d."  # noqa: E731
    flags = lambda lst, k: [r[k] for r in lst if r.get(k) is not None]  # noqa: E731
    out = {"V01": n}
    out["V02"] = rat(sum(1 for r in rows if r["qi"] and inp(r["Enr"])), n)
    out["V03"] = rat(sum(1 for r in per if r["critn"] <= 2), n)
    out["V04"] = sum(1 for r in per if r["critn"] == 1)
    f = flags(per, "int")
    out["V05"] = rat(sum(f), len(f))
    doms = {}
    for r in per:
        doms[r["Dom"]] = doms.get(r["Dom"], 0) + 1
    out["V06"] = rat(max(doms.values()) if doms else 0, n)
    out["V07"] = sum(r["pn"] or 0 for r in per)
    out["V08"] = sum(1 for r in per if r.get("DORA") == "Oui")
    out["D01"] = mean(flags(per, "jdet"))
    out["D02"] = mean(flags(per, "hdecl"))
    out["D03"] = mean(flags(per, "hpec"))
    f = flags(per, "okpec")
    out["D04"] = rat(sum(f), len(f))
    jr = flags(resd, "jres")
    out["D05"] = mean(jr)
    out["D06"] = float(np.percentile(jr, 90)) if jr else "n.d."
    f = flags(resd, "okres")
    out["D07"] = rat(sum(f), len(f))
    out["D08"] = mean(flags(clo, "jclo"))
    f = [r["resn1"] for r in resd if r["critn"] >= 3 and r["resn1"] is not None]
    out["E01"] = rat(sum(f), len(f))
    out["E02"] = rat(len(resd), n)
    backlog = [r for r in inc if r["Enr"] < b1 and not (r.get("Res") and r["Res"] < b1)]
    out["E03"] = len(backlog)
    out["E04"] = rat(sum(1 for r in backlog if r["dl"] < b1), len(backlog))
    out["E05"] = rat(sum(r["reo"] for r in resd), len(resd))
    out["E06"] = rat(sum(r["rec"] for r in per), n)
    out["E07"] = rat(sum(r["rca"] for r in clo if r["critn"] <= 2), sum(1 for r in clo if r["critn"] <= 2))
    acts = [r["actok"] for r in rows if (r["cpt"] or r["qi"]) and r.get("actok") is not None and inp(r.get("Ech"))]
    out["E08"] = rat(sum(acts), len(acts))
    out["X01"] = rat(sum(1 for r in per if r["niv"] >= 2), n)
    out["X02"] = rat(sum(1 for r in per if r["niv"] >= 3), n)
    out["X03"] = rat(sum(1 for r in per if r["niv"] == 4), n)
    f = flags(per, "escok")
    out["X04"] = rat(sum(f), len(f))
    f = flags(per, "escdel")
    out["X05"] = rat(sum(f), len(f))
    out["X06"] = mean(flags(per, "hesc"))
    out["X07"] = mean(flags(resd, "jesc"))
    reaf = [r["Reaf"] for r in per if r.get("Reaf") is not None]
    out["X08"] = rat(sum(1 for v in reaf if v >= 2), len(reaf))
    f = flags(per, "notok")
    out["X09"] = rat(sum(f), len(f))
    f = flags(per, "doraok")
    out["X10"] = rat(sum(f), len(f))
    return out


def main(path):
    now = dt.datetime.now()
    rows = [enrich(r, now) for r in D.generate(C.FONDS_DEMO)]
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Dashboard"]
    cols = {"E": (ws["E7"].value.date(), ws["E8"].value.date()), "H": (ws["H7"].value.date(), ws["H8"].value.date()),
            "I": (ws["I7"].value.date(), ws["I8"].value.date())}
    refrow = {ws.cell(r, 2).value: r for r in range(13, 60) if isinstance(ws.cell(r, 2).value, str) and len(ws.cell(r, 2).value) == 3}
    bad = 0
    for col, (a, b) in cols.items():
        exp = kpis(rows, a, b)
        for ref, v in exp.items():
            got = ws[f"{col}{refrow[ref]}"].value
            ok = (got == v) if isinstance(v, str) or isinstance(got, str) else abs(got - v) < 1e-6
            if not ok:
                bad += 1
                print(f"ECART {ref} colonne {col} : classeur={got} python={v}")
    print(f"{len(cols) * len(C.KPIS)} contrôles, {bad} écart(s)")
    return bad


if __name__ == "__main__":
    sys.exit(1 if main(sys.argv[1]) else 0)
