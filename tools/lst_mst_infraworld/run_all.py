"""Construit le classeur pour chaque taille, recalcule sous LibreOffice, collecte les résultats, puis construit le classeur final."""
import subprocess,json,os,shutil,sys
import openpyxl
HERE=os.path.dirname(os.path.abspath(__file__))
def recalc(src,outdir):
    os.makedirs(outdir,exist_ok=True)
    subprocess.run(['timeout','500','soffice','--headless','--convert-to','xlsx','--outdir',outdir,src],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    return os.path.join(outdir,os.path.basename(src))
def metrics(path):
    bd=openpyxl.load_workbook(path,data_only=True)
    names={}
    for n,d in bd.defined_names.items():
        sh,ref=d.attr_text.split('!')
        if ':' in ref: continue
        names[n]=bd[sh.strip("'")][ref.replace('$','')].value
    out={}
    for c in ('S1','S2','S3','S4'):
        out[c]={k:names[f'LST_{c}_{k}'] for k in ('gated','file8','susp','poche','couv','cession','pertes','an8')}
        out[c]['creux']=names[f'MST_{c}_creux']; out[c]['verdict']=names[f'Joint_{c}']
        if hasattr(out[c]['susp'],'strftime'): out[c]['susp']=out[c]['susp'].strftime('%d/%m/%Y')
    return out,names
res={}
for k in range(1,6):
    tmp=os.path.join(HERE,f'tmp_size{k}.xlsx')
    subprocess.run([sys.executable,os.path.join(HERE,'build.py'),str(k),'',tmp],check=True,cwd=HERE,stdout=subprocess.DEVNULL)
    rc=recalc(tmp,os.path.join(HERE,'recalc_tmp'))
    m,names=metrics(rc); res[str(k)]=m
    print(k,'AN',names['AN_depart'],'frais',round(names['Frais_an'],4),{c:(m[c]['verdict'],round(m[c]['file8'],3),round(m[c]['poche'],3)) for c in m})
json.dump(res,open(os.path.join(HERE,'results.json'),'w'),default=str)
final=os.path.join(HERE,'LST_MST_Openstone_Infraworld.xlsx')
subprocess.run([sys.executable,os.path.join(HERE,'build.py'),'1',os.path.join(HERE,'results.json'),final],check=True,cwd=HERE,stdout=subprocess.DEVNULL)
rc=recalc(final,os.path.join(HERE,'recalc'))
# error check
bd=openpyxl.load_workbook(rc,data_only=True); a=openpyxl.load_workbook(final); errs=0
for w in bd.worksheets:
    for row in a[w.title].iter_rows():
        for c in row:
            v0=c.value
            if hasattr(v0,'text'): v0=v0.text
            if isinstance(v0,str) and v0.startswith('='):
                v=w[c.coordinate].value
                if v is None or (isinstance(v,str) and v.startswith('#')):
                    errs+=1
                    if errs<8: print('ERR',w.title,c.coordinate,v,v0[:90])
print('errors',errs,'->',rc)
