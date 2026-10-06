# -*- coding: utf-8 -*-
import openpyxl,datetime as dt
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
DARK='2F4858'; TEAL='218E8E'; LIME='D9F785'; CYAN='EBFCF9'; PEACH='FFF8F2'; BLACK='1A2027'; GREY='6A7A85'; LINE='E0E0E0'
GREEN='228B22'; AMBER='DAA520'; RED='CC0000'; WHITE='FFFFFF'; INPUT_TXT='002060'
NB=' '
def F(b=False,c=BLACK,s=9,i=False): return Font(name='Calibri',size=s,bold=b,color=c,italic=i)
def fill(c): return PatternFill('solid',fgColor=c)
def al(h='left',v='center',wrap=True): return Alignment(horizontal=h,vertical=v,wrap_text=wrap)
thin=Side(style='thin',color=LINE)
B_BOT=Border(bottom=thin)
PCT='0.0 %'; PCT2='0.00 %'; PCT0='0 %'; EUR='#,##0'; MEUR='#,##0.0'; D='0'; DATE='dd/mm/yyyy'; DEC='0.00'; D3='0.000'
class Book:
    def __init__(s):
        s.wb=openpyxl.Workbook(); s.wb.remove(s.wb.active); s.names={}
    def name(s,nm,ws,coord):
        s.wb.defined_names[nm]=DefinedName(nm,attr_text=f"'{ws.title}'!{coord}")
        s.names[nm]=f"'{ws.title}'!{coord}"
class Sheet:
    def __init__(s,book,title,heading,sub,widths,ncols=12):
        s.bk=book; s.ws=book.wb.create_sheet(title); s.ws.sheet_view.showGridLines=False; s.ncols=ncols
        for i,w in enumerate(widths): s.ws.column_dimensions[L(i+1)].width=w
        c=s.ws.cell(2,2,heading); c.font=F(True,DARK,12)
        c=s.ws.cell(3,2,sub); c.font=F(False,GREY); c.alignment=al(wrap=True,v='top')
        s.ws.merge_cells(start_row=3,start_column=2,end_row=3,end_column=1+ncols); s.ws.row_dimensions[3].height=42
        s.r=5
    def cell(s,r,c,v=None,fmt=None,font=None,fl=None,alg=None,border=None):
        x=s.ws.cell(r,c)
        if v is not None: x.value=v
        x.font=font or F()
        if fmt: x.number_format=fmt
        if fl: x.fill=fill(fl)
        x.alignment=alg or al('left')
        if border: x.border=border
        return x
    def sec(s,text,note=None):
        r=s.r
        for c in range(2,2+s.ncols):
            x=s.ws.cell(r,c); x.fill=fill(TEAL)
        c=s.ws.cell(r,2,text); c.font=F(True,WHITE,9.5); c.alignment=al('left',wrap=False)
        s.ws.row_dimensions[r].height=18; s.r+=1
        if note:
            c=s.ws.cell(s.r,2,note); c.font=F(False,GREY,i=True); c.alignment=al(wrap=True,v='top')
            s.ws.merge_cells(start_row=s.r,start_column=2,end_row=s.r,end_column=1+s.ncols); s.ws.row_dimensions[s.r].height=15 if len(note)<140 else 28; s.r+=1
        return r
    def hdr(s,cols,h=28,left=1):
        r=s.r
        for i,v in enumerate(cols):
            x=s.ws.cell(r,2+i,v); x.font=F(True,WHITE); x.fill=fill(DARK); x.alignment=al('left' if i<left else 'center',wrap=True)
        s.ws.row_dimensions[r].height=h; s.r+=1; return r
    def row(s,vals,fmts=None,inputs=(),bold_first=True,h=None,aligns=None,key=(),merge=None,font_color=None):
        r=s.r
        for i,v in enumerate(vals):
            if v is None: continue
            x=s.ws.cell(r,2+i,v)
            x.font=F(bold_first and i==0, INPUT_TXT if i in inputs else (font_color or BLACK)) if i not in key else F(True)
            if fmts and i<len(fmts) and fmts[i]: x.number_format=fmts[i]
            if i in inputs: x.fill=fill(CYAN)
            if i in key: x.fill=fill(LIME)
            a=(aligns[i] if aligns and i<len(aligns) and aligns[i] else ('left' if i==0 or isinstance(v,str) and not str(v).startswith('=') else 'right'))
            x.alignment=al(a,wrap=True)
            x.border=B_BOT
        if merge:
            for a,b in merge: s.ws.merge_cells(start_row=r,start_column=a,end_row=r,end_column=b)
        if h: s.ws.row_dimensions[r].height=h
        s.r+=1; return r
    def note(s,text,h=None):
        r=s.r; c=s.ws.cell(r,2,text); c.font=F(False,GREY); c.alignment=al(wrap=True,v='top')
        s.ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=1+s.ncols)
        s.ws.row_dimensions[r].height=h or (15 if len(text)<150 else 30 if len(text)<300 else 44); s.r+=1; return r
    def blank(s,n=1): s.r+=n
    def nm(s,nm,r,c=3): s.bk.name(nm,s.ws,f'${L(c)}${r}')
    def ref(s,r,c): return f"'{s.ws.title}'!${L(c)}${r}"
def verdict_cf(ws,rng):
    from openpyxl.formatting.rule import CellIsRule,FormulaRule
    for word,col in (('Vert',GREEN),('Orange',AMBER),('Rouge',RED)):
        ws.conditional_formatting.add(rng,FormulaRule(formula=[f'LEFT({rng.split(":")[0].replace("$","")},{len(word)})="{word}"'],fill=fill(col),font=Font(name='Calibri',size=9,bold=True,color=WHITE)))
