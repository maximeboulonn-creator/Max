#!/usr/bin/env python3
"""Extrait le modèle de frais du classeur Fundcraft (v3) en JSON pour l'Operating Memorandum HTML.

Lecture cellule par cellule de 01_Input (§1 à §4), 02_Classes_de_Parts, 05_Tresorerie (§1 à §3, §8)
et 06_Documentation. Les cellules pilotées par formule (barème §3D, classe sélectionnée, paramètres
de valorisation...) deviennent des références {"ref": "..."} résolues par le moteur JavaScript, afin
que la logique du classeur (cellules grises) soit conservée telle quelle dans l'interface.

Usage : python3 -I build_fees_json.py <classeur.xlsx> <sortie.json>
"""
import json
import re
import sys

import openpyxl

INF = 999999999999  # sentinelle « sans borne » du classeur

REF_FORMULAS = {
    '=Seuil_Bareme_T1': 'seuil_t1',
    '=Seuil_Bareme_T2': 'seuil_t2',
    '=Gestion_Taux_T1': 'gestion_t1',
    '=Gestion_Taux_T2': 'gestion_t2',
    '=Gestion_Taux_T3': 'gestion_t3',
    '=Conseil_Taux_T1': 'conseil_t1',
    '=Conseil_Taux_T2': 'conseil_t2',
    '=Conseil_Taux_T3': 'conseil_t3',
    '=Distrib_Global_Classe_Selectionnee': 'distrib_global_classe',
    '=Distrib_SousDistrib_Classe_Selectionnee': 'distrib_sous_classe',
    '=IF(Classe_Selectionnee_Distribue="Oui",Nb_Distributions_Par_An*Cout_Par_Distribution_CAC,0)': 'cac_distribution',
    '=IF(Nb_Lignes_Valorisation<=10,0,ROUNDUP((Nb_Lignes_Valorisation-10)/10,0)*500)': 'valo_supplement',
    '=Cout_Marginal_Investisseur*Nb_Investisseurs': 'tenue_passif_variable',
    '=MAX(Plancher_Constitution,Taux_Constitution*ActifNet_Prev_An1)': 'constitution_montant',
    '=Plancher_Constitution': 'constitution_plancher',
}
VALO_RE = re.compile(r'^=IF\(Frequence_VL="Mensuelle",(\d+),(\d+)\)\*\(1\+Majoration_Valorisation\)$')

LINE_IDS = ['gestion', 'conseil', 'distrib_global', 'distrib_sous', 'depositaire', 'cac', 'cac_distribution',
            'rapport_annuel', 'valorisation', 'reporting_aifm', 'tenue_passif', 'finologee', 'kyc', 'constitution',
            'depo_liquidation', 'euroclear', 'solvabilite2', 'libre18', 'libre19', 'libre20']

# Compléments propres à l'Operating Memorandum (hors classeur) : bénéficiaire lisible, TVA, source Prospectus
# pour les lignes dont la colonne « Commentaire / Source » du classeur est vide.
EXTRAS = {
    'gestion': dict(beneficiaire='Société de Gestion (Fundcraft)', tva=0,
                    source="Prospectus Art. 9.2(a) : 0,18 % jusqu'à 100 M€, 0,14 % de 100 à 500 M€, 0,09 % au-delà ; minimum 60 000 € par an, facturé trimestriellement à hauteur de 15 000 € minimum ; calcul et prélèvement mensuels, prorata temporis. Barème saisi en 01_Input §3D (cellules grises ci-dessous)."),
    'conseil': dict(beneficiaire='Société de Conseil (AirFund / ONEWEALTHPLACE)', tva=0,
                    source="Prospectus Art. 9.2(b) : 0,45 % jusqu'à 100 M€, 0,35 % de 100 à 500 M€, 0,25 % au-delà ; paiement mensuel, prorata temporis, à compter du Premier Jour de Souscription. Barème saisi en 01_Input §3D."),
    'distrib_global': dict(beneficiaire='Distributeur Global (ONEWEALTHPLACE)', tva=0,
                           source="Prospectus Art. 9.7(i) · Convention-cadre de distribution globale fundcraft × OneWealthPlace du 8 juillet 2026 : 0,10 % par an sur les Parts A et B. Taux de la classe sélectionnée (02_Classes_de_Parts)."),
    'distrib_sous': dict(beneficiaire='Sous-distributeurs (via le Distributeur Global)', tva=0,
                         source="Prospectus Art. 9.7(ii) · Addendum n° 1 et Annexe 2 de la convention de distribution : 0,90 % par an sur les seules Parts A, reversés par le Distributeur Global aux sous-distributeurs. Taux de la classe sélectionnée (02_Classes_de_Parts)."),
    'depositaire': dict(tva=0.2),
    'cac': dict(tva=0.2),
    'cac_distribution': dict(beneficiaire='CAC', tva=0.2,
                             source="Pratique tarifaire générique du CAC (Aplitec, e-mail du 11/09/2026) : 1 600 € par distribution certifiée. Forfait = nombre de distributions par an × coût par distribution (01_Input §3), uniquement si la classe sélectionnée est distribuante."),
    'rapport_annuel': dict(tva=0.2),
    'valorisation': dict(tva=0.2),
    'reporting_aifm': dict(tva=0.2),
    'tenue_passif': dict(tva=0.2),
    'finologee': dict(tva=0.2),
    'kyc': dict(beneficiaire='Société de Gestion (Fundcraft)', tva=0,
                source="Prospectus Art. 9.2(f)(i) : 100 diligences KYC incluses par tranche cumulée de 100 M€ ; au-delà, 10 € par diligence. Formule du classeur : MAX(0 ; Investisseurs − 100 × (ENT(Actif Net ÷ 100 M€) + 1)) × 10 € (01_Input §3)."),
    'constitution': dict(beneficiaire='Divers (avocat, dépositaire, AMF...)', tva=0,
                         source="Prospectus Art. 9.6 : MAX(60 000 € ; 1 % HT de l'Actif Net au premier anniversaire du Premier Jour de Souscription). Ligne volontairement inactive dans le barème : les Frais de Constitution sont traités en 05_Tresorerie §8 (provision sur le 1er exercice, suivi des factures réelles)."),
    'depo_liquidation': dict(tva=0.2),
    'euroclear': dict(tva=0.2),
    'solvabilite2': dict(tva=0.2),
}
PRORATA_EXCEPTION_ROW = 78  # 05_Tresorerie : Finologee - provision sans prorata sur le mois de lancement


def cell(ws, ref):
    return ws[ref].value


def num(v):
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return None if v == INF else (int(v) if float(v).is_integer() else float(v))
    raise ValueError('valeur non numérique : %r' % (v,))


def token(v):
    """Cellule du tableau de frais : nombre, ou référence si la cellule porte une formule connue."""
    if isinstance(v, str) and v.startswith('='):
        if v in REF_FORMULAS:
            return {'ref': REF_FORMULAS[v]}
        m = VALO_RE.match(v)
        if m:
            idx = {'12000': 1, '14000': 2, '16000': 3, '18000': 4}[m.group(1)]
            return {'ref': 'valo_p%d' % idx}
        raise ValueError('formule non reconnue dans le tableau de frais : %s' % v)
    return num(v)


def iso(v):
    return v.strftime('%Y-%m-%d') if hasattr(v, 'strftime') else (str(v) if v else '')


def main(src, dst):
    wbf = openpyxl.load_workbook(src, data_only=False)
    wbv = openpyxl.load_workbook(src, data_only=True)
    inp, inpv = wbf['01_Input'], wbv['01_Input']
    cls = wbf['02_Classes_de_Parts']
    tre, trev = wbf['05_Tresorerie'], wbv['05_Tresorerie']
    doc = wbf['06_Documentation']
    recap = wbf['04_Recap']

    # ---- 01_Input §1-§3D
    fees = {
        '_doc': "Modèle de frais Fundcraft - port fidèle du classeur Fundcraft_AirFund_Modele_Frais_3.xlsx (onglets 01_Input, 02_Classes_de_Parts, 03_Detail_Depositaire_Valo, 04_Recap, 05_Tresorerie). Les valeurs ci-dessous sont celles du classeur ; les saisies de l'utilisateur sont des surcharges conservées dans le fichier.",
        'source_file': src.split('/')[-1],
        'nom_fonds': cell(inp, 'F8'),
        'echelle': {'min': num(cell(inp, 'F11')), 'max': num(cell(inp, 'F12')), 'pas': num(cell(inp, 'F13')), 'max_paliers': 150},
        'params': {
            'nb_investisseurs': num(cell(inp, 'F17')),
            'cout_marginal_investisseur': num(cell(inp, 'F18')),
            'actif_net_prev_an1': num(cell(inp, 'F19')),
            'taux_constitution': num(cell(inp, 'F20')),
            'plancher_constitution': num(cell(inp, 'F21')),
            'nb_distributions': num(cell(inp, 'F22')),
            'cout_par_distribution_cac': num(cell(inp, 'F23')),
            'nb_kyc_gratuites': num(cell(inp, 'F24')),
            'taille_tranche_kyc': num(cell(inp, 'F25')),
            'cout_kyc_supp': num(cell(inp, 'F26')),
        },
        'params_meta': {
            'nb_investisseurs': {'label': cell(inp, 'B17'), 'kind': 'int', 'note': 'alimente la Tenue du Passif (montant par porteur) et les Frais KYC'},
            'cout_marginal_investisseur': {'label': cell(inp, 'B18'), 'kind': 'eur', 'note': 'Annexe 11 de la Convention Dépositaire BFCM : 15 € par porteur et par an'},
            'actif_net_prev_an1': {'label': cell(inp, 'B19'), 'kind': 'eur', 'note': "0 dans le classeur : c'est le plancher de 60 000 € qui s'applique ; au-delà de 6 M€, 1 % de l'Actif Net dépasse le plancher"},
            'taux_constitution': {'label': cell(inp, 'B20'), 'kind': 'pct', 'note': 'Prospectus Art. 9.6'},
            'plancher_constitution': {'label': cell(inp, 'B21'), 'kind': 'eur', 'note': 'Prospectus Art. 9.6 - provisionné sur le 1er exercice (05_Tresorerie §8)'},
            'nb_distributions': {'label': cell(inp, 'B22'), 'kind': 'int', 'note': '0 : Fonds de capitalisation (Prospectus Art. 7.2) ; pilote aussi la fréquence de facturation du forfait CAC par distribution'},
            'cout_par_distribution_cac': {'label': cell(inp, 'B23'), 'kind': 'eur', 'note': 'pratique tarifaire du CAC, 1 600 € par distribution certifiée'},
            'nb_kyc_gratuites': {'label': cell(inp, 'B24'), 'kind': 'int', 'note': 'Prospectus Art. 9.2(f)(i)'},
            'taille_tranche_kyc': {'label': cell(inp, 'B25'), 'kind': 'eur', 'note': 'Prospectus Art. 9.2(f)(i)'},
            'cout_kyc_supp': {'label': cell(inp, 'B26'), 'kind': 'eur', 'note': 'Prospectus Art. 9.2(f)(i)'},
        },
        'classe_selectionnee': cell(inp, 'F29'),
        'valo': {
            'frequence_vl': cell(inp, 'F35'),
            'nb_lignes': num(cell(inp, 'F36')),
            'paliers': {'Mensuelle': [], 'Bi-mensuelle': []},
            'majoration': [[5, 0.15], [3, 0.10]],
            'supplement': {'lignes_incluses': 10, 'par_tranche': 10, 'montant': 500},
        },
        'baremes': {
            'seuil_t1': num(cell(inp, 'F42')), 'seuil_t2': num(cell(inp, 'F43')),
            'gestion': [num(cell(inp, 'E46')), num(cell(inp, 'E47')), num(cell(inp, 'E48'))],
            'conseil': [num(cell(inp, 'F46')), num(cell(inp, 'F47')), num(cell(inp, 'F48'))],
        },
        'reference': {'taux_bp': 187, 'note': "Taux de référence (hors classeur) auquel la charge directe du Fonds est comparée pour le seuil de rentabilité : 187 bp = couche FPS annoncée au DIC des Parts A (1,871 %) ; 200 bp = plafond de commissions indirectes de l'Art. 9.5. Sans valeur contractuelle."},
    }
    for col in ('M', 'Q', 'U', 'Y'):
        m = VALO_RE.match(cell(inp, col + '61'))
        fees['valo']['paliers']['Mensuelle'].append(int(m.group(1)))
        fees['valo']['paliers']['Bi-mensuelle'].append(int(m.group(2)))
    m = re.match(r'^=IF\(F37>=(\d+),([\d.]+),IF\(F37>=(\d+),([\d.]+),0\)\)$', cell(inp, 'F38'))
    fees['valo']['majoration'] = [[int(m.group(1)), float(m.group(2))], [int(m.group(3)), float(m.group(4))]]

    # ---- 02_Classes_de_Parts
    classes = []
    for r in range(7, 19):
        classes.append({'nom': cls['C%d' % r].value or '', 'distrib_global': num(cls['D%d' % r].value) or 0,
                        'sous_distrib': num(cls['E%d' % r].value) or 0, 'distribuante': cls['F%d' % r].value or 'Non',
                        'commentaire': cls['G%d' % r].value or ''})
    fees['classes'] = classes

    # ---- 05_Tresorerie §3 : fréquences, règles spéciales
    freq_rows = {}
    for i in range(20):
        v = tre['D%d' % (20 + i)].value
        freq_rows[i] = ('derive' if isinstance(v, str) and v.startswith('=') else v)
    prorata_except = set()
    for i in range(20):
        f = tre['D%d' % (67 + i)].value or ''
        if '$E$10' not in f:
            prorata_except.add(i)
    assert prorata_except == {PRORATA_EXCEPTION_ROW - 67}, prorata_except
    distrib_rule = set()
    for i in range(20):
        f = tre['D%d' % (93 + i)].value or ''
        if 'SUMPRODUCT' not in f:
            distrib_rule.add(i)

    # ---- 01_Input §4 : tableau de frais (20 lignes)
    MODES = {'Barème (Forfait/Taux/Palier)': 'bareme', 'Taux Classe sélectionnée': 'classe',
             'Frais KYC (paliers de 100M€)': 'kyc', 'Résiduel vs Gestion': 'residuel'}
    lines = []
    for i in range(20):
        r = 53 + i
        lid = LINE_IDS[i]
        ex = EXTRAS.get(lid, {})
        recap_formula = recap['E%d' % (21 + i)].value or ''
        line = {
            'id': lid, 'n': i + 1,
            'actif': cell(inp, 'C%d' % r) or 'Non',
            'nom': cell(inp, 'D%d' % r) or '',
            'beneficiaire': cell(inp, 'E%d' % r) or ex.get('beneficiaire', ''),
            'recurrent': cell(inp, 'F%d' % r) or 'Oui',
            'mode': MODES[cell(inp, 'G%d' % r)],
            'classe_taux': 'sous' if 'Distrib_SousDistrib_Classe_Selectionnee*E$20' in recap_formula else 'global',
            'plancher': token(cell(inp, 'H%d' % r)),
            'plafond': token(cell(inp, 'I%d' % r)),
            'tranches': [],
            'source': cell(inp, 'AD%d' % r) or ex.get('source', ''),
            'tva': ex.get('tva', 0),
            'freq': freq_rows[i] if freq_rows[i] != 'derive' else 'Annuel',
            'regle_facturation': 'par_distribution' if i in distrib_rule else 'standard',
            'prorata_lancement': 'Non' if i in prorata_except else 'Oui',
        }
        for k in range(5):
            c0 = 10 + 4 * k  # J, N, R, V, Z
            cols = [openpyxl.utils.get_column_letter(c0 + j) for j in range(4)]
            line['tranches'].append({
                'inf': token(cell(inp, cols[0] + str(r))), 'sup': token(cell(inp, cols[1] + str(r))),
                'nature': cell(inp, cols[2] + str(r)) or 'Taux', 'valeur': token(cell(inp, cols[3] + str(r))),
            })
        lines.append(line)
    fees['lines'] = lines

    # ---- 05_Tresorerie §1, §2, §8
    aum = []
    c = 4
    while True:
        v = trev.cell(row=15, column=c).value
        if v is None or v == '':
            break
        aum.append(num(v))
        c += 1
    factures = []
    for r in range(161, 171):
        lab, amt, df, dp = tre['B%d' % r].value, tre['C%d' % r].value, tre['D%d' % r].value, tre['E%d' % r].value
        if (lab or amt) and amt:
            factures.append({'id': 'x%d' % (r - 160), 'label': lab or '', 'amount': num(amt), 'inv': iso(df), 'pay': iso(dp)})
    fees['treso'] = {
        'lancement': iso(tre['E5'].value), 'fin_exercice_1': iso(tre['E6'].value), 'horizon': int(tre['E7'].value),
        'aum_mensuel': aum,
        'aum_note': "Actif Net du Fonds mois par mois, en euros : réel pour les mois clos, prévisionnel pour les suivants. Dans le classeur, un mois sans Actif Net ne produit aucun montant ; ici la dernière valeur saisie est reconduite sur les mois suivants (affichée en grisé) pour que le prévisionnel couvre tout l'horizon.",
        'factures_constitution': factures,
    }

    # ---- 06_Documentation + 00_Demarrage (légende)
    docs = []
    for b, rows in (('01_INPUT - 4. FRAIS & PROVISIONS DU FONDS', (5, 6, 7, 8)), ('02_CLASSES_DE_PARTS', (12, 13, 14)),
                    ('03_DETAIL_DEPOSITAIRE_VALO', (18, 19, 20)), ('05_TRESORERIE - FRAIS DE CONSTITUTION', (24, 25))):
        docs.append({'title': b, 'items': [{'k': doc['B%d' % r].value, 'v': doc['C%d' % r].value} for r in rows]})
    fees['documentation'] = docs
    fees['ref_labels'] = {
        'seuil_t1': 'Seuil tranche 1 → 2 (§3D)', 'seuil_t2': 'Seuil tranche 2 → 3 (§3D)',
        'gestion_t1': 'Commission de Gestion, tranche 1 (§3D)', 'gestion_t2': 'Commission de Gestion, tranche 2 (§3D)', 'gestion_t3': 'Commission de Gestion, tranche 3 (§3D)',
        'conseil_t1': 'Commission de Conseil, tranche 1 (§3D)', 'conseil_t2': 'Commission de Conseil, tranche 2 (§3D)', 'conseil_t3': 'Commission de Conseil, tranche 3 (§3D)',
        'distrib_global_classe': 'Taux Distributeur Global de la classe sélectionnée (§3B)', 'distrib_sous_classe': 'Taux Sous-distributeurs de la classe sélectionnée (§3B)',
        'cac_distribution': 'Nombre de distributions × coût par distribution, si la classe distribue (§3)',
        'valo_p1': 'Palier CIC 0–40 M€ selon la fréquence de VL, majoration comprise (§3C)', 'valo_p2': 'Palier CIC 40–80 M€ (§3C)', 'valo_p3': 'Palier CIC 80–120 M€ (§3C)', 'valo_p4': 'Palier CIC > 120 M€ (§3C)',
        'valo_supplement': 'Supplément de 500 € par tranche de 10 lignes au-delà de 10 (§3C)',
        'tenue_passif_variable': 'Coût marginal par investisseur × nombre d’investisseurs (§3)',
        'constitution_montant': 'MAX(plancher ; taux × Actif Net au 1er anniversaire) (§3)', 'constitution_plancher': 'Plancher des Frais de Constitution (§3)',
    }
    fees['sensitivities'] = [
        {'id': 'gestion-1bp', 'label': '± 1 bp sur la Commission de Gestion (3 tranches)', 'ops': [{'path': 'baremes.gestion.%d' % k, 'delta': 0.0001} for k in range(3)]},
        {'id': 'conseil-5bp', 'label': '± 5 bp sur la Commission de Conseil (3 tranches)', 'ops': [{'path': 'baremes.conseil.%d' % k, 'delta': 0.0005} for k in range(3)]},
        {'id': 'depo-5k', 'label': '± 5 000 € sur le plancher du Dépositaire', 'ops': [{'path': 'lines.4.plancher', 'delta': 5000}]},
        {'id': 'cac-1k', 'label': '± 1 000 € sur le Commissaire aux Comptes', 'ops': [{'path': 'lines.5.tranches.0.valeur', 'delta': 1000}]},
        {'id': 'inv-20', 'label': '± 20 investisseurs au registre', 'ops': [{'path': 'params.nb_investisseurs', 'delta': 20}]},
    ]
    with open(dst, 'w', encoding='utf-8') as f:
        json.dump(fees, f, ensure_ascii=False, indent=1)
    print('lignes :', len(lines), '| classes :', sum(1 for c in classes if c['nom']), '| AUM mensuels :', aum, '| factures :', len(factures))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
